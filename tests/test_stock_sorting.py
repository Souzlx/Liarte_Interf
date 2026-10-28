"""Whole-recorte stock ordering, known values and identity-safe publication."""

from dataclasses import replace
from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory
from threading import Event
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from PySide6.QtCore import QThread, QUrl
from PySide6.QtTest import QSignalSpy

from liarte_core.inventory import SCALE, analyze_inventory
from liarte_core.inventory_io import write_snapshot
from liarte_core.query import checkpoint
from liarte_desktop import inventory as desktop
from tests.test_audit_desktop import expanded
from tests.test_inventory import sample, selection
from tests.test_qt import APP, pump


def position(code, *, quantity=1, unit="KG", cost=100, turnover=None,
             source_days=None, last_sale=None, branch="01", warehouse="01"):
    original = analyze_inventory(sample(), selection()).rows[0]
    return replace(
        original,
        balance=replace(original.balance, product=code, quantity=quantity * SCALE,
                        unit=unit, cost=cost, branch=branch, warehouse=warehouse),
        turnover=turnover, source_days=source_days, last_sale=last_sale,
    )


def source_with_giro(count=120):
    data = expanded(count)
    return replace(data, closings=tuple(
        replace(closing, turnover=float(index))
        for index, closing in enumerate(data.closings)
    ))


class StockPositionOrderingTests(unittest.TestCase):
    def test_unclassified_abc_sentinels_remain_after_known_classes_in_both_directions(self):
        rows = tuple(replace(position(code), abc=abc) for code, abc in (
            ("KNOWN_B", "B"), ("NO_CAPITAL", "—"), ("KNOWN_A", "A"),
            ("EMPTY", ""), ("KNOWN_C", "C"),
        ))
        for descending, expected in (
            (False, ["KNOWN_A", "KNOWN_B", "KNOWN_C", "EMPTY", "NO_CAPITAL"]),
            (True, ["KNOWN_C", "KNOWN_B", "KNOWN_A", "EMPTY", "NO_CAPITAL"]),
        ):
            with self.subTest(descending=descending):
                ordered = desktop._order_rows(rows, 7, descending)
                self.assertEqual([row.balance.product for row in ordered], expected)
                self.assertEqual(set(ordered), set(rows))

    def test_product_names_sort_without_case_or_accent_bias(self):
        rows = tuple(replace(position(code), name=name) for code, name in (
            ("A", "Zinco"), ("B", "ácido"), ("C", "Agua"), ("D", "ÁGUA"),
        ))
        for descending, expected in (
            (False, ["B", "C", "D", "A"]),
            (True, ["A", "C", "D", "B"]),
        ):
            with self.subTest(descending=descending):
                self.assertEqual(
                    [r.balance.product for r in desktop._order_rows(rows, 1, descending)], expected
                )

    def test_zero_and_known_source_values_lead_negative_and_missing_in_both_directions(self):
        rows = (
            position("TEN", turnover=10, source_days=10),
            position("TWO", turnover=2, source_days=2),
            position("ZERO", turnover=0, source_days=0),
            position("MISSING"),
            position("NEGATIVE", turnover=-2, source_days=-2),
        )
        for column in (9, 10):
            for descending, expected in (
                (False, ["ZERO", "TWO", "TEN", "NEGATIVE", "MISSING"]),
                (True, ["TEN", "TWO", "ZERO", "NEGATIVE", "MISSING"]),
            ):
                with self.subTest(column=column, descending=descending):
                    ordered = desktop._order_rows(rows, column, descending)
                    self.assertEqual([r.balance.product for r in ordered], expected)
                    self.assertEqual(set(ordered), set(rows))

    def test_unknown_last_sale_never_masquerades_as_an_old_date(self):
        rows = (
            position("UNKNOWN"),
            position("RECENT", last_sale=date(2026, 3, 30)),
            position("OLD", last_sale=date(2025, 1, 1)),
        )
        for descending, expected in (
            (False, ["OLD", "RECENT", "UNKNOWN"]),
            (True, ["RECENT", "OLD", "UNKNOWN"]),
        ):
            with self.subTest(descending=descending):
                self.assertEqual(
                    [r.balance.product for r in desktop._order_rows(rows, 8, descending)], expected
                )

    def test_quantity_order_keeps_units_in_separate_numeric_groups(self):
        rows = (
            position("KG_TWO", quantity=2),
            position("L_HUNDRED", quantity=100, unit="L"),
            position("KG_TEN", quantity=10),
            position("L_ONE", quantity=1, unit="L"),
        )
        for descending, expected in (
            (False, ["KG_TWO", "KG_TEN", "L_ONE", "L_HUNDRED"]),
            (True, ["KG_TEN", "KG_TWO", "L_HUNDRED", "L_ONE"]),
        ):
            with self.subTest(descending=descending):
                self.assertEqual(
                    [r.balance.product for r in desktop._order_rows(rows, 5, descending)], expected
                )

    def test_exact_cost_and_position_ties_do_not_depend_on_sort_direction_or_input_order(self):
        rows = (
            position("SAME", warehouse="02", unit="L", cost=2**53 + 2),
            position("SAME", warehouse="01", unit="L", cost=2**53 + 2),
            position("SAME", warehouse="01", unit="KG", cost=2**53 + 2),
            position("HIGH", cost=2**53 + 3),
        )
        for descending in (False, True):
            with self.subTest(descending=descending):
                ordered = desktop._order_rows(rows, 6, descending)
                self.assertEqual(ordered[0 if descending else -1].balance.product, "HIGH")
                ties = [r.balance for r in ordered if r.balance.product == "SAME"]
                self.assertEqual([(b.warehouse, b.unit) for b in ties],
                                 [("01", "KG"), ("01", "L"), ("02", "L")])
                self.assertEqual(ordered, desktop._order_rows(tuple(reversed(rows)), 6, descending))


class StockPositionSortingControllerTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.c = desktop.InventoryController()
        self.addCleanup(self.c.shutdown)

    def idle(self):
        self.assertTrue(pump(
            lambda: not self.c.presentation["busy"] and not self.c._debounce.isActive(), 15
        ), self.c.presentation)

    def load(self, data=None):
        path = Path(self.temp.name) / "source.liarte"
        write_snapshot(data or source_with_giro(), path)
        self.c.openFile(QUrl.fromLocalFile(str(path)))
        self.idle()
        self.assertTrue(self.c.presentation["loaded"], self.c.presentation["error"])

    def test_explicit_sort_uses_all_cached_positions_before_paging_without_reanalysis(self):
        self.load()
        self.c.setPage(2)
        before = self.c._analysis
        analysis_changes = QSignalSpy(self.c.analysisChanged)
        original = desktop._order_rows
        threads = []

        def order(*args, **kwargs):
            threads.append(QThread.currentThread() == APP.thread())
            return original(*args, **kwargs)

        with (
            patch.object(desktop, "_order_rows", side_effect=order),
            patch.object(desktop, "load_inventory", side_effect=AssertionError("Source reread")),
            patch.object(desktop, "analyze_inventory", side_effect=AssertionError("Reanalysis")),
            patch.object(desktop, "analyze_monthly_inventory", side_effect=AssertionError("Monthly reanalysis")),
        ):
            self.c.sortPositions(9, True)
            self.idle()
        self.assertEqual(threads, [False])
        self.assertIs(self.c._analysis, before)
        self.assertEqual(analysis_changes.count(), 0)
        self.assertEqual(self.c.presentation["page"], 1)
        self.assertEqual(self.c.table.rows[0].balance.product, "P000119")
        self.assertEqual([r.turnover for r in self.c.table.rows], list(range(119, 69, -1)))
        self.c.setPage(3)
        self.assertEqual([r.turnover for r in self.c.table.rows], list(range(19, -1, -1)))
        self.assertEqual((self.c.presentation["sortColumn"], self.c.presentation["descending"]), (9, True))

    def test_header_and_explicit_order_share_state_and_repeating_choice_is_a_noop(self):
        self.load()
        self.c.sortPositions(9, False)
        self.idle()
        self.assertEqual(self.c.table.rows[0].turnover, 0)
        self.c.setPage(2)
        with patch.object(desktop, "_order_rows", side_effect=AssertionError("Repeated sort")):
            self.c.sortPositions(9, False)
            for column, descending in ((-1, True), (11, True), (True, False), (9, "false")):
                self.c.sortPositions(column, descending)
        self.assertEqual(self.c.presentation["page"], 2)
        self.c.sortByColumn(9)
        self.idle()
        self.assertEqual(self.c.presentation["sortColumn"], 9)
        self.assertTrue(self.c.presentation["descending"])
        self.assertEqual(self.c.table.rows[0].turnover, 119)

    def test_sort_preserves_visible_position_identity_and_its_detail(self):
        self.load(sample())
        index = next(i for i, row in enumerate(self.c.table.rows)
                     if row.balance.product == "0001" and row.balance.warehouse == "02")
        self.c.selectRow(index)
        old_detail = self.c.detail
        old_balance = self.c.table.rows[index].balance
        self.c.sortPositions(5, False)
        self.assertEqual(self.c.presentation["selectedRow"], 0)
        self.assertEqual(self.c.table.rows[0].balance, old_balance)
        self.assertEqual(self.c.detail, old_detail)

    def test_sort_clears_selection_when_that_position_leaves_the_first_page(self):
        self.load()
        self.c.selectRow(0)
        self.assertEqual(self.c.detail["code"], "P000000")
        self.c.sortPositions(9, True)
        self.idle()
        self.assertEqual(self.c.presentation["selectedRow"], -1)
        self.assertEqual(self.c.detail, {})
        self.assertEqual(self.c.table.rows[0].balance.product, "P000119")

    def test_selection_changed_while_worker_sorts_remains_matched_to_its_detail(self):
        self.load()
        self.c.selectRow(0)
        entered, release = Event(), Event()
        original = desktop._order_rows

        def delayed(rows, column, descending, cancel=None):
            entered.set()
            self.assertTrue(release.wait(5))
            return original(rows, column, descending, cancel)

        try:
            with patch.object(desktop, "_order_rows", side_effect=delayed):
                self.c.sortPositions(9, False)
                self.assertTrue(pump(entered.is_set))
                self.c.selectRow(1)
                self.assertEqual(self.c.detail["code"], "P000001")
                release.set()
                self.idle()
        finally:
            release.set()
        selected = self.c.presentation["selectedRow"]
        self.assertEqual(self.c.table.rows[selected].balance.product, "P000001")
        self.assertEqual(self.c.detail["code"], "P000001")

    def test_cancelling_worker_sort_preserves_page_rows_selection_and_applied_order(self):
        self.load()
        self.c.setPage(2)
        self.c.selectRow(3)
        old_rows, old_detail = self.c.table.rows, self.c.detail
        entered = Event()

        def delayed(rows, column, descending, cancel=None):
            entered.set()
            self.assertTrue(cancel.wait(5))
            checkpoint(cancel)

        with patch.object(desktop, "_order_rows", side_effect=delayed):
            self.c.sortPositions(9, True)
            self.assertTrue(pump(entered.is_set))
            self.c.cancel()
            self.idle()
        self.assertEqual(self.c.presentation["page"], 2)
        self.assertEqual(self.c.table.rows, old_rows)
        self.assertEqual(self.c.presentation["selectedRow"], 3)
        self.assertEqual(self.c.detail, old_detail)
        self.assertEqual((self.c.presentation["sortColumn"], self.c.presentation["descending"]), (6, True))

    def test_excel_positions_follow_full_order_even_when_a_later_page_is_visible(self):
        self.load()
        self.c.sortPositions(9, True)
        self.idle()
        self.c.setPage(2)
        expected = [r.balance.product for r in self.c._ordered]
        output = Path(self.temp.name) / "ordered.xlsx"
        self.c.exportExcel(QUrl.fromLocalFile(str(output)))
        self.idle()
        self.assertTrue(output.is_file(), self.c.presentation["error"])
        book = load_workbook(output, read_only=True, data_only=True)
        try:
            codes = [row[0] for row in book["Estoque"].iter_rows(min_row=4, values_only=True)]
        finally:
            book.close()
        self.assertEqual(codes, expected)
        self.assertEqual(len(codes), 120)
        self.assertEqual(codes[0], "P000119")
        self.assertEqual(codes[-1], "P000000")
        self.assertEqual(self.c.presentation["page"], 2)


if __name__ == "__main__":
    unittest.main()
