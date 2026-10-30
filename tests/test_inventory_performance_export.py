"""Product exports keep full rankings, units, missing facts and investigation rules."""

from dataclasses import replace
from pathlib import Path
from tempfile import TemporaryDirectory
from threading import Event
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from liarte_core.inventory import SCALE, StockSelection, analyze_inventory
from liarte_core.inventory_export import export_inventory
from liarte_core.inventory_monthly import analyze_monthly_inventory
from liarte_core.inventory_performance import analyze_product_performance, product_performance_detail
from liarte_core.inventory_performance_export import append_performance_sheets
from liarte_core.query import Cancelled
from tests.test_inventory_monthly import monthly_sample


class InventoryPerformanceExportTests(unittest.TestCase):
    def setUp(self):
        self.data = monthly_sample()
        self.selection = StockSelection("2026-01", "2026-02")
        self.analysis = analyze_inventory(self.data, self.selection)
        self.months = analyze_monthly_inventory(self.data, self.selection, self.analysis)
        self.report = analyze_product_performance(self.months, self.selection, window=1)

    def test_complete_export_preserves_units_missing_stock_and_selected_indicators(self):
        detail = product_performance_detail(
            self.months, self.selection, "A", "KG", window=1, report=self.report
        )
        with TemporaryDirectory() as directory:
            path = Path(directory) / "report.xlsx"
            export_inventory(
                self.data, self.analysis, self.selection, path, Event(),
                monthly=self.months, performance=self.report, performance_detail=detail,
            )
            book = load_workbook(path)
            try:
                sheet = book["Desempenho dos produtos"]
                values = list(sheet.values)
                rows = [dict(zip(values[0], row)) for row in values[1:]]
                self.assertEqual(len(rows), len(self.report["rows"]))
                keys = {(row["Código"], row["UM"]): row for row in rows}
                a = keys["A", "KG"]
                self.assertEqual(a["Vendas líquidas recentes"], 20)
                self.assertEqual(a["Variação (%)"], 100)
                self.assertEqual(a["Quantidade no fechamento"], 170)
                self.assertEqual(a["Estoque a custo (R$)"], 205)
                self.assertIsNone(keys["C", "KG"]["Quantidade no fechamento"])
                self.assertIn(("B", "L"), keys)
                self.assertEqual(a["Data da base"].date(), self.data.as_of)
                self.assertIn("start=2026-01", a["Filtros"])
                self.assertIn("não mede idade", a["Definição e limites"])
                investigation = list(book["Investigação do produto"].values)
                self.assertEqual(investigation[0], (
                    "Indicador", "Valor", "Unidade", "Períodos (AAAA-MM)", "Disponibilidade"
                ))
                items = {row[0]: row for row in investigation[1:]}
                self.assertEqual(items["Vendas líquidas · recente"][1:4], (20, "KG", "2026-02"))
                self.assertEqual(items["Variação das vendas líquidas"][1:3], (100, "%"))
                self.assertEqual(items["Estoque a custo"][1:3], (205, "R$"))
                self.assertNotIn("Hipótese a verificar", items)
                self.assertNotIn("O que os dados mostram", items)
                self.assertTrue(any("idade do lote" in str(row[1]) for row in investigation))
            finally:
                book.close()

    def test_no_performance_keeps_existing_export_contract(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "old.xlsx"
            export_inventory(self.data, self.analysis, self.selection, path, Event(), monthly=self.months)
            book = load_workbook(path)
            try:
                self.assertNotIn("Desempenho dos produtos", book.sheetnames)
                self.assertNotIn("Investigação do produto", book.sheetnames)
            finally:
                book.close()

    def test_precision_formula_text_and_splitting_preserve_all_rows(self):
        report = dict(self.report, rows=tuple(
            dict(row, code="=1+1", name="@Produto", sold=10**22 + 1)
            for row in self.report["rows"]
        ))
        with TemporaryDirectory() as directory:
            path = Path(directory) / "exact.xlsx"
            with patch("liarte_core.inventory_performance_export._MAX_ROWS", 3):
                export_inventory(
                    self.data, self.analysis, self.selection, path, Event(),
                    monthly=self.months, performance=report,
                )
            book = load_workbook(path)
            try:
                names = [name for name in book.sheetnames if name.startswith("Desempenho dos produtos")]
                self.assertGreater(len(names), 1)
                self.assertEqual(sum(book[name].max_row - 1 for name in names), len(report["rows"]))
                for name in names:
                    for row in book[name].iter_rows(min_row=2):
                        self.assertEqual(row[0].data_type, "s")
                        self.assertEqual(row[6].value, "10000000000000000.000001")
                        self.assertIsNotNone(row[6].comment)
            finally:
                book.close()

    def test_cancel_does_not_publish_or_replace_destination(self):
        cancel = Event()
        cancel.set()
        with TemporaryDirectory() as directory:
            path = Path(directory) / "report.xlsx"
            path.write_bytes(b"previous report")
            with self.assertRaises(Cancelled):
                export_inventory(
                    self.data, self.analysis, self.selection, path, cancel,
                    monthly=self.months, performance=self.report,
                )
            self.assertEqual(path.read_bytes(), b"previous report")

    def test_negative_and_unavailable_comparisons_are_not_excel_zeroes(self):
        data = replace(self.data, sales=tuple(
            replace(sale, quantity=-SCALE) if sale.product == "A" and sale.day.month == 2 else sale
            for sale in self.data.sales
        ))
        months = analyze_monthly_inventory(data, self.selection)
        report = analyze_product_performance(months, self.selection, window=1)
        book = Workbook()
        append_performance_sheets(book, report, None, data, self.selection, Event())
        row = next(row for row in book["Desempenho dos produtos"].values if row[0] == "A")
        self.assertLess(row[6], 0)
        self.assertIsNone(row[8])

    def test_numeric_investigation_preserves_exact_costs_and_literal_source_text(self):
        detail = product_performance_detail(
            self.months, self.selection, "A", "KG", window=1, report=self.report
        )
        detail = dict(detail, code="=1+1", name="@Produto", indicators=tuple(
            dict(item, value=10**22 + 1) if item["key"] == "stockCost" else item
            for item in detail["indicators"]
        ))
        book = Workbook()
        append_performance_sheets(book, self.report, detail, self.data, self.selection, Event())
        sheet = book["Investigação do produto"]
        rows = {row[0].value: row for row in sheet.iter_rows(min_row=2)}
        self.assertEqual(rows["Código"][1].value, "=1+1")
        self.assertEqual(rows["Código"][1].data_type, "s")
        self.assertEqual(rows["Produto"][1].value, "@Produto")
        self.assertEqual(rows["Estoque a custo"][1].value, "100000000000000000000.01")
        self.assertIsNotNone(rows["Estoque a custo"][1].comment)
        self.assertEqual(rows["Estoque a custo"][1].number_format, "@")

    def test_numeric_investigation_keeps_unknown_values_blank_and_known_zero_numeric(self):
        detail = product_performance_detail(
            self.months, self.selection, "A", "KG", window=1, report=self.report
        )
        detail = dict(detail, indicators=tuple(
            dict(item, value=None, reason="Sem fechamento conhecido.") if item["key"] == "stockCost"
            else dict(item, value=0) if item["key"] == "sold" else item
            for item in detail["indicators"]
        ))
        book = Workbook()
        append_performance_sheets(book, self.report, detail, self.data, self.selection, Event())
        values = {row[0]: row for row in book["Investigação do produto"].values}
        self.assertIsNone(values["Estoque a custo"][1])
        self.assertEqual(values["Estoque a custo"][4], "Sem fechamento conhecido.")
        self.assertEqual(values["Vendas líquidas · recente"][1], 0)
