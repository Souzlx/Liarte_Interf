"""Regressions for audited XLSX bounds, publication and exact totals."""

from dataclasses import replace
from datetime import date
from decimal import Decimal, localcontext
import io
import json
from pathlib import Path
import tempfile
from threading import Event
import unittest
from unittest.mock import PropertyMock, patch
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from liarte_core.inventory import (
    Balance,
    Closing,
    InventoryData,
    Product,
    Sale,
    StockSelection,
    analyze_inventory,
)
from liarte_core.inventory_export import _append_monthly_sheets, export_inventory
from liarte_core.inventory_monthly import analyze_monthly_inventory
from liarte_engine import liarte_dados as data
from liarte_engine import powerbi_importacoes as bi
from liarte_engine import powerbi_source as source
from liarte_engine.comexstat_importacoes import ProductNcm, to_number
from liarte_engine.liarte_analise import table_page
from tests import test_powerbi
from tests.test_inventory_monthly_export import monthly_export_data


class AuditInventoryExportTests(unittest.TestCase):
    def test_monthly_append_does_not_rescan_preceding_cells_for_row_numbers(self):
        dataset = monthly_export_data()
        month = analyze_monthly_inventory(dataset, StockSelection("2026-01", "2026-02"))[1]
        month = {**month, "stockRows": month["stockRows"] * 400}
        book = Workbook()
        try:
            with patch.object(
                Worksheet,
                "max_row",
                new_callable=PropertyMock,
                side_effect=AssertionError("quadratic row scan"),
            ):
                _append_monthly_sheets(book, [month], Event())
            self.assertEqual(
                book["Produtos mensais"]._current_row,
                len(month["stockRows"]) + len(month["saleRows"]) + 1,
            )
        finally:
            book.close()

    def test_invalid_xml_text_preserves_previous_export(self):
        selection = StockSelection("2026-01", "2026-04")
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "stock.xlsx"
            for character in ("\x01", "\x0b", "\ufffe", "\uffff", "\ud800"):
                with self.subTest(character=repr(character)):
                    original = monthly_export_data()
                    dataset = replace(
                        original,
                        products=(
                            replace(original.products[0], name="Produto" + character),
                            *original.products[1:],
                        ),
                    )
                    path.write_bytes(b"previous export")
                    with self.assertRaisesRegex(ValueError, "XML/Excel"):
                        export_inventory(
                            dataset, analyze_inventory(dataset, selection), selection, path, Event()
                        )
                    self.assertEqual(path.read_bytes(), b"previous export")

    def test_extreme_scaled_values_are_exact_labelled_text_in_current_and_monthly_sheets(self):
        quantity = cost = 10**16 + 1
        dataset = InventoryData(
            date(2026, 1, 31),
            "Synthetic precision",
            "a" * 64,
            (Product("A", "Produto", "Grupo", "MP", "KG"),),
            (Balance("A", "01", "01", "KG", quantity, cost),),
            (Closing("A", "01", "01", "KG", "2026-01", quantity, cost, 1, 30),),
            (Sale("A", "01", "KG", date(2026, 1, 5), quantity, cost),),
        )
        selection = StockSelection("2026-01", "2026-01")
        with tempfile.TemporaryDirectory() as directory, localcontext() as context:
            context.prec = 6
            path = Path(directory) / "exact.xlsx"
            export_inventory(
                dataset, analyze_inventory(dataset, selection), selection, path, Event()
            )
            book = load_workbook(path)
            try:
                for cell, expected in (
                    (book["Estoque"]["G4"], "10000000000.000001"),
                    (book["Estoque"]["H4"], "100000000000000.01"),
                    (book["Análise mensal"]["D2"], "100000000000000.01"),
                    (book["Produtos mensais"]["H2"], "10000000000.000001"),
                ):
                    self.assertEqual(cell.value, expected)
                    self.assertEqual(cell.data_type, "s")
                    self.assertEqual(cell.number_format, "@")
                    self.assertIn("15 dígitos", cell.comment.text)
                rules = dict(book["Leitura dos indicadores"].iter_rows(min_row=2, values_only=True))
                self.assertIn("somas automáticas", rules["Precisão numérica do Excel"])
            finally:
                book.close()

    def test_failed_staged_validation_preserves_previous_export(self):
        dataset = monthly_export_data()
        selection = StockSelection("2026-01", "2026-04")
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "stock.xlsx"
            path.write_bytes(b"last good workbook")
            with patch(
                "liarte_core.inventory_export._validate_staged_workbook",
                side_effect=ValueError("bad XML"),
            ):
                with self.assertRaisesRegex(ValueError, "bad XML"):
                    export_inventory(
                        dataset, analyze_inventory(dataset, selection), selection, path, Event()
                    )
            self.assertEqual(path.read_bytes(), b"last good workbook")
            self.assertEqual(list(Path(directory).iterdir()), [path])


class AuditSourceWorkbookTests(unittest.TestCase):
    def setUp(self):
        fixture = test_powerbi.PowerBiTests(methodName="runTest")
        fixture.setUp()
        self.addCleanup(fixture.tearDown)
        self.excel, self.output = fixture.excel, fixture.output

    def rewrite(self, path, transform):
        with zipfile.ZipFile(self.excel) as original:
            entries = [(info, original.read(info.filename)) for info in original.infolist()]
        with zipfile.ZipFile(self.excel, "w", zipfile.ZIP_DEFLATED) as output:
            for info, content in entries:
                output.writestr(info, transform(content) if info.filename == path else content)

    def rejected_without_publication(self):
        self.output.write_bytes(b"previous BI package")
        with self.assertRaises(bi.PowerBIProjectError):
            bi.create_powerbi_package(self.excel, self.output, "2026-01", "2026-03")
        self.assertEqual(self.output.read_bytes(), b"previous BI package")

    def test_sparse_reference_outside_excel_limits_is_rejected_before_dense_allocation(self):
        self.rewrite("xl/worksheets/sheet1.xml", lambda raw: raw.replace(b'r="M2"', b'r="ZZZZZZ2"'))
        self.rejected_without_publication()

    def test_valid_last_excel_column_never_inflates_consumed_row(self):
        self.rewrite("xl/worksheets/sheet1.xml", lambda raw: raw.replace(b'r="M2"', b'r="XFD2"'))
        manifest = bi._validate_excel_source(self.excel)
        with zipfile.ZipFile(self.excel) as archive:
            rows = list(source._rows(archive, manifest["general"][0][1], [], None))
        self.assertEqual(len(rows[1][1]), 12)
        self.assertTrue(all(len(row) <= 13 for _, row in rows))

    def test_corrupt_content_types_is_rejected_without_replacing_previous_package(self):
        self.rewrite("[Content_Types].xml", lambda _: b"corrupt")
        self.rejected_without_publication()

    def test_workbook_relationship_to_missing_part_is_rejected(self):
        self.rewrite(
            "xl/_rels/workbook.xml.rels",
            lambda raw: raw.replace(b"/xl/worksheets/sheet1.xml", b"/xl/worksheets/missing.xml"),
        )
        self.rejected_without_publication()

    def test_duplicate_cell_reference_is_rejected(self):
        self.rewrite("xl/worksheets/sheet1.xml", lambda raw: raw.replace(b'r="M2"', b'r="L2"'))
        self.rejected_without_publication()

    def test_shared_string_negative_index_is_rejected(self):
        raw = (
            f'<worksheet xmlns="{source.NS[1:-1]}"><sheetData><row r="1">'
            '<c r="A1" t="s"><v>-1</v></c></row></sheetData></worksheet>'
        ).encode()
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("sheet.xml", raw)
        with zipfile.ZipFile(buffer) as archive, self.assertRaisesRegex(ValueError, "Índice"):
            list(source._rows(archive, "sheet.xml", ["last string"], None))

    def test_oversized_metadata_budget_is_checked_before_parsing(self):
        with patch.object(source, "MAX_METADATA_BYTES", 1):
            self.rejected_without_publication()

    def test_utf16_entity_declaration_is_rejected(self):
        xml = (
            '<?xml version="1.0" encoding="utf-16"?><!DOCTYPE worksheet '
            '[<!ENTITY x "expansion">]><worksheet xmlns="' + source.NS[1:-1] + '"/>'
        )
        self.rewrite("xl/worksheets/sheet1.xml", lambda _: xml.encode("utf-16"))
        self.rejected_without_publication()


class AuditExactTotalsTests(unittest.TestCase):
    @staticmethod
    def row(month="01", value="9007199254740992.01"):
        return {
            "coNcm": "29161220",
            "year": "2026",
            "monthNumber": month,
            "country": "China",
            "state": "SP",
            "urf": "Porto",
            "metricFOB": value,
            "metricFreight": "0",
            "metricInsurance": "0",
            "metricCIF": value,
            "metricKG": "9007199254740993",
        }

    def manifest(self, directory, rows):
        path = Path(directory) / "source.xlsx"
        path.write_bytes(b"hash-only manifest fixture")
        result = data.create_manifest(
            excel_path=path,
            powerbi_path=None,
            version="audit",
            catalog_version="audit",
            pairs=[ProductNcm("Produto", "29161220")],
            start="2026-01",
            end="2026-02",
            general_rows=rows,
            municipal_rows=[],
            number=to_number,
        )
        return json.loads(result.read_bytes())

    def test_manifest_preserves_large_exact_cents_and_integer_kg_in_small_decimal_context(self):
        with tempfile.TemporaryDirectory() as directory, localcontext() as context:
            context.prec = 6
            manifest = self.manifest(directory, [self.row()])
        totals = manifest["general"]["totals"]
        self.assertEqual(totals["metricCIF"], "9007199254740992.01")
        self.assertEqual(totals["metricKG"], 9007199254740993)
        self.assertEqual(manifest["schema"], 2)
        self.assertIn("decimal strings", manifest["numeric_encoding"])

    def test_usual_manifest_totals_remain_json_numbers(self):
        with tempfile.TemporaryDirectory() as directory:
            totals = self.manifest(directory, [self.row(value="12.34")])["general"]["totals"]
        self.assertEqual(totals["metricCIF"], 12.34)
        self.assertIsInstance(totals["metricKG"], int)

    def test_one_cent_comparison_survives_above_float_precision(self):
        with localcontext() as context:
            context.prec = 6
            compared = data.period_comparisons(
                [self.row(value="9007199254740992.00"), self.row("02")],
                "2026-01",
                "2026-02",
                to_number,
            )["monthly"]
        metric = compared["metrics"]["cif"]
        self.assertEqual(metric["change"], 0.01)
        self.assertEqual(metric["current"], "9007199254740992.01")
        self.assertGreater(metric["ratio"], 0)

    def test_decimal_table_sort_is_numeric_and_keeps_missing_values_last(self):
        rows = [{"Valor": value} for value in (Decimal("2"), None, Decimal("10"), Decimal("1.5"))]
        result = table_page(rows, ["Valor"], sort="Valor")
        self.assertEqual(
            [row["Valor"] for row in result["rows"]],
            [Decimal("1.5"), Decimal("2"), Decimal("10"), None],
        )
