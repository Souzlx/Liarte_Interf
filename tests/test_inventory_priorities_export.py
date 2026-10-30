"""Priority exports preserve full product/unit evidence and exact workbook values."""

from copy import deepcopy
from dataclasses import fields
from pathlib import Path
from tempfile import TemporaryDirectory
from threading import Event
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from liarte_core.inventory import SCALE, StockSelection
from liarte_core.inventory_priorities import analyze_product_priorities
from liarte_core.inventory_priorities_export import export_product_priorities
from liarte_core.query import Cancelled


def priority_report():
    selection = StockSelection("2026-01", "2026-02", branch="01")
    return analyze_product_priorities({
        "available": True, "window": 1, "units": ("KG", "L", "UN"),
        "selectionStart": selection.start, "selectionEnd": selection.end,
        "selectionKey": tuple(getattr(selection, field.name) for field in fields(selection)),
        "baseLabel": "01/2026", "targetLabel": "02/2026", "targetPeriod": "2026-02",
        "scope": "Filial 01; unidades separadas.",
        "notice": "Janelas completas, com diferentes números de dias.", "reason": "",
        "rows": (
            {"code": "A", "name": "Produto A", "unit": "KG", "stockQuantity": 10 * SCALE,
             "stockCost": 12345, "sold": 0, "previous": 5 * SCALE, "change": -5 * SCALE,
             "percentage": -100, "coverage": None, "coverageReason": "Sem demanda líquida positiva.",
             "turnover": 0, "status": "no_positive", "negativeStockPositionCount": 0},
            {"code": "B", "name": "Produto B", "unit": "L", "stockQuantity": 20 * SCALE,
             "stockCost": 25000, "sold": None, "previous": None, "change": None,
             "percentage": None, "coverage": None, "coverageReason": "Vendas não observadas.",
             "turnover": None, "status": "insufficient", "negativeStockPositionCount": 0},
            {"code": "A", "name": "Produto A em unidades", "unit": "UN", "stockQuantity": 0,
             "stockCost": 0, "sold": 2 * SCALE, "previous": 4 * SCALE, "change": -2 * SCALE,
             "percentage": -50, "coverage": 0, "coverageReason": "Mesmas filiais.",
             "turnover": 2, "status": "down", "negativeStockPositionCount": 0},
        ),
    }, coverage_days=30)


class InventoryPrioritiesExportTests(unittest.TestCase):
    def setUp(self):
        self.report = priority_report()
        self.book = Workbook()
        self.addCleanup(self.book.close)

    def rows(self, sheet=None):
        sheet = sheet or self.book["Prioridades dos produtos"]
        values = list(sheet.values)
        return [dict(zip(values[0], row, strict=True)) for row in values[1:]]

    def test_all_units_export_once_per_identity_with_every_reason_in_the_same_row(self):
        original = deepcopy(self.report)
        layouts = export_product_priorities(self.book, self.report)
        rows = self.rows()
        self.assertEqual(len(rows), 3)
        self.assertEqual({(r["Código"], r["UM"]) for r in rows},
                         {("A", "KG"), ("A", "UN"), ("B", "L")})
        a = next(row for row in rows if row["Código"] == "A" and row["UM"] == "KG")
        self.assertIn("Vendas em queda", a["Situações para revisar"])
        self.assertIn("Saldo sem venda líquida", a["Situações para revisar"])
        self.assertIn("Maior capital", a["Situações para revisar"])
        self.assertIn("devoluções", a["Motivos observados"])
        self.assertIn("uso interno", a["Próxima verificação"])
        self.assertEqual(a["Quantidade no fechamento"], 10)
        self.assertEqual(a["Estoque a custo (R$)"], 123.45)
        self.assertEqual(a["Vendas líquidas anteriores"], 5)
        self.assertEqual(a["Diferença nas vendas"], -5)
        self.assertEqual(self.report, original)
        self.assertEqual(len(layouts["Prioridades dos produtos"]), 25)

    def test_known_zero_remains_numeric_and_missing_metrics_remain_empty(self):
        export_product_priorities(self.book, self.report)
        rows = {(r["Código"], r["UM"]): r for r in self.rows()}
        a, b, zero = rows["A", "KG"], rows["B", "L"], rows["A", "UN"]
        self.assertEqual(a["Vendas líquidas recentes"], 0)
        self.assertEqual(a["Giro fonte · mediana (x/mês)"], 0)
        self.assertEqual(zero["Quantidade no fechamento"], 0)
        self.assertEqual(zero["Estoque a custo (R$)"], 0)
        self.assertEqual(zero["Cobertura estimada (dias)"], 0)
        for name in ("Vendas líquidas recentes", "Vendas líquidas anteriores", "Diferença nas vendas",
                     "Variação (%)", "Cobertura estimada (dias)", "Giro fonte · mediana (x/mês)"):
            self.assertIsNone(b[name], name)

    def test_reference_windows_scope_threshold_and_limits_are_traceable(self):
        export_product_priorities(self.book, self.report)
        for row in self.rows():
            self.assertEqual(row["Mês do fechamento"], "2026-02")
            self.assertEqual(row["Janela anterior"], "01/2026")
            self.assertEqual(row["Janela recente"], "02/2026")
            self.assertEqual(row["Período do recorte"], "2026-01 a 2026-02")
            self.assertIn("branch=01", row["Fonte, base e filtros"])
            self.assertEqual(row["Limite de cobertura (dias)"], 30)
            self.assertEqual(row["Abrangência"], self.report["scope"])
            self.assertEqual(row["Comparação disponível"], "Sim")
            for term in ("não é faturamento", "não mede idade", "não comprova ausência de vendas",
                         "sem os limites de busca", "não devem ser somadas"):
                self.assertIn(term, row["Definição e limites"])

    def test_optional_source_base_date_and_explicit_filters_stay_in_each_exported_row(self):
        report = dict(self.report, sourceName="Base sintética.pbix", asOf="2026-03-15",
                      filters="branch=01; unit=KG")
        export_product_priorities(self.book, report)
        for row in self.rows():
            reference = row["Fonte, base e filtros"]
            self.assertIn("Fonte: Base sintética.pbix", reference)
            self.assertIn("Data da base: 2026-03-15", reference)
            self.assertIn("Filtros: branch=01; unit=KG", reference)

    def test_exact_values_and_formula_like_text_survive_save_and_splitting(self):
        rows = tuple(dict(row, code="=1+1", name="@Produto", sold=10**22 + 1,
                          stockCost=10**18 + 1) for row in self.report["rows"])
        report = dict(self.report, rows=rows)
        with patch("liarte_core.inventory_priorities_export._MAX_ROWS", 3):
            layouts = export_product_priorities(self.book, report)
        self.assertEqual(set(layouts), {"Prioridades dos produtos", "Prioridades dos produtos 2"})
        with TemporaryDirectory() as directory:
            path = Path(directory) / "priorities.xlsx"
            self.book.save(path)
            reopened = load_workbook(path)
            try:
                names = [name for name in reopened.sheetnames if name in layouts]
                self.assertEqual(sum(reopened[name].max_row - 1 for name in names), len(rows))
                for name in names:
                    for row in reopened[name].iter_rows(min_row=2):
                        self.assertEqual(row[0].value, "=1+1")
                        self.assertEqual(row[0].data_type, "s")
                        self.assertEqual(row[1].data_type, "s")
                        self.assertEqual(row[12].value, "10000000000000000.000001")
                        self.assertEqual(row[8].value, "10000000000000000.01")
                        for cell in (row[8], row[12]):
                            self.assertEqual(cell.number_format, "@")
                            self.assertIsNotNone(cell.comment)
            finally:
                reopened.close()

    def test_empty_result_still_records_headers_and_context_without_inventing_a_product(self):
        report = dict(self.report, rows=(), reason="Histórico insuficiente.")
        export_product_priorities(self.book, report)
        sheet = self.book["Prioridades dos produtos"]
        self.assertEqual(sheet.max_row, 1)
        self.assertEqual(sheet.max_column, 25)
        self.assertEqual(sheet.auto_filter.ref, "A1:Y1")
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertIn("Histórico insuficiente", sheet.cell(1, 25).comment.text)
        self.assertIn("02/2026", sheet.cell(1, 25).comment.text)

    def test_pre_cancelled_export_does_not_add_a_partial_sheet(self):
        cancel = Event()
        cancel.set()
        before = list(self.book.sheetnames)
        with self.assertRaises(Cancelled):
            export_product_priorities(self.book, self.report, cancel=cancel)
        self.assertEqual(self.book.sheetnames, before)

    def test_cancellation_between_products_stops_the_writer(self):
        cancel = Event()

        def rows():
            yield self.report["rows"][0]
            cancel.set()
            yield self.report["rows"][1]

        with self.assertRaises(Cancelled):
            export_product_priorities(self.book, dict(self.report, rows=rows()), cancel=cancel)
        self.assertEqual(self.book["Prioridades dos produtos"].max_row, 2)

    def test_invalid_selection_reference_is_rejected_before_any_sheet_is_created(self):
        with self.assertRaisesRegex(ValueError, "referência dos filtros"):
            export_product_priorities(self.book, dict(self.report, selectionKey=("2026-01",)))
        self.assertEqual(self.book.sheetnames, ["Sheet"])
