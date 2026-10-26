"""Excel de estoque do recorte exibido, com unidades e regras explícitas."""

from __future__ import annotations

import os
import math
import tempfile
from dataclasses import asdict
from datetime import date
from decimal import Decimal
import zipfile
from pathlib import Path
from threading import Event

from .inventory import SCALE, InventoryData, StockAnalysis, StockSelection
from .inventory_alerts import analyze_stock_alerts
from .query import checkpoint

HEADERS = (
    "Código",
    "Produto",
    "Grupo",
    "Filial",
    "Arm.",
    "UM",
    "Saldo",
    "Custo (R$)",
    "ABC",
    "Últ. venda",
    "Dias sem venda",
    "Giro fonte (x/mês)",
    "Dias fonte",
    "Mês giro",
    "Cobertura filial (d)",
    "CV vendas mensais",
)

_ALERTS_MAX_ROWS = 1_048_576


class _ExactText(str):
    """Decimal text used only beyond Excel's 15 significant numeric digits."""


def _append_values(sheet, values, row):
    # openpyxl converts str subclasses to plain str. Mark exact values before
    # binding them, so their warning survives the normal formatting pass.
    from openpyxl.comments import Comment
    from .inventory_io import _text

    for index, value in enumerate(values, 1):
        if isinstance(value, str):
            _text(value, f"{sheet.title}!linha {row}, coluna {index}", empty=True, max_length=32767)
    sheet.append(values)
    for index, value in enumerate(values, 1):
        if isinstance(value, _ExactText):
            sheet.cell(row, index).comment = Comment(
                "Valor decimal exato em texto: excede os 15 dígitos significativos "
                "do Excel. Unidade conforme o cabeçalho.",
                "LIARTE",
            )


def _excel_scaled(value, scale):
    if value is None:
        return None
    places = len(str(scale)) - 1
    exact = Decimal((int(value < 0), tuple(map(int, str(abs(value)))), -places))
    significant = len(str(abs(value)).rstrip("0"))
    if significant > 15:
        return _ExactText(format(exact, "f"))
    return float(exact)


def _validate_staged_workbook(path, cancel):
    from openpyxl import load_workbook
    from liarte_engine.powerbi_source import validate_xlsx_container, read_shared_strings, _rows

    def check():
        checkpoint(cancel)

    with zipfile.ZipFile(path) as archive:
        sheets = validate_xlsx_container(archive, check)
        strings = read_shared_strings(archive, check)
        for part in sheets.values():
            for _ in _rows(archive, part, strings, check, 25, allow_formulas=True):
                pass
    check()
    # Read-only reopen also checks Excel's workbook/style reader contract. XML
    # rows were already checked above without trusting dense sheet dimensions.
    reopened = load_workbook(path, read_only=True, data_only=False)
    try:
        if reopened.sheetnames != list(sheets):
            raise ValueError("As abas do Excel salvo divergem das abas validadas.")
        for sheet in reopened:
            check()
            next(sheet.iter_rows(min_row=1, max_row=1, max_col=25, values_only=True), None)
    finally:
        reopened.close()


def _append_monthly_sheets(workbook, months, cancel):
    """Exporta todas as linhas mensais, sem limitar às classificações da tela.

    Resumos monetários ocupam uma única linha por mês; as quantidades ficam em
    linhas distintas por unidade para que não exista um total misturando KG/L/UN.
    """
    money_format = '"R$" #,##0.00;[Red]("R$" #,##0.00)'
    quantity_format = "#,##0.000;[Red](#,##0.000)"
    layouts = {}

    class Table:
        def __init__(self, title, headers, widths):
            self.title, self.headers, self.widths = title, headers, widths
            self.part = 0
            self.sheet = None
            self.new_sheet()

        def new_sheet(self):
            self.part += 1
            title = self.title if self.part == 1 else self.title + " " + str(self.part)
            self.sheet = workbook.create_sheet(title)
            self.sheet.append(self.headers)
            self.row_count = 1
            layouts[title] = self.widths

        def append(self, values, formats=()):
            checkpoint(cancel)
            # A base aceita até dois milhões de fatos. Dividir antes do limite
            # do XLSX conserva todas as linhas, inclusive nas bases maiores.
            if self.row_count >= 1_048_576:
                self.new_sheet()
            self.row_count += 1
            _append_values(self.sheet, values, self.row_count)
            for column, fmt in formats:
                self.sheet.cell(self.row_count, column).number_format = fmt

    summary = Table(
        "Análise mensal",
        (
            "Mês",
            "Registro",
            "UM",
            "Estoque a custo (R$)",
            "Quantidade no fechamento",
            "Vendas das filiais (quantidade)",
            "Giro fonte · mediana (x/mês)",
            "Posições elegíveis para giro",
            "Giros inválidos · saldo positivo",
            "Produtos no fechamento",
            "Variação do custo · posições pareadas (R$)",
            "Variação total do custo (R$)",
            "Variação das vendas (quantidade)",
            "Vendas em mês parcial",
            "Posições novas",
            "Posições ausentes",
            "Abrangência dos fechamentos",
            "Abrangência das vendas",
            "Comparação do estoque",
            "Comparação das vendas",
            "Observações",
            "Fechamento em mês parcial",
            "Vendas observadas",
            "Custo dos registros novos (R$)",
            "Custo dos registros ausentes (R$)",
        ),
        (
            15,
            26,
            10,
            25,
            25,
            27,
            27,
            22,
            26,
            24,
            29,
            28,
            27,
            23,
            20,
            20,
            48,
            48,
            54,
            54,
            68,
            26,
            23,
            29,
            30,
        ),
    )
    changes = Table(
        "Variações mensais",
        ("Mês", "Comparação", "Análise", "UM", "Código", "Produto", "Variação", "Contexto"),
        (15, 20, 33, 10, 18, 52, 27, 90),
    )
    products = Table(
        "Produtos mensais",
        (
            "Mês",
            "Registro",
            "Código",
            "Produto",
            "Filial",
            "Armazém",
            "UM",
            "Quantidade no fechamento",
            "Vendas das filiais (quantidade)",
            "Estoque a custo (R$)",
            "Giro fonte (x/mês)",
            "Dias fonte",
            "Última venda observada",
            "Dias sem venda observada",
            "Observações",
        ),
        (15, 26, 18, 52, 14, 14, 10, 27, 30, 27, 23, 20, 25, 27, 90),
    )
    for month in months:
        period = month["period"]
        stock_cost = month["stockCost"]
        matched = month["matchedStockChange"]
        total_change = month["stockChange"]
        new_cost = month["newPositionCost"]
        missing_cost = month["missingPositionCost"]
        summary.append(
            [
                period,
                "Resumo mensal",
                None,
                _excel_scaled(stock_cost, 100),
                None,
                None,
                month["turnover"],
                month["turnoverCount"],
                month["turnoverInvalidCount"],
                month["productCount"],
                _excel_scaled(matched, 100),
                _excel_scaled(total_change, 100),
                None,
                "Sim" if month["partial"] else "Não",
                month["newPositionCount"],
                month["missingPositionCount"],
                month["scope"],
                month["salesScope"],
                month["comparison"],
                month["salesComparison"],
                month["notice"],
                "Sim" if month["stockPartial"] else "Não",
                "Sim" if month["salesObserved"] else "Não",
                _excel_scaled(new_cost, 100),
                _excel_scaled(missing_cost, 100),
            ],
            (
                (4, money_format),
                (7, "0.00"),
                (11, money_format),
                (12, money_format),
                (24, money_format),
                (25, money_format),
            ),
        )
        units = sorted(set(month["stockQuantity"]) | set(month["salesQuantity"]))
        for unit in units:
            quantity = month["stockQuantity"].get(unit)
            sold = month["salesQuantity"].get(unit)
            sales_change = month["salesChange"].get(unit)
            summary.append(
                [
                    period,
                    "Quantidade por unidade",
                    unit,
                    None,
                    _excel_scaled(quantity, SCALE),
                    _excel_scaled(sold, SCALE),
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    _excel_scaled(sales_change, SCALE),
                ],
                ((5, quantity_format), (6, quantity_format), (13, quantity_format)),
            )
        changes.append(
            [
                period,
                month["previousLabel"],
                "Abrangência da comparação",
                None,
                None,
                None,
                None,
                f"{month['newPositionCount']} posições novas; "
                f"{month['missingPositionCount']} ausentes. "
                + month["comparison"]
                + " Registros ausentes não são presumidos como saldo zero.",
            ],
        )
        for item in month["changes"]:
            changes.append(
                [
                    period,
                    month["previousLabel"],
                    "Custo · posições pareadas",
                    None,
                    item["code"],
                    item["name"],
                    _excel_scaled(item["value"], 100),
                    item["detail"],
                ],
                ((7, money_format),),
            )
        for unit, entries in sorted(month["salesChangesByUnit"].items()):
            for item in entries:
                changes.append(
                    [
                        period,
                        month["previousLabel"],
                        "Vendas das filiais",
                        unit,
                        item["code"],
                        item["name"],
                        _excel_scaled(item["value"], SCALE),
                        item["detail"],
                    ],
                    ((7, quantity_format),),
                )
        for item in month["stockRows"]:
            # A recência é da filial/produto/unidade e usa o fim deste mês,
            # limitado à data da base; não atribuir a venda a um armazém.
            last_sale = item["lastSale"]
            products.append(
                [
                    period,
                    "Fechamento",
                    item["code"],
                    item["name"],
                    item["branch"],
                    item["warehouse"],
                    item["unit"],
                    _excel_scaled(item["quantity"], SCALE),
                    None,
                    _excel_scaled(item["cost"], 100),
                    item["turnover"],
                    item["days"],
                    date.fromisoformat(last_sale) if last_sale else None,
                    item["daysSinceSale"],
                    "Dias sem venda observada na filial; não medem idade física nem consumo interno. "
                    + ("Sem saída positiva observada até este mês." if not last_sale else ""),
                ],
                (
                    (8, quantity_format),
                    (10, money_format),
                    (11, "0.00"),
                    (12, "0.00"),
                    (13, "dd/mm/yyyy"),
                ),
            )
        for item in month["saleRows"]:
            products.append(
                [
                    period,
                    "Vendas da filial",
                    item["code"],
                    item["name"],
                    item["branch"],
                    None,
                    item["unit"],
                    None,
                    _excel_scaled(item["quantity"], SCALE),
                    None,
                    None,
                    None,
                    None,
                    None,
                    "Vendas líquidas por filial e unidade, sem alocação por armazém. "
                    "Quantidade não representa consumo interno ou faturamento.",
                ],
                ((9, quantity_format),),
            )
    return layouts


def _append_alert_sheets(workbook, alerts, selection, cancel):
    """Write every alert occurrence, without the presentation's ranking limit."""
    headers = (
        "Alerta",
        "Código",
        "Produto",
        "Filial",
        "UM",
        "Saldo selecionado",
        "Capital positivo (R$)",
        "Última venda observada",
        "Dias sem venda observada",
        "Cobertura filial (dias)",
        "Limite configurado (dias)",
        "Data da base",
        "Período de vendas",
        "Abrangência",
        "Definição",
    )
    widths = (30, 16, 50, 12, 10, 23, 25, 24, 25, 25, 25, 18, 24, 48, 90)
    layouts = {}
    part, row_number, sheet = 0, 0, None

    def new_sheet():
        nonlocal part, row_number, sheet
        part += 1
        title = "Alertas" + (" " + str(part) if part > 1 else "")
        sheet = workbook.create_sheet(title)
        sheet.append(headers)
        layouts[title] = widths
        row_number = 1

    new_sheet()
    categories = (
        (
            "Sem venda acima do limite",
            alerts.stale,
            alerts.policy.stale_days,
            "Saldo selecionado positivo e última venda observada há mais de "
            + str(alerts.policy.stale_days)
            + " dias na data da base. Não mede idade física ou obsolescência.",
        ),
        (
            "Cobertura abaixo do limite",
            alerts.low_coverage,
            alerts.policy.coverage_days,
            "Cobertura conhecida da filial inferior a "
            + str(alerts.policy.coverage_days)
            + " dias, pela demanda líquida do período. Não é previsão garantida de ruptura.",
        ),
        (
            "Sem venda registrada",
            alerts.no_sale,
            None,
            "Saldo selecionado positivo sem saída positiva observada na filial e unidade. "
            "Não comprova ausência de venda por um número de dias nem obsolescência.",
        ),
    )
    for label, entries, threshold, definition in categories:
        for item in entries:
            checkpoint(cancel)
            if row_number >= _ALERTS_MAX_ROWS:
                new_sheet()
            row_number += 1
            scope = "Vendas da filial e unidade; saldo das posições selecionadas"
            if selection.warehouse:
                scope += "; armazém " + selection.warehouse + " selecionado"
            _append_values(
                sheet,
                [
                    label,
                    item.code,
                    item.name,
                    item.branch,
                    item.unit,
                    _excel_scaled(item.quantity, SCALE),
                    _excel_scaled(item.cost, 100),
                    item.last_sale,
                    item.days_since_sale,
                    item.coverage,
                    threshold,
                    alerts.as_of,
                    selection.start + " a " + selection.end,
                    scope,
                    definition,
                ],
                row_number,
            )
            sheet.cell(row_number, 6).number_format = "#,##0.000;[Red](#,##0.000)"
            sheet.cell(row_number, 7).number_format = '"R$" #,##0.00;[Red]("R$" #,##0.00)'
            for column in (8, 12):
                sheet.cell(row_number, column).number_format = "dd/mm/yyyy"
            sheet.cell(row_number, 10).number_format = "0.00"
    checkpoint(cancel)
    return layouts


def export_inventory(
    data: InventoryData,
    analysis: StockAnalysis,
    selection: StockSelection,
    path: Path,
    cancel: Event,
    *,
    monthly=None,
    alerts=None,
    comparison=None,
):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from .inventory_io import _text
    from openpyxl.utils import get_column_letter
    from .inventory_monthly import analyze_monthly_inventory

    if monthly is None:
        monthly = analyze_monthly_inventory(data, selection, analysis=analysis, cancel=cancel)
    if alerts is None:
        alerts = analyze_stock_alerts(data, selection, analysis, cancel=cancel)
    checkpoint(cancel)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Estoque"
    sheet.append(["LIARTE · Estoque · Base " + data.as_of.strftime("%d/%m/%Y")])
    sheet.append(["Saldos atuais; período de vendas " + selection.start + " a " + selection.end])
    sheet.append(HEADERS)
    stock_row_count = 3
    for n, row in enumerate(analysis.rows, start=4):
        checkpoint(cancel)
        b = row.balance
        _append_values(
            sheet,
            [
                b.product,
                row.name,
                row.group,
                b.branch,
                b.warehouse,
                b.unit,
                _excel_scaled(b.quantity, SCALE),
                _excel_scaled(b.cost, 100),
                row.abc,
                row.last_sale,
                row.days_since_sale,
                row.turnover,
                row.source_days,
                row.source_period or None,
                row.coverage,
                row.cv,
            ],
            n,
        )
        stock_row_count = n
        for col in (7,):
            sheet.cell(n, col).number_format = "#,##0.000;[Red](#,##0.000)"
        sheet.cell(n, 8).number_format = '"R$" #,##0.00;[Red]("R$" #,##0.00)'
        sheet.cell(n, 10).number_format = "dd/mm/yyyy"
        for col in (12, 13, 15):
            sheet.cell(n, col).number_format = "0.00"
        sheet.cell(n, 16).number_format = "0.0%"
    sheet.freeze_panes = "C4"
    sheet.auto_filter.ref = f"A3:P{stock_row_count}"
    widths = (16, 50, 26, 12, 10, 9, 19, 22, 9, 15, 19, 21, 16, 15, 23, 23)
    for i, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    history = workbook.create_sheet("Fechamentos")
    history.append(["Mês", "Custo (R$)", "Abrangência"])
    for n, point in enumerate(analysis.trend, start=2):
        checkpoint(cancel)
        _append_values(
            history,
            [
                point["period"],
                _excel_scaled(point["raw"], 100),
                analysis.historical_scope,
            ],
            n,
        )
        history.cell(n, 2).number_format = '"R$" #,##0.00'
    sales = workbook.create_sheet("Vendas por filial")
    sales.append(
        [
            "Mês",
            "UM",
            "Vendas por filial · produtos das posições atuais selecionadas · "
            "sem alocação por armazém",
            "Mês parcial",
        ]
    )
    for n, point in enumerate(analysis.sales_months, start=2):
        checkpoint(cancel)
        _append_values(
            sales,
            [
                point["period"],
                point["unit"],
                _excel_scaled(point["quantity"], SCALE),
                "Sim" if point["partial"] else "Não",
            ],
            n,
        )
        sales.cell(n, 3).number_format = "#,##0.000"
    table_layouts = _append_monthly_sheets(workbook, monthly, cancel)
    table_layouts.update(_append_alert_sheets(workbook, alerts, selection, cancel))
    if comparison is not None:
        from .inventory_compare_export import append_comparison_sheets

        table_layouts.update(
            append_comparison_sheets(workbook, comparison, data, selection, cancel)
        )
    rules = workbook.create_sheet("Leitura dos indicadores")
    rules.append(["Indicador", "Definição"])
    for entry in (
        ("Fonte", data.source_name),
        ("SHA-256 da fonte", data.source_hash),
        ("Data da base", data.as_of.isoformat()),
        ("Qualidade da fonte", " ".join(data.notes) or "Sem ocorrências na validação estrutural."),
        (
            "Giro fonte",
            "GiroEstoqueMensalQtd de Fact_Estoque, último fechamento da base; não é recalculado nem somado entre posições.",
        ),
        ("Dias fonte", "DiasEstoqueQtd da mesma posição e fechamento. Campo original do BI."),
        (
            "Cobertura filial",
            "Saldo de todas as posições da filial e produto / média diária de vendas na mesma unidade e período. Indisponível com filtro de armazém, saldo negativo ou sem vendas positivas. Inclui dias sem venda; não mede consumo interno.",
        ),
        (
            "CV",
            "Desvio padrão amostral / média das quantidades vendidas por mês completo. Mínimo de três meses; sem cálculo para média nula/negativa ou meses com quantidade líquida negativa.",
        ),
        (
            "Vendas por filial · posições atuais",
            "A aba Vendas por filial considera somente produtos, filiais e unidades das posições "
            "ATUAIS selecionadas na aba Estoque. Produtos sem posição atual não entram nessa série. "
            "Com filtro de armazém, as posições atuais definem a coorte, mas suas vendas continuam "
            "sendo da filial inteira, pois a fonte não identifica o armazém da venda. "
            "A aba Análise mensal possui universo histórico próprio, descrito em Vendas mensais.",
        ),
        (
            "Série de vendas",
            "Meses observados sem venda recebem zero. Meses anteriores à primeira venda da base ficam em branco. Meses de início/fim incompletos são identificados como parciais e não entram na média mensal, desvio padrão e CV.",
        ),
        (
            "Última venda",
            "Última data de saída com quantidade positiva até a data da base, por produto, filial e unidade; não é última movimentação de estoque.",
        ),
        (
            "ABC",
            "Participação no capital positivo do recorte antes de filtrar ABC/atividade/saldo: A até atingir 80%; B até 95%; C restante. O produto que cruza o limite permanece na classe anterior.",
        ),
        (
            "Fechamentos",
            "Meses ausentes permanecem em branco. Abrangência do histórico pode diferir do estoque atual; valores de meses diferentes não são somados como saldo.",
        ),
        (
            "Unidades",
            "Quantidades de unidades diferentes não são somadas. Valores a custo em reais; saldos negativos são preservados.",
        ),
        (
            "Análise mensal",
            "Estoque e giro pertencem ao fechamento de cada mês, sem reconstrução do saldo atual. "
            "O resumo monetário aparece uma vez por mês; quantidades ficam em linhas separadas por unidade. "
            "Produtos mensais inclui todas as posições históricas e todas as vendas da seleção, "
            "inclusive produtos sem saldo atual; não é limitado ao ranking exibido na tela.",
        ),
        (
            "Giro mediano mensal",
            "Mediana dos indicadores não negativos da fonte em posições com saldo positivo no mês. "
            "O número de posições elegíveis e os giros negativos são informados separadamente. "
            "Não é o giro consolidado da empresa nem é somado entre posições.",
        ),
        (
            "Variações mensais",
            "Compara meses consecutivos do período selecionado, sem pular lacunas. "
            "Contribuições de custo usam somente posições presentes nos dois fechamentos. "
            "Posições novas/ausentes são alterações de registros, não entradas/saídas físicas presumidas. "
            "A variação total corresponde à variação pareada mais o custo dos registros novos "
            "menos o custo dos registros ausentes. "
            "Contribuições indicam onde o valor variou; não comprovam compras, consumo, "
            "transferências, ajustes ou suas causas.",
        ),
        (
            "Vendas mensais",
            "Nas abas Análise mensal e Produtos mensais, a quantidade líquida vem dos registros "
            "históricos de produto, filial e unidade, inclusive produtos sem posição atual. "
            "ABC, atividade e saldo restringem esse universo à coorte atual selecionada, quando aplicados. "
            "Não há faturamento nesta análise. Sem atribuição a armazéns: o filtro de armazém "
            "indisponibiliza as vendas nessas duas abas, enquanto Vendas por filial mantém a série "
            "dos produtos das posições atuais selecionadas, com vendas da filial inteira. "
            "Meses sem cobertura ficam em branco; meses parciais são identificados e não geram comparação mensal.",
        ),
        (
            "Recência mensal",
            "Última saída positiva observada até o fim do mês, limitada à data da base, "
            "por produto/filial/unidade. Dias sem venda não medem idade física ou obsolescência. "
            "Consumo interno e histórico de entradas/lotes não estão disponíveis nesta base.",
        ),
        (
            "Alertas · limites aplicados",
            f"Tempo sem venda: mais de {alerts.policy.stale_days} dias. "
            f"Cobertura baixa: inferior a {alerts.policy.coverage_days} dias. "
            f"Referência da recência: {alerts.as_of.isoformat()}. "
            f"Período de demanda: {selection.start} a {selection.end}.",
        ),
        (
            "Alertas · tempo sem venda",
            "Agrupa o saldo e o custo das posições atuais selecionadas por produto, filial e "
            "unidade. Saldo agrupado positivo e dias desde a última saída positiva observada "
            "estritamente acima do limite. Não mede idade de lote, consumo ou obsolescência. "
            "A data de referência é a da base, independentemente do período de demanda.",
        ),
        (
            "Alertas · sem venda registrada",
            "Categoria informativa separada: saldo selecionado positivo sem última saída "
            "positiva observada. A ausência de histórico não é convertida em zero ou infinito "
            "e não comprova que o limite de tempo sem venda foi ultrapassado.",
        ),
        (
            "Alertas · cobertura baixa",
            "Usa a cobertura conhecida de todas as posições da filial, produto e unidade, "
            "pela demanda líquida no período selecionado. Cobertura zero pode ser sinalizada; "
            "valores ausentes ou negativos não são tratados como zero. A regra fica "
            "indisponível com filtro de armazém, pois vendas não têm essa atribuição. "
            "Não é previsão garantida de ruptura nem medida de consumo interno.",
        ),
        (
            "Alertas · disponibilidade da cobertura",
            ("Disponível no recorte. " if alerts.coverage_available else "Indisponível com filtro de armazém. ")
            + str(alerts.coverage_unavailable_count)
            + " grupos sem cobertura calculável.",
        ),
        (
            "Alertas · ocorrências e capital",
            f"Tempo sem venda: {len(alerts.stale)}; cobertura baixa: {len(alerts.low_coverage)}; "
            f"sem venda registrada: {len(alerts.no_sale)}. Cada linha representa um produto, "
            "filial, unidade e regra. Um mesmo grupo pode aparecer em regras diferentes; "
            "somar ocorrências ou capital entre regras duplica esses grupos. O capital é "
            "o custo positivo após agregar as posições selecionadas; quantidades de unidades "
            "diferentes não são somadas. A aba inclui todas as ocorrências, mesmo além da "
            "lista resumida da interface.",
        ),
        (
            "Precisão numérica do Excel",
            "Valores com mais de 15 dígitos significativos são texto decimal exato, "
            "identificado por comentário na célula. As unidades continuam as do cabeçalho; "
            "essas células textuais não entram nas somas automáticas do Excel. "
            "Demais quantidades e valores monetários permanecem numéricos.",
        ),
        ("Recorte", "; ".join(f"{k}={v}" for k, v in asdict(selection).items() if v)),
    ):
        rules.append(entry)
    if comparison is not None:
        for entry in (
            (
                "Comparação mensal · par escolhido",
                f"Base: {comparison['base']}; comparado: {comparison['target']}. "
                "Meses arbitrários do recorte, inclusive ordem invertida. "
                "Um intervalo entre os meses não é interpolado.",
            ),
            (
                "Comparação mensal · disponibilidade",
                "Estoque exige dois fechamentos conhecidos e completos; vendas exigem dois "
                "meses observados e completos, independentemente dos fechamentos. "
                "Meses parciais ou sem observação mantêm a diferença indisponível. "
                "Percentuais são calculados somente com valor da base positivo.",
            ),
            (
                "Comparação mensal · reconciliação",
                "Diferença total observada = diferença nas posições pareadas + valor dos "
                "registros só no mês comparado − valor dos registros só na base. "
                "O pareamento de estoque conserva produto, filial, armazém e unidade. "
                "Ausência de posição não representa saldo zero; mudanças de registros "
                "não comprovam compras, consumo, transferências ou suas causas.",
            ),
            (
                "Comparação mensal · contribuições",
                "Inclui todos os grupos, além do limite visual da interface. Produto e "
                "Filial são perspectivas independentes da mesma diferença e não devem ser "
                "somadas. Estoque a custo reúne todas as unidades selecionadas; quantidades "
                "e vendas permanecem separadas por unidade. Valores desconhecidos ficam vazios.",
            ),
        ):
            rules.append(entry)
    edge = Side(style="thin", color="E3E9EF")
    alignment = Alignment(vertical="center", wrap_text=True)
    normal_font = Font(name="Calibri", size=11, color="20334A")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    border = Border(bottom=edge)
    header_fill = PatternFill("solid", fgColor="103F70")
    for ws in workbook:
        header_row = 3 if ws is sheet else 1
        if ws is not sheet:
            ws.column_dimensions["A"].width = 27
            ws.column_dimensions["B"].width = 88 if ws is rules else 24
            ws.column_dimensions["C"].width = 60
            ws.freeze_panes = "A2"
        if ws.title in table_layouts:
            for column, width in enumerate(table_layouts[ws.title], 1):
                ws.column_dimensions[get_column_letter(column)].width = width
            ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
        for row in ws:
            checkpoint(cancel)
            for cell in row:
                if isinstance(cell.value, str):
                    _text(cell.value, f"{ws.title}!{cell.coordinate}", empty=True, max_length=32767)
                    if cell.comment is not None and cell.comment.author == "LIARTE":
                        cell.number_format = "@"
                    cell.data_type = "s"  # Não interpretar descrições como fórmulas.
                cell.alignment = alignment
                cell.font = normal_font
                cell.border = border
                if cell.row == header_row:
                    cell.fill = header_fill
                    cell.font = header_font
            ws.row_dimensions[row[0].row].height = (
                34
                if ws is sheet
                else max(30, math.ceil(len(str(row[1].value or "")) / 78) * 15 + 12)
                if ws is rules
                else 25
            )
        ws.print_title_rows = f"{header_row}:{header_row}"
        ws.sheet_view.showGridLines = False
        if ws is sales:
            ws.row_dimensions[header_row].height = 60
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(prefix=".estoque-", suffix=".xlsx", dir=path.parent)
    os.close(fd)
    try:
        workbook.save(name)
        checkpoint(cancel)
        _validate_staged_workbook(Path(name), cancel)
        checkpoint(cancel)
        os.replace(name, path)
    finally:
        workbook.close()
        if os.path.exists(name):
            os.unlink(name)
