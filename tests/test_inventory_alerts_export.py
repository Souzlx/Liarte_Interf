"""Alert workbooks retain complete occurrences, provenance and exact values."""

from dataclasses import replace
from datetime import date
from decimal import localcontext
from pathlib import Path
from tempfile import TemporaryDirectory
from threading import Event
import unittest
from unittest.mock import PropertyMock, patch

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from liarte_core.inventory import (
    SCALE,
    Balance,
    InventoryData,
    Product,
    Sale,
    StockSelection,
    analyze_inventory,
)
from liarte_core.inventory_alerts import AlertPolicy, analyze_stock_alerts
from liarte_core.inventory_export import _append_alert_sheets, export_inventory
from liarte_core.query import Cancelled


def alert_export_data():
    codes = tuple(f"P{index:03d}" for index in range(16))
    return InventoryData(
        date(2026, 4, 15),
        "Synthetic alert export",
        "a" * 64,
        tuple(Product(code, "Produto " + code, "Grupo", "PA", "KG") for code in codes)
        + (Product("NONE", "Sem histórico", "Grupo", "PA", "UN"),),
        tuple(Balance(code, "01", "01", "KG", SCALE, 100) for code in codes)
        + (
            Balance("P000", "01", "02", "KG", SCALE, 100),
            Balance("P000", "01", "01", "L", 3 * SCALE, 300),
            Balance("NONE", "01", "01", "UN", 2 * SCALE, 200),
        ),
        (),
        tuple(Sale(code, "01", "KG", date(2026, 1, 1), 100 * SCALE, 100) for code in codes)
        + (Sale("P000", "01", "L", date(2026, 1, 1), 100 * SCALE, 100),),
    )


class InventoryAlertsExportTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "alerts.xlsx"
        self.data = alert_export_data()
        self.selected = StockSelection("2026-01", "2026-04")

    def export(self, *, data=None, selected=None, alerts=None):
        data = self.data if data is None else data
        selected = self.selected if selected is None else selected
        analysis = analyze_inventory(data, selected)
        export_inventory(
            data, analysis, selected, self.path, Event(), monthly=(), alerts=alerts
        )
        book = load_workbook(self.path)
        self.addCleanup(book.close)
        return book

    @staticmethod
    def rows(sheet):
        values = list(sheet.iter_rows(values_only=True))
        return [dict(zip(values[0], row)) for row in values[1:]]

    def test_default_export_keeps_every_occurrence_and_deduplicates_warehouses(self):
        book = self.export()
        rows = self.rows(book["Alertas"])
        stale = [row for row in rows if row["Alerta"] == "Sem venda acima do limite"]
        low = [row for row in rows if row["Alerta"] == "Cobertura abaixo do limite"]
        self.assertEqual(len(stale), 17)
        self.assertEqual(len(low), 17)
        self.assertGreater(len(stale), 12)
        group = [row for row in stale if row["Código"] == "P000"]
        self.assertEqual({row["UM"] for row in group}, {"KG", "L"})
        self.assertEqual(len(group), 2)
        quantities = {row["UM"]: row["Saldo selecionado"] for row in group}
        self.assertEqual(quantities, {"KG": 2, "L": 3})
        self.assertEqual(
            {(row["Filial"], row["UM"]) for row in rows if row["Código"] == "P000"},
            {("01", "KG"), ("01", "L")},
        )
        self.assertTrue(all(row["Data da base"].date() == self.data.as_of for row in rows))
        self.assertTrue(all(row["Período de vendas"] == "2026-01 a 2026-04" for row in rows))
        self.assertEqual(book["Alertas"].freeze_panes, "A2")
        self.assertEqual(book["Alertas"].auto_filter.ref, f"A1:O{len(rows) + 1}")

    def test_custom_applied_policy_is_exported_without_recalculation(self):
        analysis = analyze_inventory(self.data, self.selected)
        alerts = analyze_stock_alerts(
            self.data, self.selected, analysis, AlertPolicy(stale_days=100, coverage_days=2)
        )
        with patch(
            "liarte_core.inventory_export.analyze_stock_alerts",
            side_effect=AssertionError("The controller's applied result must be reused"),
        ):
            book = self.export(alerts=alerts)
        rows = self.rows(book["Alertas"])
        stale = [row for row in rows if row["Alerta"] == "Sem venda acima do limite"]
        low = [row for row in rows if row["Alerta"] == "Cobertura abaixo do limite"]
        self.assertEqual(len(stale), 17)
        self.assertEqual(len(low), 15)
        self.assertTrue(all(row["Limite configurado (dias)"] == 100 for row in stale))
        self.assertTrue(all(row["Limite configurado (dias)"] == 2 for row in low))
        rules = dict(book["Leitura dos indicadores"].iter_rows(min_row=2, values_only=True))
        self.assertIn("mais de 100 dias", rules["Alertas · limites aplicados"])
        self.assertIn("inferior a 2 dias", rules["Alertas · limites aplicados"])
        self.assertIn("duplica", rules["Alertas · ocorrências e capital"])

    def test_missing_sale_is_separate_information_with_blank_time_and_coverage(self):
        book = self.export()
        unknown = [row for row in self.rows(book["Alertas"]) if row["Código"] == "NONE"]
        self.assertEqual(len(unknown), 1)
        self.assertEqual(unknown[0]["Alerta"], "Sem venda registrada")
        for field in (
            "Última venda observada", "Dias sem venda observada",
            "Cobertura filial (dias)", "Limite configurado (dias)",
        ):
            self.assertIsNone(unknown[0][field])
        self.assertIn("Não comprova", unknown[0]["Definição"])

    def test_known_zero_coverage_remains_numeric_zero_instead_of_missing_data(self):
        data = InventoryData(
            self.data.as_of, "Synthetic zero coverage", "c" * 64,
            (Product("ZERO", "Sem saldo", "Grupo", "PA", "UN"),),
            (Balance("ZERO", "01", "01", "UN", 0, 0),), (),
            (Sale("ZERO", "01", "UN", date(2026, 4, 1), SCALE, 100),),
        )
        book = self.export(data=data)
        rows = self.rows(book["Alertas"])
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["Alerta"], "Cobertura abaixo do limite")
        self.assertEqual(rows[0]["Cobertura filial (dias)"], 0)
        self.assertEqual(rows[0]["Saldo selecionado"], 0)
        self.assertEqual(book["Alertas"]["J2"].data_type, "n")

    def test_warehouse_selection_exports_no_low_coverage_and_explains_why(self):
        book = self.export(selected=replace(self.selected, warehouse="01"))
        rows = self.rows(book["Alertas"])
        self.assertFalse(any(row["Alerta"] == "Cobertura abaixo do limite" for row in rows))
        self.assertTrue(any(row["Alerta"] == "Sem venda acima do limite" for row in rows))
        rules = dict(book["Leitura dos indicadores"].iter_rows(min_row=2, values_only=True))
        self.assertIn("Indisponível com filtro de armazém", rules["Alertas · disponibilidade da cobertura"])
        self.assertTrue(all("armazém 01" in row["Abrangência"] for row in rows))

    def test_extreme_quantities_and_cost_remain_exact_text_and_names_are_not_formulas(self):
        value = 10**16 + 1
        data = InventoryData(
            self.data.as_of, "Synthetic precision", "b" * 64,
            (Product("=1+1", "=HYPERLINK(1)", "Grupo", "PA", "UN"),),
            (Balance("=1+1", "01", "01", "UN", value, value),), (), (),
        )
        with localcontext() as context:
            context.prec = 6
            book = self.export(data=data)
        sheet = book["Alertas"]
        self.assertEqual(sheet["B2"].value, "=1+1")
        self.assertEqual(sheet["C2"].value, "=HYPERLINK(1)")
        self.assertEqual(sheet["B2"].data_type, "s")
        self.assertEqual(sheet["C2"].data_type, "s")
        for address, expected in (
            ("F2", "10000000000.000001"), ("G2", "100000000000000.01"),
        ):
            cell = sheet[address]
            self.assertEqual(cell.value, expected)
            self.assertEqual(cell.number_format, "@")
            self.assertEqual(cell.data_type, "s")
            self.assertIn("15 dígitos", cell.comment.text)

    def test_sheet_splitting_keeps_all_occurrences_and_does_not_scan_max_row(self):
        analysis = analyze_inventory(self.data, self.selected)
        alerts = analyze_stock_alerts(self.data, self.selected, analysis)
        book = Workbook()
        self.addCleanup(book.close)
        with (
            patch("liarte_core.inventory_export._ALERTS_MAX_ROWS", 5),
            patch.object(
                Worksheet, "max_row", new_callable=PropertyMock,
                side_effect=AssertionError("quadratic row scan"),
            ),
        ):
            layouts = _append_alert_sheets(book, alerts, self.selected, Event())
        self.assertEqual(len(layouts), 9)
        self.assertEqual(
            sum(book[name]._current_row - 1 for name in layouts),
            len(alerts.stale) + len(alerts.low_coverage) + len(alerts.no_sale),
        )
        self.assertTrue(all(book[name]._current_row <= 5 for name in layouts))
        self.assertTrue(all(book[name]["A1"].value == "Alerta" for name in layouts))

    def test_cancel_during_alert_append_preserves_previous_workbook(self):
        analysis = analyze_inventory(self.data, self.selected)
        alerts = analyze_stock_alerts(self.data, self.selected, analysis)
        self.path.write_bytes(b"previous export")
        cancel = Event()

        def interrupted(*args, **kwargs):
            cancel.set()
            return _append_alert_sheets(*args, **kwargs)

        with (
            patch("liarte_core.inventory_export._append_alert_sheets", side_effect=interrupted),
            self.assertRaises(Cancelled),
        ):
            export_inventory(
                self.data, analysis, self.selected, self.path, cancel, monthly=(), alerts=alerts
            )
        self.assertEqual(self.path.read_bytes(), b"previous export")
        self.assertEqual(list(Path(self.temp.name).iterdir()), [self.path])


if __name__ == "__main__":
    unittest.main()
