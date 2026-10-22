"""Regras de integridade e apresentação do BI, sem depender do Desktop ou da rede."""

from __future__ import annotations

from decimal import Decimal
import hashlib
import json
from pathlib import Path
import tempfile
import unittest
import zipfile
from openpyxl import Workbook, load_workbook
from liarte_engine import powerbi_importacoes as bi
from liarte_engine import powerbi_semantic as semantic
from liarte_engine import powerbi_visuals as visuals
from liarte_engine.powerbi_source import audit_source, code, number, period


class PowerBiTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.excel = self.root / "IMPORTACAO_2026-01_2026-03.xlsx"
        self.output = self.root / "BI.zip"
        book = Workbook()
        general = book.active
        general.title = "Base geral"
        general.append(bi.EXPECTED_IMPORT_HEADERS)
        for month, ncm, fob, freight, kg in [
            ("01", "29161220", 1000, 100, 100),
            ("02", "29161220", 3000, 300, 1000),
            ("03", "29161220", 2000, 200, 400),
            ("03", "29161230", 4000, 400, 200),
        ]:
            general.append(
                [
                    month,
                    ncm,
                    "Descrição oficial",
                    "SP",
                    "PORTO DE SANTOS",
                    "China",
                    fob,
                    freight,
                    0,
                    fob + freight,
                    kg,
                    "G10" if ncm == "29161220" else "G12",
                    (fob + freight) / kg,
                ]
            )
        municipal = book.create_sheet("Base municipal")
        municipal.append(bi.EXPECTED_MUNICIPAL_HEADERS)
        for state, fob in [("SP", 1000), ("RS", 2000)]:
            municipal.append(
                [
                    "03",
                    "Bom Jesus",
                    state,
                    "2916",
                    "Ácidos carboxílicos",
                    "China",
                    fob,
                    100,
                    "G10",
                    "29161220",
                    fob / 100,
                ]
            )
        products = book.create_sheet("Produtos")
        products.append(bi.EXPECTED_PRODUCT_HEADERS)
        for p, n in [("G10", "29161220"), ("G10-alias", "29161220"), ("G12", "29161230")]:
            products.append([p, n])
        book.save(self.excel)
        book.close()

    def tearDown(self):
        self.temp.cleanup()

    def change(self, sheet, cell, value):
        book = load_workbook(self.excel)
        book[sheet][cell] = value
        book.save(self.excel)
        book.close()

    def generate(self, **kwargs):
        return bi.create_powerbi_package(self.excel, self.output, "2026-01", "2026-03", **kwargs)

    def unpack(self):
        self.generate()
        project = self.root / "project"
        with zipfile.ZipFile(self.output) as z:
            z.extractall(project)
        return project

    def audit(self):
        return audit_source(self.excel, bi._validate_excel_source(self.excel), "2026-01", "2026-03")

    def test_01_exact_source_controls_and_months(self):
        c = self.audit()
        self.assertEqual(
            c["general"]["totals"],
            {
                "rows": 4,
                "zero_weight_rows": 0,
                "kg": 1700,
                "fob_cents": 1000000,
                "freight_cents": 100000,
                "insurance_cents": 0,
                "cif_cents": 1100000,
            },
        )
        self.assertEqual(c["municipal"]["totals"]["fob_cents"], 300000)
        self.assertEqual(c["catalog"], {"relations": 3, "ncms": 2, "products": 3})
        self.assertEqual(c["general"]["monthly"]["2026-02"]["cif_cents"], 330000)
        self.assertEqual(c["sha256"], hashlib.sha256(self.excel.read_bytes()).hexdigest())

    def test_02_embedded_excel_is_byte_identical(self):
        self.generate()
        with zipfile.ZipFile(self.output) as z:
            self.assertEqual(z.read("Dados/" + self.excel.name), self.excel.read_bytes())
            c = json.loads(z.read("CONTROLE_FONTE.json"))
            self.assertEqual(c["sha256"], hashlib.sha256(self.excel.read_bytes()).hexdigest())
            self.assertIn("VERIFICAR_MEDIDAS.dax", z.namelist())

    def test_03_cif_discrepancy_is_rejected(self):
        self.change("Base geral", "J2", 1099)
        with self.assertRaisesRegex(bi.PowerBIProjectError, "CIF diverge"):
            self.generate()

    def test_04_monetary_values_are_not_guessed_or_rounded(self):
        for value in [
            "1100",
            "1.100,00",
            Decimal("NaN"),
            Decimal("Infinity"),
            Decimal("-1"),
            Decimal("1.001"),
        ]:
            with self.subTest(value=value), self.assertRaises(ValueError):
                number(value, True)
        self.assertEqual(number(Decimal("1.01"), True), 101)
        self.change("Base geral", "G2", "1000")
        with self.assertRaisesRegex(bi.PowerBIProjectError, "textual"):
            self.generate()

    def test_05_fractional_kg_is_rejected(self):
        self.change("Base geral", "K2", 100.5)
        with self.assertRaisesRegex(bi.PowerBIProjectError, "fracionário"):
            self.generate()

    def test_06_formulas_in_data_are_rejected(self):
        self.change("Base geral", "J2", "=G2+H2+I2")
        with self.assertRaisesRegex(bi.PowerBIProjectError, "Fórmula"):
            self.generate()

    def test_07_duplicate_fact_grain_is_rejected(self):
        book = load_workbook(self.excel)
        sheet = book["Base geral"]
        sheet.append([c.value for c in sheet[2]])
        book.save(self.excel)
        book.close()
        with self.assertRaisesRegex(bi.PowerBIProjectError, "duplicado"):
            self.generate()

    def test_08_missing_dimensions_and_unknown_codes_are_rejected(self):
        self.change("Base geral", "F2", None)
        with self.assertRaisesRegex(bi.PowerBIProjectError, "Dimensão"):
            self.generate()
        self.change("Base geral", "F2", "China")
        self.change("Base geral", "B2", "99999999")
        with self.assertRaisesRegex(bi.PowerBIProjectError, "fora do cadastro"):
            self.generate()
        with self.assertRaises(ValueError):
            code("29X61220", 8)

    def test_09_months_cross_year_require_explicit_year(self):
        self.assertEqual(period("2026-01. Janeiro", "2025-12", "2026-03"), "2026-01")
        for value, start, end in [
            ("01", "2025-12", "2026-03"),
            ("13", "2026-01", "2026-12"),
            ("2026-04", "2026-01", "2026-03"),
        ]:
            with self.subTest(value=value), self.assertRaises(ValueError):
                period(value, start, end)

    def test_10_empty_dataset_has_valid_zero_controls(self):
        book = load_workbook(self.excel)
        for name in ["Base geral", "Base municipal"]:
            book[name].delete_rows(2, book[name].max_row)
        book.save(self.excel)
        book.close()
        c = self.audit()
        self.assertEqual(c["general"]["totals"]["rows"], 0)
        self.assertEqual(c["general"]["totals"]["kg"], 0)
        self.assertEqual(self.generate()["pages"], 6)

    def test_11_source_failure_preserves_existing_destination(self):
        self.output.write_bytes(b"previous")
        self.change("Base geral", "G2", -1)
        with self.assertRaises(bi.PowerBIProjectError):
            self.generate()
        self.assertEqual(self.output.read_bytes(), b"previous")

    def test_12_cancel_preserves_destination_and_cleans_temporary(self):
        self.output.write_bytes(b"previous")

        class Cancelled(Exception):
            pass

        def progress(value, _):
            if value >= 65:
                raise Cancelled()

        with self.assertRaises(Cancelled):
            self.generate(progress_callback=progress)
        self.assertEqual(self.output.read_bytes(), b"previous")
        self.assertEqual(list(self.root.glob("*.tmp")), [])

    def test_13_source_hash_tampering_is_detected(self):
        project = self.unpack()
        control = project / "CONTROLE_FONTE.json"
        c = json.loads(control.read_text())
        c["sha256"] = "0" * 64
        control.write_text(json.dumps(c))
        with self.assertRaisesRegex(bi.PowerBIProjectError, "hash"):
            bi.validate_powerbi_project(project)

    def visual(self, project, kind):
        for p in project.glob("*.Report/definition/pages/*/visuals/*/visual.json"):
            doc = json.loads(p.read_text())
            if doc["visual"]["visualType"] == kind:
                return p, doc
        self.fail("Visual ausente: " + kind)

    def test_14_unknown_field_is_detected(self):
        project = self.unpack()
        p, doc = self.visual(project, "lineChart")
        projection = doc["visual"]["query"]["queryState"]["Y"]["projections"][0]
        projection["queryRef"] = "Importacoes.Inexistente"
        projection["field"]["Measure"]["Property"] = "Inexistente"
        p.write_text(json.dumps(doc))
        with self.assertRaisesRegex(bi.PowerBIProjectError, "inexistente"):
            bi.validate_powerbi_project(project)

    def test_15_implicit_numeric_projection_is_detected(self):
        project = self.unpack()
        p, doc = self.visual(project, "lineChart")
        projection = doc["visual"]["query"]["queryState"]["Y"]["projections"][0]
        projection["queryRef"] = "Importacoes.Quilograma Líquido"
        projection["field"] = {
            "Column": {
                "Expression": {"SourceRef": {"Entity": "Importacoes"}},
                "Property": "Quilograma Líquido",
            }
        }
        p.write_text(json.dumps(doc))
        with self.assertRaisesRegex(bi.PowerBIProjectError, "implícita"):
            bi.validate_powerbi_project(project)

    def test_16_invalid_geometry_is_detected(self):
        project = self.unpack()
        p, doc = self.visual(project, "lineChart")
        original = doc["position"].copy()
        for change in [
            {"x": float("nan")},
            {"width": -1},
            {"x": 1440},
            {"x": 24, "y": 22, "width": 100, "height": 43},
        ]:
            doc["position"] = original | change
            p.write_text(json.dumps(doc))
            with self.subTest(change=change), self.assertRaises(bi.PowerBIProjectError):
                bi.validate_powerbi_project(project)

    def test_17_relationship_cardinality_tampering_is_detected(self):
        project = self.unpack()
        p = project / (bi.MODEL_DIR_NAME + "/definition/relationships.tmdl")
        p.write_text(p.read_text().replace("toCardinality: one", "toCardinality: many", 1))
        with self.assertRaisesRegex(bi.PowerBIProjectError, "cardinalidades"):
            bi.validate_powerbi_project(project)

    def test_18_municipal_slicers_cannot_share_general_group(self):
        project = self.unpack()
        for p in project.glob("*.Report/definition/pages/*/visuals/*/visual.json"):
            d = json.loads(p.read_text())
            v = d["visual"]
            g = v.get("syncGroup", {}).get("groupName", "")
            if g.startswith("municipal_"):
                v["syncGroup"]["groupName"] = "geral_" + g.removeprefix("municipal_")
                p.write_text(json.dumps(d))
                break
        with self.assertRaisesRegex(bi.PowerBIProjectError, "grupos separados"):
            bi.validate_powerbi_project(project)

    def test_19_money_and_grain_are_explicit(self):
        model = {t.name: t for t in semantic.tables()}
        self.assertEqual(len(model), 11)
        self.assertTrue(
            all(
                c.dtype == "decimal" and c.hidden
                for c in model["Importacoes"].columns
                if c.name in ["US$ FOB", "US$ Frete", "US$ Seguro", "Valor US$ CIF"]
            )
        )
        price = next(c for c in model["Importacoes"].columns if c.name == "US$ CIF/kg")
        self.assertTrue(price.hidden)
        self.assertEqual(price.dtype, "double")
        self.assertIn("summarizeBy: none", semantic.table_tmdl(model["Importacoes"]))
        self.assertFalse(
            {"Código NCM", "URF", "Valor US$ CIF"} & {c.name for c in model["Municipios"].columns}
        )
        self.assertEqual(sum(both for *_, both in semantic.relationships()), 1)
        self.assertIn(
            "Chave município",
            next(
                m.expression for m in model["Municipios"].measures if m.name == "Municípios ativos"
            ),
        )

    def test_20_layout_and_ids_are_deterministic_and_compact(self):
        pages = visuals.pages()
        self.assertEqual(len(pages), 6)
        self.assertEqual(sum(len(s) for _, s in pages), 85)
        for title, specs in pages:
            for i, spec in enumerate(specs):
                a = visuals.visual_payload(spec, title, i)
                self.assertEqual(a, visuals.visual_payload(spec, title, i))
                if spec.visual_type == "cardVisual":
                    self.assertLess(spec.height, visuals.CANVAS_HEIGHT * 0.15)

    def test_21_weighted_measures_and_monthly_context_are_preserved(self):
        ms = {m.name: m.expression for t in semantic.tables() for m in t.measures}
        self.assertEqual(len(ms), 57)
        self.assertNotIn("AVERAGE(", "".join(ms.values()).upper())
        self.assertNotIn("DATEADD(", "".join(ms.values()).upper())
        self.assertIn("DIVIDE([Valor CIF], [Peso líquido (kg)])", ms["CIF por kg"])
        self.assertIn("COUNTROWS(Janela) = 3", ms["CIF por kg MM3"])
        self.assertIn("MesesComPeso = 3", ms["CIF por kg MM3"])
        self.assertIn("'Calendario'[Índice mês] = MesAtual", ms["CIF por kg vs mês anterior"])
        self.assertIn("ALLSELECTED", ms["Concentração Top 5 Países"])
        self.assertNotIn("Municipios", ms["Valor CIF"])

    def test_22_m_path_and_tmdl_mask_escaping(self):
        self.assertEqual(semantic.mtext('C:\\a#(tab)\\b".xlsx'), '"C:\\a#(#)(tab)\\b"".xlsx"')
        self.assertEqual(semantic.property_text(semantic.MONEY), '"""US$"" #,##0.00"')
        self.assertIn("Replace('#(', '#(#)(')", bi._powershell_launcher())
        self.assertIn("MatchEvaluator", bi._powershell_launcher())
        project = self.unpack()
        definition = json.loads((project / bi.REPORT_DIR_NAME / "definition.pbir").read_text())
        self.assertNotIn("byConnection", definition["datasetReference"])

    def test_23_detail_tables_are_flat_single_fact_column_queries(self):
        for title, specs in visuals.pages():
            for spec in specs:
                if spec.visual_type != "tableEx":
                    continue
                source = "Municipios" if title == "Análise Municipal" else "Importacoes"
                roles = dict(spec.roles)
                self.assertEqual(set(roles), {"Values"})
                fields = roles["Values"]
                self.assertTrue(all(f.table == source and f.kind == "column" for f in fields))
                names = {f.name for f in fields}
                self.assertTrue({"Data", "Países", "Quilograma Líquido"} <= names)
                required = (
                    {"Município / UF", "Código SH4"}
                    if source == "Municipios"
                    else {"Código NCM", "UF do Produto", "URF"}
                )
                self.assertTrue(required <= names)
                payload = visuals.visual_payload(spec, title, 1)["visual"]
                self.assertEqual(
                    payload["objects"]["total"][0]["properties"]["totals"]["expr"]["Literal"][
                        "Value"
                    ],
                    "false",
                )
                self.assertEqual(
                    payload["visualContainerObjects"]["stylePreset"][0]["properties"]["name"][
                        "expr"
                    ]["Literal"]["Value"],
                    "'None'",
                )

    def test_24_details_recompute_price_from_values_not_cached_ratio(self):
        shared = {name: source for name, source, _ in semantic.expressions("2026-01", "2026-03")}
        for key, price, numerator in [
            ("BaseGeral", "US$ CIF/kg", "Valor US$ CIF"),
            ("BaseMunicipal", "US$ FOB/kg", "US$ FOB"),
        ]:
            query = shared[key]
            self.assertIn("Table.RemoveColumns", query)
            self.assertIn('"' + price + '"', query)
            self.assertIn("if [Quilograma Líquido] = 0 then null", query)
            self.assertIn('[#"' + numerator + '"] / [Quilograma Líquido]', query)
            self.assertNotIn("Number.Round(", query)
        # The exporter retains the Excel itself byte-for-byte even with a stale
        # display ratio. M recomputes this column in the model, not in the source.
        self.change("Base geral", "M2", 99999)
        self.generate()
        with zipfile.ZipFile(self.output) as z:
            self.assertEqual(z.read("Dados/" + self.excel.name), self.excel.read_bytes())

    def test_25_price_guidance_prioritizes_single_ncm_selection(self):
        expression = next(
            m.expression for m in semantic.general_measures() if m.name == "Contexto de preços"
        )
        self.assertLess(expression.index("NOT HASONEVALUE"), expression.index("[CIF sem peso] > 0"))
        self.assertIn("único NCM", expression)
        self.assertIn("[CIF sem peso] > 0", expression)

    def test_26_uf_slicer_has_room_for_default_label(self):
        for title, specs in visuals.pages():
            for spec in specs:
                if spec.visual_type == "slicer" and spec.title == "UF":
                    self.assertGreaterEqual(spec.width, 104, title)
            slicers = [s for s in specs if s.visual_type == "slicer"]
            self.assertEqual(slicers[-1].x + slicers[-1].width, 1416)

    def test_27_city_label_does_not_append_same_uf_twice(self):
        shared = {n: q for n, q, _ in semantic.expressions("2026-01", "2026-03")}
        query = shared["BaseMunicipal"]
        self.assertIn(
            'Text.EndsWith([Município fiscal], " - " & [UF do Município], Comparer.OrdinalIgnoreCase)',
            query,
        )
        dimension = next(t for t in semantic.tables() if t.name == "MunicipiosFiscais")
        self.assertIn('"Município / UF"', dimension.source)
        self.assertNotIn("Table.AddColumn", dimension.source)
        self.assertIn('"Chave município"', dimension.source)

    def test_28_logistics_overview_uses_full_freight_not_incomplete_ratio(self):
        specs = dict(visuals.pages())["Custos Logísticos"]
        chart = next(s for s in specs if s.visual_type == "lineChart")
        self.assertIn("US$ mi", chart.title)
        self.assertEqual(dict(chart.roles)["Y"][0].name, "Frete gráfico (mi)")
        self.assertIn("Frete por kg", {f.name for f in dict(chart.roles)["Tooltips"]})
        self.assertIn("Registros sem peso", {f.name for f in dict(chart.roles)["Tooltips"]})

    def test_29_detail_validator_rejects_mixed_dimension_and_measure_query(self):
        project = self.unpack()
        p, doc = self.visual(project, "tableEx")
        role = doc["visual"]["query"]["queryState"]["Values"]["projections"]
        role[0] = visuals.projection(visuals.meas("Valor CIF"))
        p.write_text(json.dumps(doc))
        with self.assertRaisesRegex(bi.PowerBIProjectError, "mesma base"):
            bi.validate_powerbi_project(project)

    def test_30_axis_measures_preserve_value_and_explicit_chart_units(self):
        measures = {m.name: m for t in semantic.tables() for m in t.measures}
        for name, base in [
            ("CIF gráfico (mi)", "CIF (US$ mi)"),
            ("Frete gráfico (mi)", "Frete (US$ mi)"),
            ("FOB municipal gráfico (mi)", "FOB municipal (US$ mi)"),
        ]:
            self.assertEqual(measures[name].expression, "[" + base + "]")
            self.assertEqual(measures[name].fmt, "#,##0.0")
            self.assertTrue(measures[name].hidden)
        self.assertEqual(measures["Valor CIF"].fmt, semantic.MONEY)
        self.assertEqual(measures["CIF (US$ mi)"].fmt, semantic.MILLION)

    def test_31_nominal_trends_use_zero_baseline_and_price_is_explicit(self):
        for title, specs in visuals.pages():
            for spec in specs:
                if spec.visual_type != "lineChart":
                    continue
                objects = visuals.visual_payload(spec, title, 1)["visual"]["objects"]
                props = objects["valueAxis"][0]["properties"]
                if title == "Produtos e NCM":
                    self.assertIn("escala automática", spec.title)
                    self.assertNotIn("start", props)
                else:
                    self.assertEqual(props["start"]["expr"]["Literal"]["Value"], "0D")
