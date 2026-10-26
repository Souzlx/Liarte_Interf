"""Accepted comparisons export complete contributions and explicit missing values."""

from dataclasses import replace
from datetime import date
from decimal import localcontext
from pathlib import Path
from tempfile import TemporaryDirectory
from threading import Event
import unittest
from unittest.mock import PropertyMock, patch

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from liarte_core.inventory import (
    SCALE, Balance, Closing, InventoryData, Product, Sale, StockSelection, analyze_inventory,
)
from liarte_core.inventory_compare import compare_inventory_months
from liarte_core.inventory_compare_export import append_comparison_sheets
from liarte_core.inventory_export import export_inventory
from liarte_core.inventory_monthly import analyze_monthly_inventory
from liarte_core.query import Cancelled


def comparison_export_data():
    codes = tuple(f"P{index:03d}" for index in range(16))
    products = tuple(Product(code, "Produto " + code, "Grupo", "PA", "KG") for code in codes)
    closings = []
    sales = []
    for index, code in enumerate(codes):
        branch = "01" if index % 2 == 0 else "02"
        for period, quantity, cost in (
            ("2026-01", 10 + index, 1000 + 100 * index),
            ("2026-03", 12 + index, 1200 + 100 * index),
        ):
            closings.append(Closing(
                code, branch, "01", "KG", period, quantity * SCALE, cost, 1, 30,
            ))
        sales.extend((
            Sale(code, branch, "KG", date(2026, 1, 1), (index + 1) * SCALE, 100),
            Sale(code, branch, "KG", date(2026, 3, 2),
                 (-1 if index == 15 else 2 * (index + 1)) * SCALE, 100),
        ))
    return InventoryData(
        date(2026, 4, 15), "Comparador sintético.pbix", "a" * 64,
        products + (
            Product("OLD", "Somente na base", "Grupo", "PA", "KG"),
            Product("NEW", "Somente no comparado", "Grupo", "PA", "UN"),
        ),
        (Balance("P000", "01", "01", "KG", 10 * SCALE, 1000),),
        tuple(closings) + (
            Closing("OLD", "04", "01", "KG", "2026-01", 2 * SCALE, 200, 1, 30),
            Closing("NEW", "03", "01", "UN", "2026-03", 6 * SCALE, 600, 1, 30),
            Closing("P000", "01", "02", "L", "2026-01", 3 * SCALE, 300, 1, 30),
            Closing("P000", "01", "02", "L", "2026-03", 5 * SCALE, 500, 1, 30),
        ),
        tuple(sales) + (
            Sale("P000", "01", "L", date(2026, 1, 1), 2 * SCALE, 100),
            Sale("P000", "01", "L", date(2026, 3, 2), 3 * SCALE, 100),
        ),
    )


class InventoryCompareExportTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "comparison.xlsx"
        self.data = comparison_export_data()
        self.selection = StockSelection("2026-01", "2026-04")

    def prepare(self, data=None, selection=None, base="2026-01", target="2026-03"):
        data = self.data if data is None else data
        selection = self.selection if selection is None else selection
        analysis = analyze_inventory(data, selection)
        monthly = analyze_monthly_inventory(data, selection, analysis=analysis)
        return data, selection, analysis, monthly, compare_inventory_months(monthly, base, target)

    def export(self, prepared=None, *, include=True):
        data, selection, analysis, monthly, comparison = prepared or self.prepare()
        export_inventory(
            data, analysis, selection, self.path, Event(), monthly=monthly,
            comparison=comparison if include else None,
        )
        book = load_workbook(self.path)
        self.addCleanup(book.close)
        return book

    @staticmethod
    def rows(sheet):
        values = list(sheet.iter_rows(values_only=True))
        return [dict(zip(values[0], row)) for row in values[1:]]

    def test_no_comparison_keeps_the_previous_export_without_choosing_a_pair(self):
        prepared = self.prepare()
        with patch(
            "liarte_core.inventory_compare.compare_inventory_months",
            side_effect=AssertionError("Export must not select comparison months"),
        ):
            book = self.export(prepared, include=False)
        self.assertNotIn("Comparação mensal", book.sheetnames)
        self.assertNotIn("Contribuições", book.sheetnames)
        self.assertIn("Análise mensal", book.sheetnames)
        self.assertIn("Alertas", book.sheetnames)

    def test_complete_stock_and_quantity_contributions_reconcile_and_keep_scope(self):
        book = self.export()
        summary = self.rows(book["Comparação mensal"])
        stock = next(row for row in summary if row["Indicador"] == "Estoque a custo")
        self.assertEqual((stock["Valor na base"], stock["Valor no comparado"]), (285, 323))
        self.assertEqual(stock["Diferença observada"], 38)
        self.assertEqual(stock["Diferença pareada"], 34)
        self.assertEqual(stock["Valor só no comparado"], 6)
        self.assertEqual(stock["Valor só na base"], 2)
        self.assertEqual(stock["Comparação disponível"], "Sim")
        self.assertEqual(stock["Mês base"], "2026-01")
        self.assertEqual(stock["Mês comparado"], "2026-03")
        self.assertEqual(stock["Data da base"].date(), self.data.as_of)
        self.assertIn("start=2026-01", stock["Filtros"])
        rows = self.rows(book["Contribuições"])
        products = [row for row in rows if row["Indicador"] == "Estoque a custo" and row["Dimensão"] == "Produto"]
        self.assertEqual(len(products), 18)
        self.assertGreater(len(products), 12)
        self.assertEqual(sum(row["Contribuição observada"] for row in products), 38)
        branches = [row for row in rows if row["Indicador"] == "Estoque a custo" and row["Dimensão"] == "Filial"]
        self.assertEqual({row["Código"] for row in branches}, {"01", "02", "03", "04"})
        self.assertEqual(sum(row["Contribuição observada"] for row in branches), 38)
        new = next(row for row in products if row["Código"] == "NEW")
        old = next(row for row in products if row["Código"] == "OLD")
        self.assertIsNone(new["Valor na base"])
        self.assertIsNone(old["Valor no comparado"])
        self.assertEqual(new["Contribuição observada"], 6)
        self.assertEqual(old["Contribuição observada"], -2)
        quantities = {row["UM"]: row for row in summary if row["Indicador"] == "Quantidade em estoque"}
        self.assertEqual(quantities["KG"]["Diferença observada"], 30)
        self.assertEqual(quantities["L"]["Diferença observada"], 2)
        self.assertIsNone(quantities["UN"]["Valor na base"])
        self.assertEqual(quantities["UN"]["Valor no comparado"], 6)
        self.assertIsNone(quantities["UN"]["Diferença observada"])
        rules = dict(book["Leitura dos indicadores"].iter_rows(min_row=2, values_only=True))
        self.assertIn("não comprovam", rules["Comparação mensal · reconciliação"])
        self.assertEqual(book["Contribuições"].freeze_panes, "A2")
        self.assertEqual(book["Contribuições"].auto_filter.ref, f"A1:T{len(rows) + 1}")

    def test_sales_returns_zero_and_missing_month_remain_distinct(self):
        book = self.export()
        summary = self.rows(book["Comparação mensal"])
        sales = {row["UM"]: row for row in summary if row["Indicador"] == "Vendas líquidas"}
        self.assertEqual(sales["KG"]["Valor na base"], 136)
        self.assertEqual(sales["KG"]["Valor no comparado"], 239)
        self.assertEqual(sales["KG"]["Diferença observada"], 103)
        self.assertEqual(sales["L"]["Diferença observada"], 1)
        self.assertEqual(sales["UN"]["Valor na base"], 0)
        self.assertEqual(sales["UN"]["Valor no comparado"], 0)
        self.assertEqual(sales["UN"]["Diferença observada"], 0)
        self.assertIsNone(sales["UN"]["Variação (%)"])
        returns = [row for row in self.rows(book["Contribuições"]) if (
            row["Indicador"] == "Vendas líquidas" and row["Dimensão"] == "Produto"
            and row["Código"] == "P015" and row["UM"] == "KG"
        )]
        self.assertEqual(len(returns), 1)
        self.assertEqual(returns[0]["Valor no comparado"], -1)
        unavailable = self.export(self.prepare(target="2026-04"))
        for row in self.rows(unavailable["Comparação mensal"]):
            self.assertEqual(row["Comparação disponível"], "Não")
            self.assertIsNone(row["Diferença observada"])
            self.assertTrue(row["Motivo e abrangência"])

    def test_warehouse_comparison_keeps_sales_unavailable_and_describes_filter(self):
        book = self.export(self.prepare(selection=replace(self.selection, warehouse="01")))
        rows = self.rows(book["Comparação mensal"])
        self.assertTrue(any(row["Indicador"] == "Estoque a custo" and row["Comparação disponível"] == "Sim" for row in rows))
        for row in rows:
            if row["Indicador"] == "Vendas líquidas":
                self.assertEqual(row["Comparação disponível"], "Não")
                self.assertIsNone(row["Valor na base"])
                self.assertIsNone(row["Diferença observada"])
                self.assertIn("warehouse=01", row["Filtros"])
        self.assertFalse(any(row["Indicador"] == "Vendas líquidas" for row in self.rows(book["Contribuições"])))

    def test_reverse_pair_is_exported_as_accepted_without_recomputing_it(self):
        prepared = self.prepare(base="2026-03", target="2026-01")
        with patch(
            "liarte_core.inventory_compare.compare_inventory_months",
            side_effect=AssertionError("Accepted comparison must be reused"),
        ):
            book = self.export(prepared)
        stock = next(row for row in self.rows(book["Comparação mensal"]) if row["Indicador"] == "Estoque a custo")
        self.assertEqual((stock["Mês base"], stock["Mês comparado"]), ("2026-03", "2026-01"))
        self.assertEqual(stock["Diferença observada"], -38)

    def test_extreme_values_remain_exact_and_product_text_is_never_a_formula(self):
        value = 10**16 + 1
        data = InventoryData(
            date(2026, 3, 31), "Synthetic precision", "b" * 64,
            (Product("=1+1", "=HYPERLINK(1)", "Grupo", "PA", "UN"),), (),
            (
                Closing("=1+1", "01", "01", "UN", "2026-01", value, value, 1, 30),
                Closing("=1+1", "01", "01", "UN", "2026-03", 2 * value, 2 * value, 1, 30),
            ),
            (Sale("=1+1", "01", "UN", date(2026, 1, 1), SCALE, 100),),
        )
        prepared = self.prepare(data=data, selection=StockSelection("2026-01", "2026-03"))
        with localcontext() as context:
            context.prec = 6
            book = self.export(prepared)
        summary = book["Comparação mensal"]
        self.assertEqual(summary["E2"].value, "100000000000000.01")
        self.assertEqual(summary["E2"].number_format, "@")
        self.assertIn("15 dígitos", summary["E2"].comment.text)
        contributions = book["Contribuições"]
        self.assertEqual(contributions["C2"].value, "=1+1")
        self.assertEqual(contributions["D2"].value, "=HYPERLINK(1)")
        self.assertEqual(contributions["C2"].data_type, "s")
        self.assertEqual(contributions["D2"].data_type, "s")
        self.assertEqual(contributions["H2"].value, "100000000000000.01")

    def test_tables_split_without_dropping_rows_or_scanning_max_row_per_append(self):
        data, selected, _, _, comparison = self.prepare()
        book = Workbook()
        self.addCleanup(book.close)
        expected = sum(
            len(metric[dimension])
            for metric in [comparison["stock"], *comparison["quantityByUnit"].values(), *comparison["salesByUnit"].values()]
            for dimension in ("products", "branches")
        )
        with (
            patch("liarte_core.inventory_compare_export._MAX_ROWS", 5),
            patch.object(Worksheet, "max_row", new_callable=PropertyMock, side_effect=AssertionError("quadratic row scan")),
        ):
            layouts = append_comparison_sheets(book, comparison, data, selected, Event())
        summary_sheets = [name for name in layouts if name.startswith("Comparação mensal")]
        contribution_sheets = [name for name in layouts if name.startswith("Contribuições")]
        self.assertGreater(len(summary_sheets), 1)
        self.assertGreater(len(contribution_sheets), 1)
        self.assertEqual(sum(book[name]._current_row - 1 for name in contribution_sheets), expected)
        self.assertEqual(sum(book[name]._current_row - 1 for name in summary_sheets), 1 + 2 * len(comparison["units"]))
        self.assertTrue(all(book[name]._current_row <= 5 for name in layouts))

    def test_cancel_during_comparison_preserves_the_existing_workbook(self):
        data, selected, analysis, monthly, comparison = self.prepare()
        self.path.write_bytes(b"previous export")
        cancel = Event()

        def interrupted(*args, **kwargs):
            cancel.set()
            return append_comparison_sheets(*args, **kwargs)

        with (
            patch("liarte_core.inventory_compare_export.append_comparison_sheets", side_effect=interrupted),
            self.assertRaises(Cancelled),
        ):
            export_inventory(data, analysis, selected, self.path, cancel, monthly=monthly, comparison=comparison)
        self.assertEqual(self.path.read_bytes(), b"previous export")
        self.assertEqual(list(Path(self.temp.name).iterdir()), [self.path])
