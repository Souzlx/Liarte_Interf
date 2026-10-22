"""Contratos de dados abertos, precisão, interrupções e dez clientes concorrentes.

As massas abaixo são sintéticas e não são evidência de acesso ao MDIC. Os testes
de integração externa são executados separadamente e documentados como tal.
"""

from __future__ import annotations

import csv
import io
import os
import subprocess
import sys
import tempfile
import time
import unittest
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from threading import Event, Lock
from unittest.mock import patch
from urllib.error import HTTPError

from liarte_core.catalog import catalog
from liarte_core.comex import ComexProvider
from liarte_core.comex_central import CentralBroker, CentralServer
from liarte_core.comex_open_data import (
    BASE,
    PUBLICATION,
    GENERAL_HEADER,
    MUNICIPAL_HEADER,
    OfficialFiles,
    OfficialRedirects,
    OpenDataError,
    OpenDataRepository,
    download_lock,
    csv_rows,
    integer,
)
from liarte_core.comex_source import OfficialSourceClient
from liarte_core.domain import Period, Selection
from liarte_core.query import Cancelled
from liarte_engine import comexstat_importacoes as engine
from liarte_engine.comex_access import RequestCoordinator
from liarte_engine.comex_gateway_client import GatewayConnection, save_connection
from tests.test_comex import Transport
from tests.test_comex_access import TOKEN


def csv_bytes(fields, rows):
    output = io.StringIO(newline="")
    writer = csv.writer(output, delimiter=";", lineterminator="\n")
    writer.writerow(fields)
    writer.writerows(rows)
    return output.getvalue().encode("latin-1")


class Download(io.BytesIO):
    status = 200

    def __init__(self, raw, url, *, length=None):
        super().__init__(raw)
        self.url = url
        self.headers = {
            "Content-Length": str(len(raw) if length is None else length),
            "ETag": '"revision-1"',
        }

    def geturl(self):
        return self.url


class FilesTransport:
    def __init__(self):
        codes = sorted(p["ncm"] for p in catalog())
        headings = sorted({c[:4] for c in codes})
        self.documents = {
            PUBLICATION: (
                "<p>Atualizado em 06/10/2026 15h29</p>"
                "<h2>Últimos dados disponíveis: janeiro - setembro de 2026</h2>"
            ).encode(),
            BASE + "tabelas/NCM.csv": csv_bytes(
                ["CO_NCM", "NO_NCM_POR"], [(c, "Produto " + c) for c in codes]
            ),
            BASE + "tabelas/NCM_SH.csv": csv_bytes(
                ["CO_SH4", "NO_SH4_POR"], [(c, "Posição " + c) for c in headings]
            ),
            BASE + "tabelas/PAIS.csv": csv_bytes(
                ["CO_PAIS", "NO_PAIS"], [(160, "China"), (249, "Estados Unidos")]
            ),
            BASE + "tabelas/UF_MUN.csv": csv_bytes(
                ["CO_MUN_GEO", "NO_MUN_MIN", "SG_UF"], [(3550308, "São Paulo", "SP")]
            ),
            BASE + "tabelas/URF.csv": csv_bytes(
                ["CO_URF", "NO_URF"], [(817800, "PORTO DE SANTOS")]
            ),
        }
        self.general = []
        self.municipal = []
        for month in range(1, 10):
            for i, code in enumerate(codes):
                # Distinct units/vias must be aggregated in the current API dimensions.
                for via in (1, 4):
                    self.general.append(
                        [
                            2026,
                            month,
                            code,
                            10,
                            160,
                            "SP",
                            via,
                            817800,
                            50,
                            1000 + i,
                            100 + month,
                            7,
                            2,
                        ]
                    )
            # Outside the embedded catalog: still part of national validation.
            self.general.append([2026, month, 10011010, 10, 160, "SP", 1, 817800, 1, 33, 11, 1, 0])
            for code in headings:
                self.municipal.append([2026, month, code, 160, "SP", 3550308, 5000, 6000])
            self.municipal.append([2026, month, "1001", 160, "SP", 3550308, 44, 22])
        self.publish()
        self.calls = []
        self.lock = Lock()
        self.failure = None

    def publish(self):
        self.documents[BASE + "comexstat-bd/ncm/IMP_2026.csv"] = csv_bytes(
            GENERAL_HEADER, self.general
        )
        self.documents[BASE + "comexstat-bd/mun/IMP_2026_MUN.csv"] = csv_bytes(
            MUNICIPAL_HEADER, self.municipal
        )
        for municipal in (False, True):
            rows = self.municipal if municipal else self.general
            header = ["CO_ANO", "CO_MES", "KG_LIQUIDO", "VL_FOB", "QT_REGISTROS"]
            control = []
            for month in range(1, 10):
                chosen = [r for r in rows if r[1] == month]
                control.append(
                    [
                        2026,
                        month,
                        sum(r[6 if municipal else 9] for r in chosen),
                        sum(r[7 if municipal else 10] for r in chosen),
                        len(chosen),
                    ]
                )
            path = (
                "mun/IMP_TOTAIS_CONFERENCIA_MUN.csv"
                if municipal
                else "ncm/IMP_TOTAIS_CONFERENCIA.csv"
            )
            self.documents[BASE + "comexstat-bd/" + path] = csv_bytes(header, control)

    def __call__(self, request, **kwargs):
        with self.lock:
            self.calls.append((request.full_url, dict(request.headers)))
        if self.failure:
            raise self.failure
        if request.full_url not in self.documents:
            raise AssertionError("Unexpected external request: " + request.full_url)
        return Download(self.documents[request.full_url], request.full_url)


def body(municipal=False, *, start="2026-01", end="2026-09", control=False):
    return {
        "flow": "import",
        "monthDetail": True,
        "period": {"from": start, "to": end},
        "filters": [
            {
                "filter": "heading" if municipal else "ncm",
                "values": [int(p["ncm"][:4] if municipal else p["ncm"]) for p in catalog()],
            }
        ],
        "details": ["heading" if municipal else "ncm"]
        if control
        else engine.MUNICIPAL_DETAILS
        if municipal
        else engine.DEFAULT_DETAILS,
        "metrics": engine.MUNICIPAL_METRICS if municipal else engine.DEFAULT_METRICS,
    }


class OpenDataTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.directory = Path(self.temp.name)
        self.transport = FilesTransport()
        self.files = OfficialFiles(self.directory / "raw", opener=self.transport)
        self.repo = OpenDataRepository(self.directory, files=self.files)

    def tearDown(self):
        self.temp.cleanup()

    def query(self, municipal=False, **kwargs):
        return self.repo.query(
            body(municipal, **kwargs), municipal, lambda: None, lambda *args: None
        )["data"]["list"]

    def test_full_catalog_nine_months_exact_and_dimensions_preserved(self):
        rows = self.query()
        self.assertEqual(len(rows), len(catalog()) * 9)
        self.assertEqual(
            sum(int(r["metricKG"]) for r in rows),
            sum(r[9] for r in self.transport.general if str(r[2]) in self.repo.codes),
        )
        for r in rows:
            self.assertEqual(
                int(r["metricCIF"]),
                int(r["metricFOB"]) + int(r["metricFreight"]) + int(r["metricInsurance"]),
            )
            self.assertEqual(r["urf"], "0817800 - PORTO DE SANTOS")
            self.assertEqual(r["__liarte_source"], "mdic-csv")

    def test_municipal_is_full_heading_not_selected_ncm_and_has_no_cif(self):
        rows = self.query(True)
        self.assertEqual(len(rows), len(self.repo.headings) * 9)
        self.assertEqual({r["headingCode"] for r in rows}, self.repo.headings)
        self.assertEqual({r["noMunMinsgUf"] for r in rows}, {"São Paulo"})
        self.assertTrue(all("metricCIF" not in r for r in rows))

    def test_national_totals_include_products_outside_catalog(self):
        self.transport.general.pop()
        self.transport.documents[BASE + "comexstat-bd/ncm/IMP_2026.csv"] = csv_bytes(
            GENERAL_HEADER, self.transport.general
        )
        with self.assertRaisesRegex(OpenDataError, "divergentes"):
            self.query()
        self.assertFalse(list((self.directory / "indices").glob("*")))

    def test_independent_raw_controls_equal_detailed_sums(self):
        rows = self.query()
        controls = self.query(control=True)
        self.assertEqual(len(rows), len(controls))
        for r, c in zip(rows, controls, strict=True):
            self.assertEqual(
                {m: r[m] for m in engine.DEFAULT_METRICS}, {m: c[m] for m in engine.DEFAULT_METRICS}
            )

    def test_absent_freight_is_not_silently_zero(self):
        path = BASE + "comexstat-bd/ncm/IMP_2026.csv"
        self.transport.documents[path] = csv_bytes(
            GENERAL_HEADER[:-2], [r[:-2] for r in self.transport.general]
        )
        with self.assertRaisesRegex(OpenDataError, "VL_FRETE"):
            self.query()

    def test_missing_unselected_month_is_detected(self):
        self.transport.general = [r for r in self.transport.general if r[1] != 1]
        self.transport.documents[BASE + "comexstat-bd/ncm/IMP_2026.csv"] = csv_bytes(
            GENERAL_HEADER, self.transport.general
        )
        with self.assertRaisesRegex(OpenDataError, "meses"):
            self.query(start="2026-09")

    def test_unknown_reference_is_rejected(self):
        self.transport.documents[BASE + "tabelas/URF.csv"] = csv_bytes(
            ["CO_URF", "NO_URF"], [(999999, "Other")]
        )
        with self.assertRaisesRegex(OpenDataError, "Código sem descrição"):
            self.query()

    def test_non_integral_or_negative_published_units_are_rejected(self):
        for value in ("1.5", "-3", "1,000", "NaN", "Infinity", "1e5", "١٢", "1" * 130):
            with self.subTest(value=value), self.assertRaises(OpenDataError):
                integer(value, "test")
        self.assertEqual(integer("9007199254740993", "test"), 9007199254740993)

    def test_truncated_csv_row_is_rejected(self):
        path = BASE + "comexstat-bd/ncm/IMP_2026.csv"
        self.transport.documents[path] += b"2026;09;29161220\n"
        with self.assertRaisesRegex(OpenDataError, "incompleta"):
            self.query()

    def test_wrong_year_is_rejected(self):
        self.transport.general[0][0] = 2025
        self.transport.publish()
        with self.assertRaisesRegex(OpenDataError, "ano inválido"):
            self.query()

    def test_cached_different_months_and_controls_do_not_redownload(self):
        self.query()
        count = len(self.transport.calls)
        self.query(start="2026-09")
        self.query(control=True)
        self.assertEqual(count, len(self.transport.calls))

    def test_publication_change_revalidates_annual_file_references_and_controls(self):
        self.query()
        revision = {"period": "2026-09", "updated": "2026-10-07"}
        self.repo.query(body(), False, lambda: None, lambda *a: None, revision=revision)
        urls = [u for u, _ in self.transport.calls]
        for suffix in ("tabelas/NCM.csv", "ncm/IMP_2026.csv", "ncm/IMP_TOTAIS_CONFERENCIA.csv"):
            self.assertEqual(sum(u.endswith(suffix) for u in urls), 2)
        count = len(urls)
        self.repo.query(body(), False, lambda: None, lambda *a: None, revision=revision)
        self.assertEqual(count, len(self.transport.calls))

    def test_forced_refresh_revalidates_all_inputs(self):
        self.query()
        count = len(self.transport.calls)
        self.repo.query(body(), False, lambda: None, lambda *a: None, refresh=True)
        self.assertEqual(len(self.transport.calls), 2 * count)

    def test_redirect_is_refused_before_contacting_another_host(self):
        from urllib.request import Request

        handler = OfficialRedirects()
        with self.assertRaises(OpenDataError):
            handler.redirect_request(
                Request(BASE), None, 302, "Found", {}, "http://127.0.0.1/private"
            )

    def test_corrupted_index_is_rebuilt_from_verified_source(self):
        expected = self.query()
        index = next((self.directory / "indices").glob("*.json.z"))
        index.write_bytes(b"invalid")
        self.assertEqual(expected, self.query())

    def test_corrupted_raw_file_is_not_trusted_after_restart(self):
        self.query()
        source = self.directory / "raw/comexstat-bd/ncm/IMP_2026.csv"
        raw = source.read_bytes()
        source.write_bytes(raw.replace(b";SP;", b";RJ;", 1))
        repo = OpenDataRepository(
            self.directory, files=OfficialFiles(self.directory / "raw", opener=self.transport)
        )
        expected = self.query()
        actual = repo.query(body(), False, lambda: None, lambda *args: None)["data"]["list"]
        self.assertEqual(expected, actual)
        self.assertEqual(source.read_bytes(), raw)

    def test_304_requires_verified_existing_file(self):
        p, meta = self.files.fetch(
            "tabelas/PAIS.csv", check=lambda: None, notify=lambda *args: None
        )
        self.files.clock = lambda: meta["checked_at"] + 600

        def cached(request, **kwargs):
            self.assertEqual(request.headers.get("If-none-match"), '"revision-1"')
            raise HTTPError(request.full_url, 304, "Not Modified", {}, None)

        self.files.opener = cached
        self.assertEqual(
            p,
            self.files.fetch("tabelas/PAIS.csv", check=lambda: None, notify=lambda *args: None)[0],
        )

    def test_truncated_download_preserves_prior_file(self):
        p, meta = self.files.fetch(
            "tabelas/PAIS.csv", check=lambda: None, notify=lambda *args: None
        )
        previous = p.read_bytes()
        self.files.clock = lambda: meta["checked_at"] + 600
        self.files.opener = lambda request, **kw: Download(b"short", request.full_url, length=100)
        with self.assertRaises(OpenDataError):
            self.files.fetch("tabelas/PAIS.csv", check=lambda: None, notify=lambda *args: None)
        self.assertEqual(previous, p.read_bytes())
        self.assertFalse(list(p.parent.glob("*.part")))

    def test_download_limit_and_redirect_are_enforced(self):
        for response in (
            Download(b"a", "https://evil.example/file"),
            Download(b"a", BASE + "tabelas/PAIS.csv", length=1000),
        ):
            with self.subTest(response=response), tempfile.TemporaryDirectory() as temporary:
                files = OfficialFiles(temporary, opener=lambda *a, **k: response)
                with self.assertRaises(OpenDataError):
                    files.fetch(
                        "tabelas/PAIS.csv", check=lambda: None, notify=lambda *a: None, max_size=100
                    )

    def test_cancelled_download_never_publishes_partial(self):
        checks = 0

        def check():
            nonlocal checks
            checks += 1
            if checks >= 4:
                raise Cancelled()

        with self.assertRaises(Cancelled):
            self.files.fetch("tabelas/PAIS.csv", check=check, notify=lambda *args: None)
        self.assertFalse((self.directory / "raw/tabelas/PAIS.csv").exists())
        self.assertFalse(list((self.directory / "raw").rglob("*.part")))

    def test_file_server_failure_is_cooled_down_without_more_requests(self):
        self.transport.failure = HTTPError(BASE, 503, "unavailable", {}, io.BytesIO(b"failure"))
        for _ in range(3):
            with self.assertRaises(OpenDataError):
                self.files.fetch("tabelas/PAIS.csv", check=lambda: None, notify=lambda *args: None)
        self.assertEqual(len(self.transport.calls), 1)

    def test_file_server_retry_after_longer_than_an_hour_is_preserved(self):
        now = time.time()
        self.files.clock = lambda: now
        self.transport.failure = HTTPError(
            BASE, 503, "down", {"Retry-After": "7200"}, io.BytesIO(b"down")
        )
        with self.assertRaises(OpenDataError):
            self.files.fetch("tabelas/PAIS.csv", check=lambda: None, notify=lambda *a: None)
        now += 3601
        with self.assertRaises(OpenDataError):
            self.files.fetch("tabelas/PAIS.csv", check=lambda: None, notify=lambda *a: None)
        self.assertEqual(len(self.transport.calls), 1)

    def test_publication_metadata_parses_current_official_format(self):
        self.assertEqual(
            self.repo.availability(lambda: None, lambda *a: None),
            {"year": "2026", "monthNumber": "09", "updated": "2026-10-06"},
        )

    def test_changed_publication_invalid_dates_rejected(self):
        self.transport.documents[PUBLICATION] = b"invalid page"
        with self.assertRaises(OpenDataError):
            self.repo.availability(lambda: None, lambda *a: None)

    def test_municipal_availability_does_not_assume_general_coverage(self):
        path = BASE + "comexstat-bd/mun/IMP_TOTAIS_CONFERENCIA_MUN.csv"
        lines = self.transport.documents[path].splitlines(keepends=True)
        self.transport.documents[path] = b"".join(lines[:-1])
        with self.assertRaisesRegex(OpenDataError, "revisões diferentes"):
            self.repo.availability(lambda: None, lambda *a: None, municipal=True)

    def test_ten_download_clients_obtain_same_file_once(self):
        def request(_):
            files = OfficialFiles(self.directory / "raw", opener=self.transport)
            return files.fetch("tabelas/PAIS.csv", check=lambda: None, notify=lambda *a: None)[1][
                "sha256"
            ]

        with ThreadPoolExecutor(max_workers=10) as pool:
            results = list(pool.map(request, range(10)))
        self.assertEqual(len(set(results)), 1)
        self.assertEqual(len(self.transport.calls), 1)

    @unittest.skipIf(
        sys.platform == "win32", "POSIX subprocess crash; Windows locking requires native QA"
    )
    def test_download_lock_is_released_by_os_after_crash(self):
        script = "from pathlib import Path\nimport time,sys\nfrom liarte_core.comex_open_data import download_lock\nwith download_lock(Path(sys.argv[1]),lambda:None):\n print('ready',flush=True)\n time.sleep(30)\n"
        lock = self.directory / "crash.lock"
        process = subprocess.Popen(
            [sys.executable, "-c", script, str(lock)],
            stdout=subprocess.PIPE,
            text=True,
            env={**os.environ, "PYTHONPATH": os.pathsep.join(sys.path)},
        )
        try:
            self.assertEqual(process.stdout.readline().strip(), "ready")
            process.kill()
            process.wait(timeout=5)
            start = time.monotonic()
            with download_lock(lock, lambda: None):
                self.assertLess(time.monotonic() - start, 1)
        finally:
            if process.poll() is None:
                process.kill()
                process.wait(timeout=5)
            process.stdout.close()

    def test_utf8_sample_boundary_does_not_switch_to_latin1(self):
        path = self.directory / "encoding.csv"
        path.write_bytes(b"K;V\n1;" + b"a" * 65529 + "é".encode() + b"\n")
        with csv_rows(path) as rows:
            self.assertTrue(next(rows)["V"].endswith("é"))

    def test_latin1_accent_after_long_ascii_prefix_is_decoded_correctly(self):
        path = self.directory / "late_accent.csv"
        path.write_bytes(b"K;V\n" + b"1;ASCII\n" * 15000 + "2;São Paulo\n".encode("latin-1"))
        with csv_rows(path) as rows:
            last = None
            for last in rows:
                pass
        self.assertEqual(last["V"], "São Paulo")


class SourceFallbackTests(unittest.TestCase):
    setUp = OpenDataTests.setUp
    tearDown = OpenDataTests.tearDown

    def client(self):
        client = OfficialSourceClient(
            repository=self.repo,
            min_interval=0,
            max_attempts=9,
            coordinator_path=self.directory / "api.sqlite3",
        )
        client.api._network_opener = Transport()
        return client

    def test_successful_files_do_not_use_api(self):
        client = self.client()
        client.api._network_opener = lambda *a, **k: self.fail("Unexpected API request")
        data = client.request_json("POST", "/general?language=pt", body())
        self.assertEqual(len(data["data"]["list"]), len(catalog()) * 9)

    def test_new_source_exports_excel_pdf_and_powerbi_with_preserved_scope(self):
        from liarte_core.exports import ExportRequest, export_document
        from openpyxl import load_workbook
        from pypdf import PdfReader
        from zipfile import ZipFile

        def factory(**kwargs):
            client = engine.ComexStatClient(**kwargs)
            origin = OfficialSourceClient(repository=self.repo, **kwargs)

            def receive(method, path, json_body=None):
                return origin.request_json(method, path, json_body)

            client.request_json = receive
            return client

        provider = ComexProvider(
            self.directory / "client", client_factory=factory, automatic_access=True
        )
        period = Period("2026-01", "2026-09")
        snapshot = provider.load(period, Event(), lambda *a: None)
        self.assertEqual(snapshot.source_channels, ("mdic-csv",))
        selection = Selection(period, ncm="29161220", country="China")
        for kind, suffix in (("excel", ".xlsx"), ("pdf", ".pdf"), ("powerbi", ".zip")):
            destination = self.directory / ("report" + suffix)
            report = export_document(
                ExportRequest(kind, destination, snapshot, selection),
                Event(),
                lambda *a: None,
                self.directory / "audit",
            )
            self.assertTrue(destination.is_file())
            if kind == "excel":
                book = load_workbook(destination)
                try:
                    self.assertEqual(book["Base geral"].max_row, len(snapshot.general) + 1)
                    self.assertEqual(book["Base municipal"].column_dimensions["E"].width, 64)
                    values = list(book["Base geral"].values)
                    total = sum(row[list(values[0]).index("CIF")] for row in values[1:])
                    self.assertEqual(total, sum(int(row["metricCIF"]) for row in snapshot.general))
                    self.assertTrue(
                        all(
                            c.alignment.vertical == "center"
                            for sheet in book
                            for row in sheet
                            for c in row
                            if c.value is not None
                        )
                    )
                finally:
                    book.close()
            elif kind == "pdf":
                text = "".join(page.extract_text() for page in PdfReader(destination).pages)
                self.assertIn("Acrilato de etila", text)
                self.assertNotIn("29161220", text)
            else:
                self.assertGreaterEqual(report["pages"], 6)
                with ZipFile(destination) as archive:
                    self.assertIsNone(archive.testzip())
        provider.close()

    def test_verification_script_uses_valid_contract_and_always_closes_provider(self):
        import verificar_comex
        from liarte_core.comex_central import validate_submission
        from liarte_engine.comex_gateway_client import PROTOCOL
        from unittest.mock import Mock

        def factory(**kwargs):
            client = engine.ComexStatClient(**kwargs)
            origin = OfficialSourceClient(repository=self.repo, **kwargs)

            def receive(method, path, json_body=None):
                validate_submission(
                    {
                        "protocol": PROTOCOL,
                        "request_id": "a" * 32,
                        "method": method,
                        "path": path,
                        "body": json_body,
                        "revision": None,
                        "refresh": False,
                    }
                )
                return origin.request_json(method, path, json_body)

            client.request_json = receive
            return client

        provider = ComexProvider(
            self.directory / "verify", client_factory=factory, automatic_access=True
        )
        close = Mock(wraps=provider.close)
        provider.close = close
        with patch.object(verificar_comex, "ComexProvider", return_value=provider):
            result = verificar_comex.verify(
                self.directory, self.directory / "verify.json", start="2026-01", end="2026-09"
            )
        self.assertEqual(result["source_channels"], ["mdic-csv"])
        self.assertEqual(result["general_rows"], len(catalog()) * 9)
        close.assert_called_once()
        provider.availability = Mock(side_effect=OpenDataError("offline"))
        with (
            patch.object(verificar_comex, "ComexProvider", return_value=provider),
            self.assertRaises(OpenDataError),
        ):
            verificar_comex.verify(self.directory, self.directory / "failure.json")
        self.assertEqual(close.call_count, 2)

    def test_one_api_attempt_if_file_server_is_unavailable(self):
        self.transport.failure = HTTPError(BASE, 503, "down", {}, io.BytesIO(b"down"))
        client = self.client()
        api = Transport()
        client.api._network_opener = api
        result = client.request_json("POST", "/general?language=pt", body(start="2026-09"))
        self.assertTrue(result["data"]["list"])
        self.assertEqual(len(api.calls), 1)
        self.assertEqual(client.api.max_attempts, 3)

    def test_files_can_be_read_during_api_cooldown(self):
        client = self.client()
        RequestCoordinator(self.directory / "api.sqlite3").defer(60)
        client.api._network_opener = lambda *a, **k: self.fail("API accessed during cooldown")
        self.assertTrue(client.request_json("POST", "/general?language=pt", body())["data"]["list"])

    def test_both_unavailable_does_not_retry_or_ignore_api_pause(self):
        self.transport.failure = HTTPError(BASE, 503, "down", {}, io.BytesIO(b"down"))
        client = self.client()
        RequestCoordinator(self.directory / "api.sqlite3").defer(60)
        client.api._network_opener = lambda *a, **k: self.fail("API accessed during cooldown")
        with self.assertRaises(engine.ComexRateLimitError) as result:
            client.request_json("POST", "/general?language=pt", body())
        self.assertIsNone(result.exception.observed_http_status)
        self.assertEqual(len(self.transport.calls), 1)

    def test_ten_users_share_one_annual_download_and_keep_exact_results(self):
        # Actual local HTTP broker; upstream files are controlled fixtures.
        from liarte_core import comex_source

        repo = self.repo
        broker = None
        server = None
        clients = []
        with patch.object(comex_source, "OpenDataRepository", return_value=repo):
            broker = CentralBroker(self.directory / "server", interval=0)
        try:
            server = CentralServer(broker, TOKEN, port=0)
            server.start()
            connection = GatewayConnection(True, f"http://127.0.0.1:{server.port}", TOKEN)
            for i in range(10):
                directory = self.directory / str(i)
                save_connection(directory / "conexao_comex.json", connection)
                clients.append(ComexProvider(directory, automatic_access=True))

            def query(provider):
                return provider.query(
                    Selection(Period("2026-01", "2026-09")), Event(), lambda *a: None
                )

            with ThreadPoolExecutor(max_workers=10) as pool:
                results = list(pool.map(query, clients))
            self.assertTrue(all(r.totals == results[0].totals for r in results))
            self.assertTrue(all(p.snapshot.totals_verified for p in clients))
            self.assertTrue(all(p.snapshot.source_channels == ("mdic-csv",) for p in clients))
            urls = [u for u, _ in self.transport.calls]
            for name in ("ncm/IMP_2026.csv", "mun/IMP_2026_MUN.csv"):
                self.assertEqual(sum(u.endswith(name) for u in urls), 1)
            self.assertLessEqual(len(urls), 10)
        finally:
            for provider in clients:
                provider.close()
            if server is not None:
                server.close()
            elif broker is not None:
                broker.close()
