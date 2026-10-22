#!/usr/bin/env python3
"""Gera relatórios Excel e Power BI com o cadastro interno de produtos e NCMs.

Uso básico:
    python comexstat_importacoes.py

Além do Excel, o comando gera por padrão um pacote Power BI portátil. Para
gerar somente a planilha, use ``--sem-power-bi``.

No uso interativo, o programa solicita o mês inicial e o mês final. Também é
possível informar o período diretamente no comando:
    python comexstat_importacoes.py --inicio 2026-01 --fim 2026-07

Para consultar automaticamente apenas o último mês publicado:
    python comexstat_importacoes.py --ultimo-mes

Dependência:
    python -m pip install openpyxl
"""

from __future__ import annotations

import argparse
import json
import math
import os
import re
import sys
import tempfile
import time
import unicodedata
import urllib.error
import urllib.request
import zipfile
from collections import defaultdict
from contextlib import nullcontext
from copy import copy
from dataclasses import dataclass
from decimal import Decimal
from email.utils import parsedate_to_datetime
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable, Iterable, Sequence

from .liarte_dados import IncrementalCache, create_manifest
from .numeric import from_units, published_units, sum_published
from .comex_access import AccessError, RequestCoordinator

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.exceptions import InvalidFileException
except ImportError as exc:  # pragma: no cover - mensagem amigável para o usuário
    raise SystemExit(
        "Dependência ausente. Execute: python -m pip install openpyxl"
    ) from exc


API_BASE = "https://api-comexstat.mdic.gov.br"
API_GENERAL = "/general"
API_CITIES = "/cities"
API_UPDATED = "/general/dates/updated"
GENERATOR_VERSION = "5.0.0-rc.9"
MIN_PERIOD = "1997-01"
QUERY_WINDOW_MONTHS = 12
MAX_API_RESPONSE_BYTES = 100 * 1024 * 1024
EXCEL_MAX_DATA_ROWS = 1_048_575
EXCEL_CELL_CHAR_LIMIT = 32_767
DEFAULT_DETAILS = ["ncm", "state", "urf", "country"]
DEFAULT_METRICS = [
    "metricFOB",
    "metricFreight",
    "metricInsurance",
    "metricCIF",
    "metricKG",
]
MUNICIPAL_DETAILS = ["city", "state", "heading", "country"]
MUNICIPAL_METRICS = ["metricFOB", "metricKG"]
SHEET_SUMMARY_NCM = "Resumo por NCM"
SHEET_SUMMARY_MUNICIPAL = "Resumo municipal"
SHEET_GENERAL_DATA = "Base geral"
SHEET_MUNICIPAL_DATA = "Base municipal"
SHEET_PRODUCTS = "Produtos"
ABSOLUTE_NUMBER_FORMAT = "#,##0"
CURRENCY_FORMAT = '"US$" #,##0'
WEIGHT_FORMAT = '#,##0 "kg"'
UNIT_VALUE_FORMAT = '"US$" #,##0.00'
PERCENT_FORMAT = "0.00%"
SUMMARY_TAB_COLOR = "18B7B1"
DATA_TAB_COLOR = "1F4E78"
MONTH_TAB_COLOR = "5B9BD5"
REFERENCE_TAB_COLOR = "7A8996"
MONTH_NAMES = {
    1: "Janeiro",
    2: "Fevereiro",
    3: "Março",
    4: "Abril",
    5: "Maio",
    6: "Junho",
    7: "Julho",
    8: "Agosto",
    9: "Setembro",
    10: "Outubro",
    11: "Novembro",
    12: "Dezembro",
}
MONTH_ALIASES = {
    "jan": 1,
    "janeiro": 1,
    "feb": 2,
    "fev": 2,
    "fevereiro": 2,
    "mar": 3,
    "marco": 3,
    "abr": 4,
    "abril": 4,
    "apr": 4,
    "mai": 5,
    "maio": 5,
    "may": 5,
    "jun": 6,
    "junho": 6,
    "jul": 7,
    "julho": 7,
    "ago": 8,
    "agosto": 8,
    "aug": 8,
    "set": 9,
    "setembro": 9,
    "sep": 9,
    "out": 10,
    "outubro": 10,
    "oct": 10,
    "nov": 11,
    "novembro": 11,
    "dez": 12,
    "dezembro": 12,
    "dec": 12,
}

# O Comex Stat retorna o nome completo da UF. Nos artefatos de saída, a sigla
# oficial economiza espaço sem perder significado. Valores desconhecidos são
# preservados, em vez de serem convertidos por aproximação.
STATE_ABBREVIATIONS = {
    "ACRE": "AC",
    "ALAGOAS": "AL",
    "AMAPA": "AP",
    "AMAZONAS": "AM",
    "BAHIA": "BA",
    "CEARA": "CE",
    "DISTRITO FEDERAL": "DF",
    "ESPIRITO SANTO": "ES",
    "GOIAS": "GO",
    "MARANHAO": "MA",
    "MATO GROSSO": "MT",
    "MATO GROSSO DO SUL": "MS",
    "MINAS GERAIS": "MG",
    "PARA": "PA",
    "PARAIBA": "PB",
    "PARANA": "PR",
    "PERNAMBUCO": "PE",
    "PIAUI": "PI",
    "RIO DE JANEIRO": "RJ",
    "RIO GRANDE DO NORTE": "RN",
    "RIO GRANDE DO SUL": "RS",
    "RONDONIA": "RO",
    "RORAIMA": "RR",
    "SANTA CATARINA": "SC",
    "SAO PAULO": "SP",
    "SERGIPE": "SE",
    "TOCANTINS": "TO",
}

# Cadastro mestre incorporado ao aplicativo. Esta tupla elimina a dependência
# de uma planilha externa e é incluída no executável pelo próprio módulo. A
# ordem original dos vínculos é preservada para facilitar auditorias.
INTERNAL_CATALOG_VERSION = "2026-09-11"
INTERNAL_PRODUCT_NCM_DATA: tuple[tuple[str, str], ...] = (
    ("607", "39073022"),
    ("601R", "39073022"),
    ("AC7015", "38249941"),
    ("AC7050", "38249941"),
    ("AT6030", "38249941"),
    ("B1", "28151100"),
    ("B36", "29221100"),
    ("BA100", "39111029"),
    ("BA120", "39111029"),
    ("C16", "29156019"),
    ("DY-128", "39073022"),
    ("G10", "29161220"),
    ("G12", "29161230"),
    ("G13", "29161240"),
    ("G2", "29161410"),
    ("G20", "29161310"),
    ("G28", "29025000"),
    ("G3", "29161430"),
    ("G5", "29161490"),
    ("G501", "29171990"),
    ("LI4196", "28252020"),
    ("MOISIL SE-31", "39100019"),
    ("OF201-100", "39100011"),
    ("OF201-1000", "39100011"),
    ("OF201-12500", "39100012"),
    ("OF201-200", "39100011"),
    ("OF201-350", "39100011"),
    ("OF201-50", "39100011"),
    ("OF201-5000", "39100012"),
    ("OF201-60000", "39100012"),
    ("OF408", "39100011"),
    ("OF57", "39100011"),
    ("OL60", "27101931"),
    ("OLA3", "27101931"),
    ("OLP40", "27101931"),
    ("P50", "32061110"),
    ("PIB128", "39022000"),
    ("PIB30", "39022000"),
    ("PIB32", "39022000"),
    ("S103", "29153931"),
    ("S120", "29051210"),
    ("S130", "29153300"),
    ("S19", "29153100"),
    ("S2", "29141100"),
    ("S23", "29094310"),
    ("S28", "29021100"),
    ("S34", "29051220"),
    ("S45", "29141200"),
    ("S63", "29051210"),
    ("S6-IM", "29141300"),
    ("SILHB151", "28112290"),
    ("SILHL200", "28112290"),
    ("V31", "29053100"),
    ("W09-IM", "34024200"),
)
INTERNAL_PRODUCT_COUNT = 54
INTERNAL_NCM_COUNT = 35

RAW_HEADERS = [
    "Mês",
    "NCM",
    "Descrição",
    "UF",
    "URF",
    "País",
    "FOB",
    "Frete",
    "Seguro",
    "CIF",
    "Peso líq.",
    "Produtos",
    "CIF/kg",
]

MUNICIPAL_HEADERS = [
    "Mês",
    "Município fiscal",
    "UF",
    "SH4",
    "Descrição",
    "País",
    "FOB",
    "Peso líq.",
    "Produtos",
    "NCMs",
    "FOB/kg",
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ALTERNATE_ROW_FILL = PatternFill("solid", fgColor="DCE6F1")
THIN_GRAY = Side(style="thin", color="B8C7D3")
CELL_GRID = Border(
    left=THIN_GRAY,
    right=THIN_GRAY,
    top=THIN_GRAY,
    bottom=THIN_GRAY,
)
ILLEGAL_XML_CHARACTERS = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x84\x86-\x9f]")


class ComexStatError(RuntimeError):
    """Erro de comunicação ou de resposta da API do Comex Stat."""


class ComexRateLimitError(ComexStatError):
    """Pausa real da fonte, preservada separadamente do texto técnico da resposta."""

    def __init__(self, seconds, endpoint="", detail="", observed_http_status=429):
        if (
            type(seconds) not in (int, float)
            or not math.isfinite(seconds)
            or seconds < 0
        ):
            raise ValueError("Pausa da fonte inválida.")
        self.retry_after = math.ceil(seconds)
        self.retry_at = time.time() + seconds
        self.endpoint = endpoint
        self.observed_http_status = observed_http_status
        self.detail = str(detail)[:1500]
        origin = "HTTP 429" if observed_http_status == 429 else "código 429 da API"
        super().__init__(
            f"Comex Stat solicitou uma pausa ({origin}). Aguarde {self.retry_after} s antes de tentar novamente."
        )


@dataclass(frozen=True)
class ProductNcm:
    product: str
    ncm: str


def text_key(value: Any) -> str:
    """Normaliza cabeçalhos para comparação tolerante a acentos e espaços."""
    text = "" if value is None else str(value).strip()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"\s+", " ", text).upper()


def clean_text(value: Any) -> str:
    """Produz texto seguro e dentro do limite de uma célula do Excel."""
    text = "" if value is None else str(value).strip()
    text = ILLEGAL_XML_CHARACTERS.sub("", text)
    text = re.sub(r"\s+", " ", text.replace("\u00a0", " ")).strip()
    return text[:EXCEL_CELL_CHAR_LIMIT]


def state_abbreviation(value: Any) -> str:
    """Retorna a sigla oficial da UF e preserva rótulos não reconhecidos."""
    state = clean_text(value)
    if re.fullmatch(r"[A-Za-z]{2}", state):
        return state.upper()
    return STATE_ABBREVIATIONS.get(text_key(state), state)


def normalize_ncm(value: Any) -> str | None:
    """Converte NCM numérico ou formatado (ex.: 3907.30.22) em oito dígitos."""
    if value is None:
        return None
    if isinstance(value, bool):
        raise ValueError(f"NCM inválido: {value!r}")
    if isinstance(value, (int, float)):
        if isinstance(value, float) and (
            not math.isfinite(value) or not value.is_integer()
        ):
            raise ValueError(f"NCM inválido: {value!r}")
        numeric = int(value)
        if numeric < 0 or numeric > 99_999_999:
            raise ValueError(f"NCM inválido: {value!r}")
        return f"{numeric:08d}"

    text = str(value).strip()
    if not text:
        return None
    if re.fullmatch(r"\d+(?:\.0+)?", text):
        digits = text.split(".", 1)[0]
    else:
        if not re.fullmatch(r"[\d.\-/\s]+", text):
            raise ValueError(f"NCM inválido: {value!r}")
        digits = re.sub(r"\D", "", text)
    if not digits:
        return None
    # Um NCM pode perder somente o zero inicial ao ser armazenado como número
    # no Excel. Códigos com seis dígitos ou menos são incompletos e não devem
    # ser transformados silenciosamente em outro NCM.
    if len(digits) == 7:
        digits = digits.zfill(8)
    if len(digits) != 8:
        raise ValueError(f"NCM inválido: {value!r} (esperados 8 dígitos)")
    return digits


def internal_product_ncms() -> list[ProductNcm]:
    """Retorna uma cópia validada do cadastro incorporado ao aplicativo."""
    pairs: list[ProductNcm] = []
    seen: set[tuple[str, str]] = set()
    for row_number, (product_value, ncm_value) in enumerate(
        INTERNAL_PRODUCT_NCM_DATA,
        start=1,
    ):
        product = clean_text(product_value)
        ncm = normalize_ncm(ncm_value)
        if not product or ncm is None:
            raise RuntimeError(
                f"Cadastro interno inválido na linha {row_number}: produto ou NCM vazio."
            )
        key = (product, ncm)
        if key in seen:
            raise RuntimeError(
                f"Cadastro interno duplicado na linha {row_number}: {product} / {ncm}."
            )
        seen.add(key)
        pairs.append(ProductNcm(product=product, ncm=ncm))

    unique_ncms = {pair.ncm for pair in pairs}
    if len(pairs) != INTERNAL_PRODUCT_COUNT or len(unique_ncms) != INTERNAL_NCM_COUNT:
        raise RuntimeError(
            "O cadastro interno de produtos e NCMs não passou na verificação de "
            "integridade. Reinstale a versão oficial do aplicativo."
        )
    return pairs


def heading_from_ncm(value: Any) -> str:
    """Retorna a posição SH4 correspondente a um NCM de oito dígitos."""
    ncm = normalize_ncm(value)
    if ncm is None:
        raise ValueError("Não é possível obter a posição SH4 de um NCM vazio.")
    return ncm[:4]


def normalize_heading(value: Any) -> str:
    """Normaliza códigos SH4 recebidos da API para exatamente quatro dígitos."""
    if value is None or isinstance(value, bool):
        raise ValueError(f"Código SH4 inválido: {value!r}")
    text = str(value).strip()
    if re.fullmatch(r"\d+(?:\.0+)?", text):
        digits = text.split(".", 1)[0]
    else:
        if not re.fullmatch(r"[\d.\-/\s]+", text):
            raise ValueError(f"Código SH4 inválido: {value!r}")
        digits = re.sub(r"\D", "", text)
    digits = digits.zfill(4)
    if len(digits) != 4:
        raise ValueError(f"Código SH4 inválido: {value!r} (esperados 4 dígitos)")
    return digits


def validate_period(value: str) -> str:
    """Valida um período e o normaliza como AAAA-MM.

    Aceita AAAA-MM, MM/AAAA e nomes localizados, como "fevereiro de 2026".
    """
    text = str(value).strip()
    match = re.fullmatch(r"(\d{4})[-/](0?[1-9]|1[0-2])(?:-0?1)?", text)
    if match:
        return f"{int(match.group(1)):04d}-{int(match.group(2)):02d}"
    match = re.fullmatch(r"(0?[1-9]|1[0-2])[-/](\d{4})", text)
    if match:
        return f"{int(match.group(2)):04d}-{int(match.group(1)):02d}"

    normalized = text_key(text).casefold()
    normalized = re.sub(r"[.,]", " ", normalized)
    normalized = re.sub(r"\s+", " ", normalized).strip()
    match = re.fullmatch(r"([a-z]+)(?:\s+de)?\s+(\d{4})", normalized)
    if match:
        month = MONTH_ALIASES.get(match.group(1))
        if month:
            return f"{int(match.group(2)):04d}-{month:02d}"
    raise argparse.ArgumentTypeError(
        "Use AAAA-MM ou MM/AAAA, por exemplo 2026-04 ou 04/2026."
    )


def prompt_period(label: str, default: str | None = None) -> str:
    """Solicita um mês no terminal até receber um valor válido."""
    default_text = f" [{default}]" if default else ""
    while True:
        try:
            entered = input(f"{label} (AAAA-MM ou MM/AAAA){default_text}: ").strip()
        except EOFError as exc:
            raise ValueError(
                "Não foi possível ler o período pelo terminal. Informe --inicio e --fim no comando."
            ) from exc
        if not entered and default:
            return default
        try:
            return validate_period(entered)
        except argparse.ArgumentTypeError as exc:
            print(f"Valor inválido. {exc}", file=sys.stderr)


def choose_period_interactively() -> tuple[str, str]:
    """Permite selecionar o intervalo no próprio uso do script."""
    print("\nSelecione o período da consulta ao Comex Stat.", flush=True)
    start_period = prompt_period("Mês inicial")
    while True:
        end_period = prompt_period("Mês final", default=start_period)
        if period_index(start_period) <= period_index(end_period):
            return start_period, end_period
        print(
            "O mês final não pode ser anterior ao mês inicial. Tente novamente.",
            file=sys.stderr,
        )


def period_index(value: str) -> int:
    year, month = map(int, value.split("-"))
    return year * 12 + month


def period_from_index(index: int) -> str:
    """Converte o índice mensal usado internamente para AAAA-MM."""
    year, month_zero_based = divmod(index - 1, 12)
    return f"{year:04d}-{month_zero_based + 1:02d}"


def iter_periods(start_period: str, end_period: str) -> Iterable[str]:
    """Percorre todos os meses de um intervalo inclusivo."""
    start = period_index(validate_period(start_period))
    end = period_index(validate_period(end_period))
    if start > end:
        raise ValueError("O período inicial não pode ser posterior ao período final.")
    for index in range(start, end + 1):
        yield period_from_index(index)


def period_windows(
    start_period: str,
    end_period: str,
    window_months: int = QUERY_WINDOW_MONTHS,
) -> list[tuple[str, str]]:
    """Divide consultas longas em janelas menores para reduzir falhas da API."""
    if window_months < 1:
        raise ValueError("A janela de consulta deve conter ao menos um mês.")
    start = period_index(validate_period(start_period))
    end = period_index(validate_period(end_period))
    if start > end:
        raise ValueError("O período inicial não pode ser posterior ao período final.")
    windows: list[tuple[str, str]] = []
    window_start = start
    while window_start <= end:
        window_end = min(window_start + window_months - 1, end)
        windows.append((period_from_index(window_start), period_from_index(window_end)))
        window_start = window_end + 1
    return windows


def chunks(items: Sequence[str], size: int) -> Iterable[list[str]]:
    for start in range(0, len(items), size):
        yield list(items[start : start + size])


class ComexStatClient:
    """Cliente pequeno com limitação de ritmo e repetição para 429/5xx."""

    def __init__(
        self,
        timeout: float = 120.0,
        min_interval: float = 10.5,
        max_attempts: int = 6,
        *,
        use_cache: bool = False,
        cache_path: Path | None = None,
        refresh_cache: bool = False,
        source_revision: Any = None,
        progress_callback: Callable[[float | None, str], None] | None = None,
        cancel_check: Callable[[], None] | None = None,
        cancel_wait: Callable[[float], bool] | None = None,
        coordinator_path: Path | None = None,
        gateway=None,
        failure_callback=None,
        verify_official_totals: bool = True,
        window_months: int | None = QUERY_WINDOW_MONTHS,
    ) -> None:
        if not math.isfinite(timeout) or timeout <= 0:
            raise ValueError("O tempo limite da API deve ser um número positivo.")
        if not math.isfinite(min_interval) or min_interval < 0:
            raise ValueError("O intervalo da API deve ser um número não negativo.")
        if isinstance(max_attempts, bool) or not isinstance(max_attempts, int):
            raise ValueError("A quantidade de tentativas da API deve ser um inteiro.")
        if max_attempts < 1 or max_attempts > 20:
            raise ValueError(
                "A quantidade de tentativas da API deve ficar entre 1 e 20."
            )
        if not isinstance(verify_official_totals, bool):
            raise ValueError("A conferência adicional de totais deve ser booleana.")
        if window_months is not None and (
            isinstance(window_months, bool)
            or not isinstance(window_months, int)
            or window_months < 1
        ):
            raise ValueError(
                "A janela deve ser um inteiro positivo ou None para o período inteiro."
            )
        self.timeout = float(timeout)
        self.min_interval = float(min_interval)
        self.max_attempts = max_attempts
        self.last_request_at: float | None = None
        self.cache = IncrementalCache(cache_path) if use_cache else None
        self.refresh_cache = refresh_cache
        self.source_revision = source_revision
        self.progress_callback = progress_callback
        self.cancel_check = cancel_check
        self.cancel_wait = cancel_wait
        self.gateway = gateway
        self.failure_callback = failure_callback
        self.verify_official_totals = verify_official_totals
        self.window_months = window_months
        self.shared_cache_hits = 0
        self.network_requests = 0
        self._coordinator = (
            RequestCoordinator(coordinator_path)
            if coordinator_path is not None
            else None
        )
        self.headers = {
            "Accept": "application/json",
            "Content-Type": "application/json",
            "User-Agent": f"comexstat-importacoes-catalogo/{GENERATOR_VERSION}",
        }

    def close(self) -> None:
        return None

    def _check_cancel(self):
        if self.cancel_check is not None:
            self.cancel_check()

    def _wait(self, seconds):
        self._check_cancel()
        if self.cancel_wait is None:
            time.sleep(seconds)
        else:
            self.cancel_wait(seconds)
        self._check_cancel()

    def _notify(self, fraction: float | None, message: str) -> None:
        self._check_cancel()
        if self.progress_callback is not None:
            self.progress_callback(fraction, message)

    def _throttle(self) -> None:
        if self.last_request_at is None:
            return
        remaining = self.min_interval - (time.monotonic() - self.last_request_at)
        if remaining > 0:
            self._wait(remaining)

    @staticmethod
    def _retry_seconds(headers: Any, body: str, attempt: int) -> float:
        header = headers.get("Retry-After")
        if header:
            try:
                delay = float(header)
                if math.isfinite(delay) and delay >= 0:
                    return delay + 0.5
            except (ValueError, TypeError, OverflowError):
                try:
                    retry_at = parsedate_to_datetime(header)
                    if retry_at.tzinfo is None:
                        retry_at = retry_at.replace(tzinfo=timezone.utc)
                    try:
                        now = parsedate_to_datetime(headers.get("Date"))
                        if now.tzinfo is None:
                            now = now.replace(tzinfo=timezone.utc)
                    except (ValueError, TypeError, OverflowError):
                        now = datetime.now(timezone.utc)
                    return max(0.0, (retry_at - now).total_seconds()) + 0.5
                except (TypeError, ValueError, OverflowError):
                    pass
        match = re.search(r"(\d+)\s*segundo", body, flags=re.IGNORECASE)
        if match:
            # O texto genérico da API pode continuar pedindo só 10 s mesmo após
            # recusas consecutivas. Sem Retry-After HTTP, aplicar margem e recuo
            # conservadores em vez de repetir a mesma frequência que foi recusada.
            return max(
                float(match.group(1)) + 0.5, min(300.0, 60.0 * 2 ** (attempt - 1))
            )
        return min(300.0, 60.0 * 2 ** (attempt - 1))

    def _read_response(self, request):
        from .secure_http import urlopen

        opener = getattr(self, "_network_opener", None) or urlopen
        self.network_requests += 1
        with opener(request, timeout=self.timeout) as response:
            raw_status = getattr(response, "status", None)
            if raw_status is None:
                raw_status = response.getcode()
            if isinstance(raw_status, bool) or not isinstance(raw_status, int):
                raise ComexStatError(
                    "A API não informou um código de status HTTP válido."
                )
            status = raw_status
            headers = response.headers
            raw_length = headers.get("Content-Length")
            expected_length = None
            if raw_length:
                try:
                    expected_length = int(raw_length)
                    if expected_length > MAX_API_RESPONSE_BYTES:
                        raise ComexStatError(
                            "A resposta da API excedeu o limite seguro de 100 MB."
                        )
                except ValueError:
                    pass
            # Limita também a duração total de uma resposta que continue
            # chegando aos poucos; o timeout de socket, sozinho, reinicia
            # a cada leitura e pode manter a consulta aberta indefinidamente.
            deadline = time.monotonic() + max(30.0, min(300.0, self.timeout * 2))
            buffer = bytearray()
            read = getattr(response, "read1", response.read)
            while True:
                self._check_cancel()
                if time.monotonic() > deadline:
                    raise TimeoutError(
                        "O tempo total para receber a resposta da API foi excedido."
                    )
                block = read(min(256 * 1024, MAX_API_RESPONSE_BYTES + 1 - len(buffer)))
                if not block:
                    break
                buffer.extend(block)
                if len(buffer) > MAX_API_RESPONSE_BYTES:
                    break
            response_bytes = bytes(buffer)
            if (
                expected_length is not None
                and expected_length >= 0
                and len(response_bytes) != expected_length
            ):
                raise ComexStatError(
                    "A API retornou uma resposta incompleta; os dados não serão utilizados."
                )
            if len(response_bytes) > MAX_API_RESPONSE_BYTES:
                raise ComexStatError(
                    "A resposta da API excedeu o limite seguro de 100 MB."
                )
            try:
                response_text = response_bytes.decode("utf-8-sig")
            except UnicodeDecodeError as exc:
                raise ComexStatError(
                    "A API retornou texto com codificação inválida."
                ) from exc
        return status, headers, response_text

    def request_json(self, method: str, path: str, json_body=None):
        if self.gateway is not None:
            from .comex_gateway_client import request_shared

            local_request = getattr(self.gateway, "request_json", None)
            request = (
                local_request
                if local_request is not None
                else lambda *a, **kw: request_shared(self.gateway, *a, **kw)
            )
            try:
                payload, shared = request(
                    method,
                    path,
                    json_body,
                    revision=self.source_revision,
                    refresh=self.refresh_cache,
                    check=self._check_cancel,
                    wait=self._wait,
                    notify=self._notify,
                )
            except ComexRateLimitError as exc:
                if self._coordinator is not None:
                    self._coordinator.defer(exc.retry_after)
                raise
            self.shared_cache_hits += int(shared)
            return payload
        url = f"{API_BASE}{path}"
        last_error = None
        for attempt in range(1, self.max_attempts + 1):
            self._notify(
                None,
                f"Acessando o Comex Stat · tentativa {attempt}/{self.max_attempts}",
            )
            retry_delay = None
            if self._coordinator is None:
                self._throttle()
            try:
                lease = (
                    self._coordinator.acquire(
                        self.min_interval,
                        max(
                            120.0,
                            self.timeout * 2
                            + max(30.0, min(300.0, self.timeout * 2))
                            + 30,
                        ),
                        check=self._check_cancel,
                        wait=self._wait,
                        notify=self._notify,
                    )
                    if self._coordinator is not None
                    else nullcontext()
                )
                with lease as access:
                    request_data = (
                        json.dumps(json_body, ensure_ascii=False).encode("utf-8")
                        if json_body is not None
                        else None
                    )
                    request = urllib.request.Request(
                        url, data=request_data, headers=self.headers, method=method
                    )
                    try:
                        status, headers, response_text = self._read_response(request)
                    except urllib.error.HTTPError as exc:
                        status, headers = exc.code, exc.headers or {}
                        try:
                            response_text = exc.read(4096).decode(
                                "utf-8", errors="replace"
                            )
                        finally:
                            exc.close()
                    except (urllib.error.URLError, TimeoutError, OSError) as exc:
                        last_error = exc
                        if attempt < self.max_attempts:
                            retry_delay = min(60.0, 2.0**attempt)
                        status = None
                    finally:
                        self.last_request_at = time.monotonic()
                    if status is not None:
                        observed_http_status = status
                        payload = None
                        if 200 <= status < 300:
                            try:
                                payload = json.loads(response_text)
                            except json.JSONDecodeError as exc:
                                raise ComexStatError(
                                    "A API retornou uma resposta que não é JSON."
                                ) from exc
                            if not isinstance(payload, dict):
                                raise ComexStatError(
                                    "A API retornou um JSON em formato inesperado."
                                )
                            error = payload.get("error")
                            code = payload.get(
                                "code",
                                error.get("code") if isinstance(error, dict) else None,
                            )
                            if (payload.get("success") is False or error) and str(
                                code
                            ) == "429":
                                status = 429
                        if (
                            not 200 <= status < 300
                            and self.failure_callback is not None
                        ):
                            try:
                                self.failure_callback(
                                    {
                                        "http_status": status,
                                        "observed_http_status": observed_http_status,
                                        "endpoint": path,
                                        "response_excerpt": response_text[:1500],
                                        "captured_at": time.time(),
                                        "retry_after_header": str(
                                            headers.get("Retry-After") or ""
                                        )[:200],
                                    }
                                )
                            except Exception:
                                pass  # Diagnóstico não pode alterar a resposta ou a pausa da fonte.
                        if status == 429 or 500 <= status < 600:
                            delay = self._retry_seconds(headers, response_text, attempt)
                            if access is not None and (
                                status in {429, 503} or headers.get("Retry-After")
                            ):
                                access.defer(delay)
                            last_error = (
                                ComexRateLimitError(
                                    delay, path, response_text, observed_http_status
                                )
                                if status == 429
                                else ComexStatError(
                                    f"API retornou HTTP {status}: {response_text[:500]}"
                                )
                            )
                            if attempt < self.max_attempts:
                                retry_delay = delay
                        elif not 200 <= status < 300:
                            raise ComexStatError(
                                f"API retornou HTTP {status}: {response_text[:1000]}"
                            )
                        else:
                            if payload.get("success") is False or payload.get("error"):
                                raise ComexStatError(
                                    f"Consulta recusada pela API: {json.dumps(payload, ensure_ascii=False)[:1200]}"
                                )
                            return payload
            except AccessError as exc:
                raise ComexStatError(str(exc)) from exc
            if retry_delay is None:
                break
            self._notify(
                None, f"Aguardando o Comex Stat · nova tentativa em {retry_delay:.0f} s"
            )
            self._wait(retry_delay)
        if isinstance(last_error, ComexRateLimitError):
            raise last_error
        raise ComexStatError(f"Não foi possível acessar o Comex Stat: {last_error}")

    def latest_published_period(self) -> tuple[str, str | None]:
        payload = self.request_json("GET", API_UPDATED)
        data = payload.get("data") or {}
        if not isinstance(data, dict):
            raise ComexStatError(
                "A API não informou o último período em um formato válido."
            )
        year = str(data.get("year", ""))
        month = str(data.get("monthNumber", "")).zfill(2)
        try:
            period = validate_period(f"{year}-{month}")
        except argparse.ArgumentTypeError as exc:
            raise ComexStatError(
                "A API não informou um último período publicado válido."
            ) from exc
        self.source_revision = {"period": period, "updated": data.get("updated")}
        return period, data.get("updated")

    def _query_cached(
        self,
        ncms: Sequence[str],
        start_period: str,
        end_period: str,
        chunk_size: int,
        *,
        municipal: bool,
    ) -> list[dict[str, Any]]:
        if (
            isinstance(chunk_size, bool)
            or not isinstance(chunk_size, int)
            or not 1 <= chunk_size <= 200
        ):
            raise ValueError("O lote deve ser um inteiro entre 1 e 200.")
        if isinstance(ncms, (str, bytes)):
            raise ValueError("Informe uma sequência de NCMs.")
        normalized = [normalize_ncm(code) for code in ncms]
        if any(code is None for code in normalized):
            raise ValueError("A lista contém um NCM vazio.")
        start_period, end_period = (
            validate_period(start_period),
            validate_period(end_period),
        )
        if not MIN_PERIOD <= start_period <= end_period:
            raise ValueError(f"Período inválido. Os dados começam em {MIN_PERIOD}.")
        live = (
            self._query_municipal_imports_live
            if municipal
            else self._query_imports_live
        )
        if self.cache is None:
            return live(normalized, start_period, end_period, chunk_size)
        codes = (
            sorted({heading_from_ncm(code) for code in normalized})
            if municipal
            else sorted(set(normalized))
        )
        reader = (
            (
                lambda batch, first, last, size: live(
                    [ncm for ncm in normalized if heading_from_ncm(ncm) in batch],
                    first,
                    last,
                    size,
                )
            )
            if municipal
            else live
        )
        return self.cache.query(
            dataset="municipal" if municipal else "general",
            codes=codes,
            start=start_period,
            end=end_period,
            chunk_size=chunk_size,
            contract={
                "endpoint": API_BASE + (API_CITIES if municipal else API_GENERAL),
                "details": MUNICIPAL_DETAILS if municipal else DEFAULT_DETAILS,
                "metrics": MUNICIPAL_METRICS if municipal else DEFAULT_METRICS,
                "flow": "import",
                "language": "pt",
                "monthDetail": True,
                "totals_verification": int(self.verify_official_totals),
                "revision": self.source_revision,
            },
            refresh=self.refresh_cache,
            reader=reader,
            validator=validate_municipal_api_row if municipal else validate_api_row,
            code_of=(lambda row: normalize_heading(row["headingCode"]))
            if municipal
            else (lambda row: normalize_ncm(row["coNcm"])),
            period_of=row_period,
            window_months=self.window_months
            or (period_index(end_period) - period_index(start_period) + 1),
            progress_callback=self._notify,
        )

    def query_imports(
        self,
        ncms: Sequence[str],
        start_period: str,
        end_period: str,
        chunk_size: int,
    ) -> list[dict[str, Any]]:
        return self._query_cached(
            ncms, start_period, end_period, chunk_size, municipal=False
        )

    def query_municipal_imports(
        self,
        ncms: Sequence[str],
        start_period: str,
        end_period: str,
        chunk_size: int = 25,
    ) -> list[dict[str, Any]]:
        return self._query_cached(
            ncms, start_period, end_period, chunk_size, municipal=True
        )

    def _verify_totals(self, endpoint, body, detailed, *, municipal):
        """Reconcile month/code totals before admitting a batch into the cache.

        A well-formed response can still omit records. The control query uses
        the same official source with fewer dimensions, independently grouped.
        Cached filters and document exports do not repeat this network access.
        """
        field = "headingCode" if municipal else "coNcm"
        normalize = normalize_heading if municipal else normalize_ncm
        metrics = MUNICIPAL_METRICS if municipal else DEFAULT_METRICS
        identity = municipal_row_identity if municipal else api_row_identity
        signature = municipal_metric_signature if municipal else api_metric_signature
        unique = {}
        for index, row in enumerate(detailed):
            if index % 512 == 0:
                self._check_cancel()
            key = identity(row)
            if key in unique and signature(unique[key]) != signature(row):
                raise ComexStatError(
                    "A API retornou registros duplicados com valores divergentes."
                )
            unique[key] = row
        actual = {}
        for row in unique.values():
            key = (row_period(row), normalize(row[field]))
            bucket = actual.setdefault(key, [0] * len(metrics))
            for i, metric in enumerate(metrics):
                bucket[i] += validate_metric(row.get(metric), metric, "detalhamento")
        self._notify(None, "Conferindo os totais oficiais por mês e produto")
        control_body = {**body, "details": ["heading" if municipal else "ncm"]}
        payload = self.request_json(
            "POST", f"{endpoint}?language=pt", json_body=control_body
        )
        self.control_requests = getattr(self, "control_requests", 0) + 1
        data = payload.get("data")
        controls = data.get("list") if isinstance(data, dict) else None
        if not isinstance(controls, list):
            raise ComexStatError(
                "O Comex Stat não forneceu os totais de conferência. Tente novamente."
            )
        requested = {normalize(value) for value in body["filters"][0]["values"]}
        expected = {}
        for index, row in enumerate(controls):
            if index % 512 == 0:
                self._check_cancel()
            if not isinstance(row, dict):
                raise ComexStatError(
                    "O Comex Stat retornou um total de conferência inválido."
                )
            try:
                key = (row_period(row), normalize(row.get(field)))
            except (ValueError, TypeError, argparse.ArgumentTypeError) as exc:
                raise ComexStatError(
                    "Identificação inválida na conferência de totais."
                ) from exc
            if (
                not body["period"]["from"] <= key[0] <= body["period"]["to"]
                or key[1] not in requested
            ):
                raise ComexStatError(
                    "O total de conferência está fora do período ou produto solicitado."
                )
            if key in expected:
                raise ComexStatError(
                    "O Comex Stat repetiu um total de conferência. Tente novamente."
                )
            expected[key] = [
                validate_metric(row.get(m), m, "conferência oficial") for m in metrics
            ]
        for key in actual.keys() | expected.keys():
            if key not in actual or key not in expected or actual[key] != expected[key]:
                raise ComexStatError(
                    f"Os registros não conferem com os totais oficiais de {key[0]}, "
                    f"{'SH4' if municipal else 'NCM'} {key[1]}. "
                    "Nenhum resultado incompleto foi publicado. Tente novamente."
                )

    def _query_imports_live(
        self,
        ncms: Sequence[str],
        start_period: str,
        end_period: str,
        chunk_size: int,
    ) -> list[dict[str, Any]]:
        if isinstance(chunk_size, bool) or not isinstance(chunk_size, int):
            raise ValueError("O lote de NCMs deve ser um número inteiro.")
        if chunk_size < 1 or chunk_size > 200:
            raise ValueError("O lote de NCMs deve ficar entre 1 e 200.")
        if isinstance(ncms, (str, bytes)):
            raise ValueError("A lista de NCMs precisa ser uma sequência de códigos.")
        if not ncms:
            return []
        normalized_ncms: list[str] = []
        seen_ncms: set[str] = set()
        for value in ncms:
            normalized = normalize_ncm(value)
            if normalized is None:
                raise ValueError("A lista contém um NCM vazio.")
            if normalized not in seen_ncms:
                seen_ncms.add(normalized)
                normalized_ncms.append(normalized)
        start_period = validate_period(start_period)
        end_period = validate_period(end_period)
        if period_index(start_period) > period_index(end_period):
            raise ValueError(
                "O período inicial não pode ser posterior ao período final."
            )
        if period_index(start_period) < period_index(MIN_PERIOD):
            raise ValueError(f"Os dados gerais do Comex Stat começam em {MIN_PERIOD}.")

        rows: list[dict[str, Any]] = []
        batches = list(chunks(normalized_ncms, chunk_size))
        windows = (
            period_windows(start_period, end_period, self.window_months)
            if self.window_months is not None
            else [(start_period, end_period)]
        )
        tasks = [
            (batch, window_start, window_end)
            for window_start, window_end in windows
            for batch in batches
        ]
        for task_number, (batch, window_start, window_end) in enumerate(tasks, start=1):
            self._notify(
                None,
                f"Consultando dados gerais: lote {task_number}/{len(tasks)} · {window_start} a {window_end}",
            )
            print(
                f"Consultando lote {task_number}/{len(tasks)} "
                f"({len(batch)} NCMs; {window_start} a {window_end})...",
                flush=True,
            )
            body = {
                "flow": "import",
                "monthDetail": True,
                "period": {"from": window_start, "to": window_end},
                "filters": [{"filter": "ncm", "values": [int(ncm) for ncm in batch]}],
                "details": DEFAULT_DETAILS,
                "metrics": DEFAULT_METRICS,
            }
            payload = self.request_json(
                "POST",
                f"{API_GENERAL}?language=pt",
                json_body=body,
            )
            data = payload.get("data") or {}
            batch_rows = data.get("list") if isinstance(data, dict) else None
            if batch_rows is None:
                raise ComexStatError(
                    "Formato inesperado: a resposta não contém data.list."
                )
            if not isinstance(batch_rows, list):
                raise ComexStatError("Formato inesperado: data.list não é uma lista.")
            requested_ncms = set(batch)
            for row_number, row in enumerate(batch_rows, start=1):
                if not isinstance(row, dict):
                    raise ComexStatError(
                        "Formato inesperado: data.list contém um item que não é objeto "
                        f"(lote {task_number}, posição {row_number})."
                    )
                validate_api_row(
                    row,
                    requested_ncms=requested_ncms,
                    window_start=window_start,
                    window_end=window_end,
                    context=f"lote {task_number}, posição {row_number}",
                )
            if self.verify_official_totals:
                self._verify_totals(API_GENERAL, body, batch_rows, municipal=False)
            rows.extend(batch_rows)

            if self.cache is None:
                self._notify(
                    task_number / len(tasks),
                    f"Dados gerais: lote {task_number}/{len(tasks)} validado",
                )

        # Proteção contra repetição sem esconder divergências entre duas respostas.
        unique: dict[tuple[Any, ...], dict[str, Any]] = {}
        for row in rows:
            key = api_row_identity(row)
            previous = unique.get(key)
            if previous is not None:
                if api_metric_signature(previous) != api_metric_signature(row):
                    raise ComexStatError(
                        "A API retornou registros duplicados com valores divergentes "
                        f"para {key}. A geração foi interrompida para evitar dupla contagem."
                    )
                continue
            unique[key] = row
        return list(unique.values())

    def _query_municipal_imports_live(
        self,
        ncms: Sequence[str],
        start_period: str,
        end_period: str,
        chunk_size: int = 25,
    ) -> list[dict[str, Any]]:
        """Consulta o módulo municipal pelas posições SH4 do cadastro interno.

        O Comex Stat não publica NCM de oito dígitos, CIF, frete, seguro ou URF
        na base municipal. Por isso, esta consulta permanece separada dos Dados
        Gerais e solicita somente município fiscal, UF, SH4, país, FOB e peso.
        """
        if isinstance(chunk_size, bool) or not isinstance(chunk_size, int):
            raise ValueError("O lote de posições SH4 deve ser um número inteiro.")
        if chunk_size < 1 or chunk_size > 200:
            raise ValueError("O lote de posições SH4 deve ficar entre 1 e 200.")
        if isinstance(ncms, (str, bytes)):
            raise ValueError("A lista de NCMs precisa ser uma sequência de códigos.")
        if not ncms:
            return []

        headings = sorted({heading_from_ncm(value) for value in ncms})
        start_period = validate_period(start_period)
        end_period = validate_period(end_period)
        if period_index(start_period) > period_index(end_period):
            raise ValueError(
                "O período inicial não pode ser posterior ao período final."
            )
        if period_index(start_period) < period_index(MIN_PERIOD):
            raise ValueError(
                f"Os dados municipais do Comex Stat começam em {MIN_PERIOD}."
            )

        rows: list[dict[str, Any]] = []
        batches = list(chunks(headings, chunk_size))
        windows = (
            period_windows(start_period, end_period, self.window_months)
            if self.window_months is not None
            else [(start_period, end_period)]
        )
        tasks = [
            (batch, window_start, window_end)
            for window_start, window_end in windows
            for batch in batches
        ]
        for task_number, (batch, window_start, window_end) in enumerate(tasks, start=1):
            self._notify(
                None,
                f"Consultando municípios: lote {task_number}/{len(tasks)} · {window_start} a {window_end}",
            )
            print(
                f"Consultando municípios {task_number}/{len(tasks)} "
                f"({len(batch)} posições SH4; {window_start} a {window_end})...",
                flush=True,
            )
            body = {
                "flow": "import",
                "monthDetail": True,
                "period": {"from": window_start, "to": window_end},
                "filters": [{"filter": "heading", "values": list(batch)}],
                "details": MUNICIPAL_DETAILS,
                "metrics": MUNICIPAL_METRICS,
            }
            payload = self.request_json(
                "POST",
                f"{API_CITIES}?language=pt",
                json_body=body,
            )
            data = payload.get("data") or {}
            batch_rows = data.get("list") if isinstance(data, dict) else None
            if batch_rows is None:
                raise ComexStatError(
                    "Formato inesperado: a resposta municipal não contém data.list."
                )
            if not isinstance(batch_rows, list):
                raise ComexStatError(
                    "Formato inesperado: data.list da consulta municipal não é uma lista."
                )
            requested_headings = set(batch)
            for row_number, row in enumerate(batch_rows, start=1):
                if not isinstance(row, dict):
                    raise ComexStatError(
                        "Formato inesperado: a consulta municipal retornou um item "
                        f"que não é objeto (lote {task_number}, posição {row_number})."
                    )
                validate_municipal_api_row(
                    row,
                    requested_headings=requested_headings,
                    window_start=window_start,
                    window_end=window_end,
                    context=f"lote municipal {task_number}, posição {row_number}",
                )
            if self.verify_official_totals:
                self._verify_totals(API_CITIES, body, batch_rows, municipal=True)
            rows.extend(batch_rows)

            if self.cache is None:
                self._notify(
                    task_number / len(tasks),
                    f"Dados municipais: lote {task_number}/{len(tasks)} validado",
                )

        unique: dict[tuple[Any, ...], dict[str, Any]] = {}
        for row in rows:
            key = municipal_row_identity(row)
            previous = unique.get(key)
            if previous is not None:
                if municipal_metric_signature(previous) != municipal_metric_signature(
                    row
                ):
                    raise ComexStatError(
                        "A API municipal retornou registros duplicados com valores "
                        f"divergentes para {key}. A geração foi interrompida."
                    )
                continue
            unique[key] = row
        return list(unique.values())


def to_number(value: Any) -> int | float:
    if value in (None, ""):
        return 0
    if isinstance(value, bool):
        raise ComexStatError(f"Métrica numérica inválida recebida da API: {value!r}")
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        number = value
    else:
        text = str(value).strip().replace("\u00a0", "").replace(" ", "")
        if "," in text and "." in text:
            if text.rfind(",") > text.rfind("."):
                text = text.replace(".", "").replace(",", ".")
            else:
                text = text.replace(",", "")
        elif "," in text:
            text = text.replace(",", ".")
        try:
            if re.fullmatch(r"[+-]?[0-9]+", text):
                return int(text)
            number = float(text)
        except (TypeError, ValueError) as exc:
            raise ComexStatError(
                f"Métrica numérica inválida recebida da API: {value!r}"
            ) from exc
    if not math.isfinite(number):
        raise ComexStatError(f"Métrica numérica inválida recebida da API: {value!r}")
    try:
        integer = int(number)
    except (TypeError, ValueError, OverflowError) as exc:
        raise ComexStatError(
            f"Métrica numérica inválida recebida da API: {value!r}"
        ) from exc
    return integer if number.is_integer() else number


def validate_metric(
    value: Any,
    field: str,
    context: str,
) -> int:
    """Valida kg ou centavos exatos, sem arredondar ou preencher ausências."""
    if value is None or (isinstance(value, str) and not value.strip()):
        raise ComexStatError(
            f"A API retornou a métrica {field} sem valor em {context}."
        )
    try:
        return published_units(value, places=0 if field == "metricKG" else 2)
    except ValueError as exc:
        raise ComexStatError(f"Métrica {field} inválida em {context}: {exc}") from exc


def row_period(row: dict[str, Any]) -> str:
    year = str(row.get("year", ""))
    month = str(row.get("monthNumber", "")).zfill(2)
    return validate_period(f"{year}-{month}")


def validate_api_row(
    row: dict[str, Any],
    requested_ncms: set[str],
    window_start: str,
    window_end: str,
    context: str,
) -> None:
    """Confere identidade, período e métricas antes de aceitar uma linha da API."""
    # O contrato atual da API entrega os rótulos solicitados em ``details`` e o
    # código NCM, mas não garante códigos auxiliares para UF, URF e país.
    required_fields = {"coNcm", "ncm", "year", "monthNumber", *DEFAULT_DETAILS}
    missing_fields = sorted(field for field in required_fields if field not in row)
    if missing_fields:
        raise ComexStatError(
            "A API omitiu campos de identificação solicitados "
            f"({', '.join(missing_fields)}) em {context}."
        )
    try:
        ncm = normalize_ncm(row.get("coNcm"))
    except ValueError as exc:
        raise ComexStatError(f"NCM inválido recebido da API ({context}).") from exc
    if ncm is None:
        raise ComexStatError(f"Linha da API sem Código NCM ({context}).")
    if ncm not in requested_ncms:
        raise ComexStatError(
            f"A API retornou o NCM {ncm}, que não fazia parte da consulta ({context})."
        )
    try:
        period = row_period(row)
    except (argparse.ArgumentTypeError, ValueError) as exc:
        raise ComexStatError(f"Período inválido recebido da API ({context}).") from exc
    if not (
        period_index(window_start) <= period_index(period) <= period_index(window_end)
    ):
        raise ComexStatError(
            f"A API retornou o período {period} fora da janela "
            f"{window_start} a {window_end} ({context})."
        )
    missing_metrics = [field for field in DEFAULT_METRICS if field not in row]
    if missing_metrics:
        raise ComexStatError(
            f"A API omitiu métricas solicitadas ({', '.join(missing_metrics)}) em {context}."
        )
    metrics = {
        field: validate_metric(row[field], field, context) for field in DEFAULT_METRICS
    }
    expected_cif = (
        metrics["metricFOB"] + metrics["metricFreight"] + metrics["metricInsurance"]
    )
    if metrics["metricCIF"] != expected_cif:
        raise ComexStatError(
            f"A API retornou um CIF divergente de FOB + frete + seguro em {context}."
        )
    for field in ("ncm", "state", "urf", "country"):
        if not isinstance(row.get(field), str):
            raise ComexStatError(
                f"A API retornou o campo {field} com formato inválido em {context}."
            )
        if not clean_text(row.get(field)):
            raise ComexStatError(f"A API retornou o campo {field} vazio em {context}.")


def api_row_identity(row: dict[str, Any]) -> tuple[Any, ...]:
    """Identifica uma agregação mensal por NCM, UF, URF e país."""
    return (
        normalize_ncm(row.get("coNcm")),
        row_period(row),
        clean_text(row.get("state")),
        clean_text(row.get("urf")),
        clean_text(row.get("country")),
    )


def api_metric_signature(row: dict[str, Any]) -> tuple[int | float, ...]:
    """Produz uma assinatura numérica estável para detectar duplicidades conflitantes."""
    return tuple(
        validate_metric(row.get(field), field, "duplicidade")
        for field in DEFAULT_METRICS
    )


def municipal_name(row: dict[str, Any]) -> str:
    """Retorna o município fiscal conforme identificado pela base municipal."""
    return clean_text(row.get("noMunMinsgUf"))


def validate_municipal_api_row(
    row: dict[str, Any],
    requested_headings: set[str],
    window_start: str,
    window_end: str,
    context: str,
) -> None:
    """Valida uma linha do endpoint municipal sem misturá-la aos Dados Gerais."""
    required_fields = {
        "noMunMinsgUf",
        "year",
        "monthNumber",
        "state",
        "headingCode",
        "heading",
        "country",
    }
    missing_fields = sorted(field for field in required_fields if field not in row)
    if missing_fields:
        raise ComexStatError(
            "A API municipal omitiu campos de identificação solicitados "
            f"({', '.join(missing_fields)}) em {context}."
        )
    if not municipal_name(row):
        raise ComexStatError(
            f"Linha da API municipal sem município fiscal ({context})."
        )
    try:
        heading = normalize_heading(row.get("headingCode"))
    except ValueError as exc:
        raise ComexStatError(
            f"Código SH4 inválido recebido da API ({context})."
        ) from exc
    if heading not in requested_headings:
        raise ComexStatError(
            f"A API municipal retornou o SH4 {heading}, que não fazia parte "
            f"da consulta ({context})."
        )
    try:
        period = row_period(row)
    except (argparse.ArgumentTypeError, ValueError) as exc:
        raise ComexStatError(
            f"Período inválido recebido da API municipal ({context})."
        ) from exc
    if not (
        period_index(window_start) <= period_index(period) <= period_index(window_end)
    ):
        raise ComexStatError(
            f"A API municipal retornou o período {period} fora da janela "
            f"{window_start} a {window_end} ({context})."
        )
    missing_metrics = [field for field in MUNICIPAL_METRICS if field not in row]
    if missing_metrics:
        raise ComexStatError(
            "A API municipal omitiu métricas solicitadas "
            f"({', '.join(missing_metrics)}) em {context}."
        )
    for field in MUNICIPAL_METRICS:
        validate_metric(row[field], field, context)
    for field in ("noMunMinsgUf", "state", "heading", "country"):
        if not isinstance(row.get(field), str):
            raise ComexStatError(
                f"A API municipal retornou o campo {field} com formato inválido em {context}."
            )
        if not clean_text(row.get(field)):
            raise ComexStatError(
                f"A API municipal retornou o campo {field} vazio em {context}."
            )


def municipal_row_identity(row: dict[str, Any]) -> tuple[Any, ...]:
    """Identifica uma agregação municipal mensal por SH4 e país."""
    return (
        municipal_name(row),
        row_period(row),
        clean_text(row.get("state")),
        normalize_heading(row.get("headingCode")),
        clean_text(row.get("country")),
    )


def municipal_metric_signature(row: dict[str, Any]) -> tuple[int | float, ...]:
    return tuple(
        validate_metric(row.get(field), field, "duplicidade municipal")
        for field in MUNICIPAL_METRICS
    )


def month_label(period: str, include_year: bool) -> str:
    year, month_text = period.split("-")
    month = int(month_text)
    if include_year:
        return f"{year}-{month:02d}. {MONTH_NAMES[month]}"
    return f"{month:02d}. {MONTH_NAMES[month]}"


def excel_metric(row, field):
    units = validate_metric(row.get(field), field, "relatório Excel")
    if field == "metricKG":
        return units
    return units // 100 if units % 100 == 0 else from_units(units)


def raw_row(
    row: dict[str, Any],
    include_year: bool,
    products_by_ncm: dict[str, list[str]],
) -> list[Any]:
    period = row_period(row)
    ncm = normalize_ncm(row.get("coNcm"))
    if ncm is None:
        raise ComexStatError(f"Linha da API sem coNcm: {row!r}")
    cif = excel_metric(row, "metricCIF")
    kg = excel_metric(row, "metricKG")
    cif_per_kg = float(cif) / float(kg) if kg else None
    return [
        month_label(period, include_year),
        ncm,
        clean_text(row.get("ncm")),
        state_abbreviation(row.get("state")),
        clean_text(row.get("urf")),
        clean_text(row.get("country")),
        excel_metric(row, "metricFOB"),
        excel_metric(row, "metricFreight"),
        excel_metric(row, "metricInsurance"),
        cif,
        kg,
        clean_text(", ".join(products_by_ncm.get(ncm, []))),
        cif_per_kg,
    ]


def municipal_excel_row(
    row: dict[str, Any],
    include_year: bool,
    products_by_heading: dict[str, list[str]],
    ncms_by_heading: dict[str, list[str]],
) -> list[Any]:
    """Converte uma linha da base municipal mantendo sua metodologia explícita."""
    period = row_period(row)
    heading = normalize_heading(row.get("headingCode"))
    fob = excel_metric(row, "metricFOB")
    kg = excel_metric(row, "metricKG")
    return [
        month_label(period, include_year),
        municipal_name(row),
        state_abbreviation(row.get("state")),
        heading,
        clean_text(row.get("heading")),
        clean_text(row.get("country")),
        fob,
        kg,
        clean_text(", ".join(products_by_heading.get(heading, []))),
        clean_text(", ".join(ncms_by_heading.get(heading, []))),
        float(fob) / float(kg) if kg else None,
    ]


def style_header(sheet: Any, row: int = 1) -> None:
    for cell in sheet[row]:
        if cell.value is None:
            continue
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = CELL_GRID
    sheet.row_dimensions[row].height = 28


def apply_filter_and_banding(sheet: Any) -> None:
    """Aplica filtros e faixas alternadas sem objetos Table do Excel.

    O uso apenas do AutoFilter da própria planilha evita reparos de arquivos em
    algumas instalações do Excel que rejeitam tabelas estruturadas externas.
    """
    max_row, max_column = sheet.max_row, sheet.max_column
    if max_row < 1 or max_column < 1:
        return
    sheet.auto_filter.ref = sheet.dimensions
    # sheet[row_number] recalcula max_column percorrendo TODAS as células.
    # Em uma base longa isso transforma a formatação em trabalho quadrático.
    for row_number, row in enumerate(
        sheet.iter_rows(
            min_row=2,
            max_row=max_row,
            min_col=1,
            max_col=max_column,
        ),
        start=2,
    ):
        if row_number % 2 == 0:
            for cell in row:
                cell.fill = ALTERNATE_ROW_FILL


def apply_grid(sheet: Any) -> None:
    """Aplica grade e centraliza verticalmente toda a área utilizada."""
    if sheet.max_row < 1 or sheet.max_column < 1:
        return
    centered: dict[tuple[Any, ...], Alignment] = {}
    for row in sheet.iter_rows(
        min_row=1,
        max_row=sheet.max_row,
        min_col=1,
        max_col=sheet.max_column,
    ):
        for cell in row:
            cell.border = CELL_GRID
            original = cell.alignment
            if original.vertical != "center":
                key = (
                    original.horizontal,
                    original.vertical,
                    original.textRotation,
                    original.wrapText,
                    original.shrinkToFit,
                    original.indent,
                    original.relativeIndent,
                    original.justifyLastLine,
                    original.readingOrder,
                )
                alignment = centered.get(key)
                if alignment is None:
                    alignment = copy(original)
                    alignment.vertical = "center"
                    centered[key] = alignment
                cell.alignment = alignment


def force_text_columns(sheet: Any, columns: Sequence[int]) -> None:
    """Impede que códigos ou descrições iniciados por '=' virem fórmulas."""
    if sheet.max_row < 2:
        return
    for column in columns:
        for row in sheet.iter_rows(
            min_row=2,
            max_row=sheet.max_row,
            min_col=column,
            max_col=column,
        ):
            cell = row[0]
            if cell.value is not None:
                cell.value = str(cell.value)
                cell.data_type = "s"


def format_columns(sheet: Any, formats: dict[int, str]) -> None:
    """Aplica máscaras a todas as células preenchidas das colunas indicadas."""
    if sheet.max_row < 2:
        return
    for column, number_format in formats.items():
        for row in sheet.iter_rows(
            min_row=2,
            max_row=sheet.max_row,
            min_col=column,
            max_col=column,
        ):
            cell = row[0]
            if cell.value is not None:
                cell.number_format = (
                    CURRENCY_FORMAT + ".00"
                    if number_format == CURRENCY_FORMAT
                    and published_units(cell.value, places=2) % 100
                    else number_format
                )


def configure_printing(sheet: Any, landscape: bool = True) -> None:
    """Configura impressão legível sem interferir nos filtros do Excel."""
    sheet.print_title_rows = "1:1"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.orientation = "landscape" if landscape else "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.sheet_view.zoomScale = 90


def fit_columns(sheet: Any, maximum: int = 48) -> None:
    for column_cells in sheet.iter_cols(
        min_row=1,
        max_row=min(sheet.max_row, 4000),
        min_col=1,
        max_col=sheet.max_column,
    ):
        length = max(
            (len(str(cell.value)) for cell in column_cells if cell.value is not None),
            default=0,
        )
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(
            maximum, max(10, length + 2)
        )


def prepare_data_sheet(sheet: Any, rows: Sequence[list[Any]]) -> None:
    sheet.append(RAW_HEADERS)
    for row in rows:
        sheet.append(row)
    style_header(sheet)
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    force_text_columns(sheet, (1, 2, 3, 4, 5, 6, 12))
    format_columns(
        sheet,
        {
            1: "@",
            2: "@",
            3: "@",
            4: "@",
            5: "@",
            6: "@",
            7: CURRENCY_FORMAT,
            8: CURRENCY_FORMAT,
            9: CURRENCY_FORMAT,
            10: CURRENCY_FORMAT,
            11: WEIGHT_FORMAT,
            12: "@",
            13: UNIT_VALUE_FORMAT,
        },
    )
    apply_filter_and_banding(sheet)
    apply_grid(sheet)
    configure_printing(sheet)
    fit_columns(sheet)
    sheet.column_dimensions["C"].width = 46
    sheet.column_dimensions["D"].width = 12
    sheet.column_dimensions["E"].width = 36
    sheet.column_dimensions["F"].width = 22
    sheet.column_dimensions["L"].width = 34
    for column in ("G", "H", "I", "J", "K"):
        sheet.column_dimensions[column].width = max(
            sheet.column_dimensions[column].width or 0,
            18,
        )
    sheet.column_dimensions["K"].width = 18
    sheet.column_dimensions["M"].width = 15
    for row_number in range(2, sheet.max_row + 1):
        description = str(sheet.cell(row_number, 3).value or "")
        products = str(sheet.cell(row_number, 12).value or "")
        for column in (3, 12):
            sheet.cell(row_number, column).alignment = Alignment(
                wrap_text=True,
                horizontal="left",
                vertical="center",
            )
        wrapped_lines = max(
            1,
            (len(description) + 45) // 46,
            (len(products) + 33) // 34,
        )
        if wrapped_lines > 1:
            sheet.row_dimensions[row_number].height = min(60, 15 * wrapped_lines)


def create_data_sheets(
    workbook: Any,
    base_title: str,
    rows: Sequence[list[Any]],
    tab_color: str = DATA_TAB_COLOR,
) -> None:
    """Cria uma ou mais abas sem ultrapassar o limite de linhas do Excel."""
    starts = range(0, max(1, len(rows)), EXCEL_MAX_DATA_ROWS)
    for part, start in enumerate(starts, start=1):
        suffix = "" if part == 1 else f" {part}"
        title = f"{base_title[: 31 - len(suffix)]}{suffix}"
        sheet = workbook.create_sheet(title)
        sheet.sheet_properties.tabColor = tab_color
        prepare_data_sheet(
            sheet,
            rows[start : start + EXCEL_MAX_DATA_ROWS],
        )


def prepare_municipal_data_sheet(sheet: Any, rows: Sequence[list[Any]]) -> None:
    sheet.append(MUNICIPAL_HEADERS)
    for row in rows:
        sheet.append(row)
    style_header(sheet)
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    text_columns = (1, 2, 3, 4, 5, 6, 9, 10)
    force_text_columns(sheet, text_columns)
    format_columns(
        sheet,
        {
            1: "@",
            2: "@",
            3: "@",
            4: "@",
            5: "@",
            6: "@",
            7: CURRENCY_FORMAT,
            8: WEIGHT_FORMAT,
            9: "@",
            10: "@",
            11: UNIT_VALUE_FORMAT,
        },
    )
    apply_filter_and_banding(sheet)
    apply_grid(sheet)
    configure_printing(sheet)
    fit_columns(sheet)
    sheet.column_dimensions["B"].width = 30
    sheet.column_dimensions["C"].width = 10
    sheet.column_dimensions["D"].width = 10
    sheet.column_dimensions["E"].width = 64
    sheet.column_dimensions["F"].width = 22
    sheet.column_dimensions["G"].width = 18
    sheet.column_dimensions["H"].width = 18
    sheet.column_dimensions["I"].width = 36
    sheet.column_dimensions["J"].width = 30
    sheet.column_dimensions["K"].width = 15
    simple_columns = tuple(
        column
        for column in range(1, len(MUNICIPAL_HEADERS) + 1)
        if column not in {5, 9, 10}
    )
    simple_alignment = Alignment(vertical="center")
    for row_number in range(2, sheet.max_row + 1):
        description = str(sheet.cell(row_number, 5).value or "")
        products = str(sheet.cell(row_number, 9).value or "")
        ncms = str(sheet.cell(row_number, 10).value or "")
        for column in (5, 9, 10):
            sheet.cell(row_number, column).alignment = Alignment(
                wrap_text=True,
                horizontal="left",
                vertical="center",
            )
        for column in simple_columns:
            sheet.cell(row_number, column).alignment = simple_alignment
        wrapped_lines = max(
            1,
            (len(description) + 63) // 64,
            (len(products) + 35) // 36,
            (len(ncms) + 29) // 30,
        )
        if wrapped_lines > 1:
            sheet.row_dimensions[row_number].height = min(48, 13.5 * wrapped_lines)


def create_municipal_data_sheets(
    workbook: Any,
    rows: Sequence[list[Any]],
) -> None:
    """Cria a base municipal em partes, se necessário, sem misturar métricas."""
    starts = range(0, max(1, len(rows)), EXCEL_MAX_DATA_ROWS)
    for part, start in enumerate(starts, start=1):
        suffix = "" if part == 1 else f" {part}"
        title = f"{SHEET_MUNICIPAL_DATA[: 31 - len(suffix)]}{suffix}"
        sheet = workbook.create_sheet(title)
        sheet.sheet_properties.tabColor = DATA_TAB_COLOR
        prepare_municipal_data_sheet(
            sheet,
            rows[start : start + EXCEL_MAX_DATA_ROWS],
        )


def build_municipal_summary(
    sheet: Any,
    municipal_rows: Sequence[dict[str, Any]],
) -> None:
    headers = [
        "Município fiscal",
        "UF",
        "SH4",
        "Países",
        "Registros",
        "FOB",
        "Peso líq.",
        "FOB/kg",
        "% FOB",
    ]
    grouped: dict[tuple[str, str], dict[str, Any]] = {}
    for row in municipal_rows:
        key = (municipal_name(row), state_abbreviation(row.get("state")))
        current = grouped.setdefault(
            key,
            {"headings": set(), "countries": set(), "records": 0, "fob": 0, "kg": 0},
        )
        current["headings"].add(normalize_heading(row.get("headingCode")))
        country = clean_text(row.get("country"))
        if country:
            current["countries"].add(country)
        current["records"] += 1
        current["fob"] += validate_metric(
            row.get("metricFOB"), "metricFOB", "resumo municipal"
        )
        current["kg"] += validate_metric(
            row.get("metricKG"), "metricKG", "resumo municipal"
        )

    total_fob = sum(item["fob"] for item in grouped.values())
    sheet.append(headers)
    for (municipality, state), item in sorted(
        grouped.items(),
        key=lambda entry: (-entry[1]["fob"], entry[0][0].casefold()),
    ):
        sheet.append(
            [
                municipality,
                state,
                ", ".join(sorted(item["headings"])),
                len(item["countries"]),
                item["records"],
                from_units(item["fob"]),
                item["kg"],
                item["fob"] / (100 * item["kg"]) if item["kg"] else None,
                item["fob"] / total_fob if total_fob else 0,
            ]
        )
    style_header(sheet)
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    force_text_columns(sheet, (1, 2, 3))
    format_columns(
        sheet,
        {
            1: "@",
            2: "@",
            3: "@",
            4: ABSOLUTE_NUMBER_FORMAT,
            5: ABSOLUTE_NUMBER_FORMAT,
            6: CURRENCY_FORMAT,
            7: WEIGHT_FORMAT,
            8: UNIT_VALUE_FORMAT,
            9: PERCENT_FORMAT,
        },
    )
    apply_filter_and_banding(sheet)
    apply_grid(sheet)
    configure_printing(sheet)
    fit_columns(sheet)
    sheet.column_dimensions["A"].width = 30
    sheet.column_dimensions["B"].width = 10
    sheet.column_dimensions["C"].width = 12
    sheet.column_dimensions["D"].width = 10
    sheet.column_dimensions["E"].width = 12
    sheet.column_dimensions["F"].width = 18
    sheet.column_dimensions["G"].width = 18
    sheet.column_dimensions["H"].width = 15
    sheet.column_dimensions["I"].width = 10


def build_summary(
    sheet: Any,
    api_rows: Sequence[dict[str, Any]],
    products_by_ncm: dict[str, list[str]],
) -> None:
    headers = [
        "NCM",
        "Descrição",
        "Produtos",
        "Registros",
        "FOB",
        "Frete",
        "Seguro",
        "CIF",
        "Peso líq.",
        "CIF/kg",
        "% CIF",
    ]
    grouped: dict[str, dict[str, Any]] = {}
    for row in api_rows:
        ncm = normalize_ncm(row.get("coNcm"))
        if ncm is None:
            continue
        current = grouped.setdefault(
            ncm,
            {
                "description": clean_text(row.get("ncm", "")),
                "description_period": row_period(row),
                "records": 0,
                "fob": 0,
                "freight": 0,
                "insurance": 0,
                "cif": 0,
                "kg": 0,
            },
        )
        current_period = row_period(row)
        description = clean_text(row.get("ncm", ""))
        if description and current_period >= current["description_period"]:
            current["description"] = description
            current["description_period"] = current_period
        current["records"] += 1
        for key, field in (
            ("fob", "metricFOB"),
            ("freight", "metricFreight"),
            ("insurance", "metricInsurance"),
            ("cif", "metricCIF"),
            ("kg", "metricKG"),
        ):
            current[key] += validate_metric(row.get(field), field, "resumo NCM")

    # NCMs sem movimentação também aparecem no resumo, deixando a ausência explícita.
    for ncm in products_by_ncm:
        grouped.setdefault(
            ncm,
            {
                "description": "Não retornada — sem movimentação no período",
                "description_period": "",
                "records": 0,
                "fob": 0,
                "freight": 0,
                "insurance": 0,
                "cif": 0,
                "kg": 0,
            },
        )

    total_cif = sum(item["cif"] for item in grouped.values())
    sheet.append(headers)
    ordered = sorted(grouped.items(), key=lambda item: (-item[1]["cif"], item[0]))
    for ncm, item in ordered:
        cif_per_kg = item["cif"] / (100 * item["kg"]) if item["kg"] else None
        share = item["cif"] / total_cif if total_cif else 0
        sheet.append(
            [
                ncm,
                item["description"],
                clean_text(", ".join(products_by_ncm.get(ncm, []))),
                item["records"],
                from_units(item["fob"]),
                from_units(item["freight"]),
                from_units(item["insurance"]),
                from_units(item["cif"]),
                item["kg"],
                cif_per_kg,
                share,
            ]
        )

    style_header(sheet)
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    force_text_columns(sheet, (1, 2, 3))
    format_columns(
        sheet,
        {
            1: "@",
            2: "@",
            3: "@",
            4: ABSOLUTE_NUMBER_FORMAT,
            5: CURRENCY_FORMAT,
            6: CURRENCY_FORMAT,
            7: CURRENCY_FORMAT,
            8: CURRENCY_FORMAT,
            9: WEIGHT_FORMAT,
            10: UNIT_VALUE_FORMAT,
            11: PERCENT_FORMAT,
        },
    )
    for row_number in range(2, sheet.max_row + 1):
        description = str(sheet.cell(row_number, 2).value or "")
        products = str(sheet.cell(row_number, 3).value or "")
        wrapped_lines = max(
            1,
            (len(description) + 45) // 46,
            (len(products) + 33) // 34,
        )
        if wrapped_lines > 1:
            sheet.row_dimensions[row_number].height = min(60, 15 * wrapped_lines)
        for column in (2, 3):
            sheet.cell(row_number, column).alignment = Alignment(
                wrap_text=True,
                horizontal="left",
                vertical="center",
            )
    apply_filter_and_banding(sheet)
    apply_grid(sheet)
    configure_printing(sheet)
    fit_columns(sheet)
    sheet.column_dimensions["B"].width = 46
    sheet.column_dimensions["C"].width = 34
    sheet.column_dimensions["D"].width = 12
    for column in ("E", "F", "G", "H", "I"):
        sheet.column_dimensions[column].width = max(
            sheet.column_dimensions[column].width or 0,
            18,
        )
    sheet.column_dimensions["J"].width = 15
    sheet.column_dimensions["K"].width = 10


def build_mapping(sheet: Any, pairs: Sequence[ProductNcm]) -> None:
    sheet.append(["Produto", "NCM"])
    for pair in sorted(pairs, key=lambda item: (item.ncm, item.product.casefold())):
        sheet.append([clean_text(pair.product), pair.ncm])
    style_header(sheet)
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    force_text_columns(sheet, (1, 2))
    for row in sheet.iter_rows(min_row=2, max_col=2):
        for cell in row:
            cell.number_format = "@"
            cell.alignment = Alignment(horizontal="left", vertical="center")
    apply_filter_and_banding(sheet)
    apply_grid(sheet)
    configure_printing(sheet, landscape=False)
    fit_columns(sheet)


def _check_saved_row(
    cells: Sequence[Any], expected: Sequence[Any], context: str
) -> None:
    """Confere identidades e métricas, inclusive divergências que preservam totais."""
    if len(cells) != len(expected):
        raise ValueError(f"Quantidade de colunas divergente em {context}.")
    for column, (cell, source) in enumerate(zip(cells, expected, strict=True), start=1):
        value = cell.value
        if cell.data_type == "f":
            raise ValueError(f"Fórmula inesperada em {context}, coluna {column}.")
        if isinstance(source, (int, float, Decimal)) and not isinstance(source, bool):
            if isinstance(value, bool) or not isinstance(value, (int, float)):
                raise ValueError(
                    f"Valor numérico ausente/inválido em {context}, coluna {column}."
                )
            same = (
                Decimal(str(value)) == source
                if isinstance(source, (int, Decimal))
                else math.isclose(
                    value,
                    source,
                    rel_tol=1e-13,
                    abs_tol=1e-9,
                )
            )
        else:
            # Excel grava uma string vazia como célula sem conteúdo.
            same = ("" if value is None else value) == (
                "" if source is None else source
            )
        if not same:
            raise ValueError(f"Registro divergente em {context}, coluna {column}.")


def _verify_saved_data(
    workbook: Any,
    sheet_names: Sequence[str],
    headers: Sequence[str],
    expected_rows: int,
    metric_columns: Sequence[int],
    expected_totals: Sequence[float],
    expected_data: Sequence[Sequence[Any]] | None = None,
) -> None:
    """Lê cada aba uma vez, somando centavos e kg inteiros, sem tolerância."""
    source = iter(expected_data) if expected_data is not None else None
    count = 0
    totals = [0] * len(metric_columns)
    for name in sheet_names:
        rows = workbook[name].iter_rows(min_row=1, min_col=1, max_col=len(headers))
        header = next(rows, ())
        if tuple(cell.value for cell in header) != tuple(headers):
            raise ValueError(f"Cabeçalhos divergentes na aba {name}.")
        for row in rows:
            count += 1
            context = f"{name}, linha {count + 1}"
            if source is not None:
                expected = next(source, None)
                if expected is None:
                    raise ValueError(f"Linhas adicionais na aba {name}.")
                _check_saved_row(row, expected, context)
            for index, column in enumerate(metric_columns):
                cell = row[column - 1]
                value = cell.value
                if (
                    isinstance(value, bool)
                    or not isinstance(value, (int, float))
                    or not math.isfinite(value)
                ):
                    raise ValueError(f"Métrica inválida em {context}, coluna {column}.")
                totals[index] += published_units(
                    value, places=0 if index == len(metric_columns) - 1 else 2
                )
    if count != expected_rows or (
        source is not None and next(source, None) is not None
    ):
        raise ValueError(
            f"Quantidade de linhas divergente: esperadas {expected_rows}, encontradas {count}."
        )
    for index, (column, actual, expected) in enumerate(
        zip(metric_columns, totals, expected_totals, strict=True)
    ):
        places = 0 if index == len(metric_columns) - 1 else 2
        if actual != published_units(expected, places=places):
            raise ValueError(
                f"Total divergente na coluna {column}. O relatório não foi publicado."
            )


def _verify_grouped_summary(
    summary,
    source,
    key_columns,
    source_keys,
    metric_columns,
    source_metrics,
    records_column,
):
    """A grand total alone cannot detect a value moved to another NCM/city."""
    expected = {}
    for row in source:
        key = tuple(row[i] for i in source_keys)
        bucket = expected.setdefault(key, [0] * (len(source_metrics) + 1))
        bucket[-1] += 1
        for i, column in enumerate(source_metrics):
            bucket[i] += published_units(
                row[column], places=0 if i == len(source_metrics) - 1 else 2
            )
    seen = set()
    zero = [0] * (len(source_metrics) + 1)
    for row in summary:
        key = tuple(row[i] for i in key_columns)
        if key in seen:
            raise ValueError("Identificação repetida no resumo do Excel.")
        seen.add(key)
        actual = [
            published_units(row[c], places=0 if i == len(metric_columns) - 1 else 2)
            for i, c in enumerate(metric_columns)
        ]
        actual.append(published_units(row[records_column]))
        if actual != expected.get(key, zero):
            raise ValueError(
                f"O resumo do Excel divergiu da base para {' / '.join(map(str, key))}."
            )
    if expected.keys() - seen:
        raise ValueError("O resumo do Excel omitiu registros presentes na base.")


def validate_generated_workbook(
    workbook_path: Path,
    pairs: Sequence[ProductNcm],
    expected_rows: int,
    selected_periods: Sequence[str],
    include_year: bool,
    monthly_sheets: bool,
    expected_totals: Sequence[float],
    expected_municipal_rows: int = 0,
    expected_municipal_totals: Sequence[float] = (0.0, 0.0),
    *,
    expected_general_data: Sequence[Sequence[Any]] | None = None,
    expected_municipal_data: Sequence[Sequence[Any]] | None = None,
) -> None:
    """Confere bases, resumos, cadastro e abas mensais antes da publicação."""
    try:
        workbook = load_workbook(
            workbook_path, read_only=True, data_only=False, keep_links=False
        )
    except (OSError, zipfile.BadZipFile, InvalidFileException) as exc:
        raise ValueError(
            "O Excel gerado não passou na verificação de integridade."
        ) from exc
    try:
        names = workbook.sheetnames
        required = {SHEET_SUMMARY_NCM, SHEET_SUMMARY_MUNICIPAL, SHEET_PRODUCTS}
        missing = sorted(required.difference(names))
        if missing:
            raise ValueError(
                "O Excel gerado não contém as abas obrigatórias: " + ", ".join(missing)
            )
        if {"Parametros", "Parâmetros"}.intersection(names):
            raise ValueError(
                "O Excel gerado contém uma aba de parâmetros não prevista."
            )
        general = [
            name
            for name in names
            if name == SHEET_GENERAL_DATA or name.startswith(SHEET_GENERAL_DATA + " ")
        ]
        municipal = [
            name
            for name in names
            if name == SHEET_MUNICIPAL_DATA
            or name.startswith(SHEET_MUNICIPAL_DATA + " ")
        ]
        if not general or not municipal:
            raise ValueError("O Excel gerado deve conter Base geral e Base municipal.")
        _verify_saved_data(
            workbook,
            general,
            RAW_HEADERS,
            expected_rows,
            (7, 8, 9, 10, 11),
            expected_totals,
            expected_general_data,
        )
        _verify_saved_data(
            workbook,
            municipal,
            MUNICIPAL_HEADERS,
            expected_municipal_rows,
            (7, 8),
            expected_municipal_totals,
            expected_municipal_data,
        )

        expected_mapping = [
            (clean_text(pair.product), pair.ncm)
            for pair in sorted(
                pairs, key=lambda item: (item.ncm, item.product.casefold())
            )
        ]
        mapping = workbook[SHEET_PRODUCTS].iter_rows(
            min_row=1, max_col=2, values_only=True
        )
        if next(mapping, ()) != ("Produto", "NCM") or list(mapping) != expected_mapping:
            raise ValueError(
                f"A aba {SHEET_PRODUCTS} divergiu do cadastro interno de produtos e NCMs."
            )
        summary = workbook[SHEET_SUMMARY_NCM].iter_rows(
            min_row=2, max_col=11, values_only=True
        )
        summary_rows = list(summary)
        if len(summary_rows) != len({pair.ncm for pair in pairs}) or {
            row[0] for row in summary_rows
        } != {pair.ncm for pair in pairs}:
            raise ValueError(
                f"A aba {SHEET_SUMMARY_NCM} não contém todos os NCMs consultados."
            )
        for offset, expected in enumerate(expected_totals, start=4):
            places = 0 if offset == 8 else 2
            total = sum(
                published_units(row[offset], places=places) for row in summary_rows
            )
            if total != published_units(expected, places=places):
                raise ValueError(
                    f"Total divergente no resumo NCM, coluna {offset + 1}."
                )
        municipal_summary = list(
            workbook[SHEET_SUMMARY_MUNICIPAL].iter_rows(
                min_row=2, max_col=9, values_only=True
            )
        )
        for offset, expected in zip((5, 6), expected_municipal_totals, strict=True):
            places = 0 if offset == 6 else 2
            total = sum(
                published_units(row[offset], places=places) for row in municipal_summary
            )
            if total != published_units(expected, places=places):
                raise ValueError(
                    f"Total divergente no resumo municipal, coluna {offset + 1}."
                )

        if expected_general_data is not None:
            _verify_grouped_summary(
                summary_rows,
                expected_general_data,
                (0,),
                (1,),
                (4, 5, 6, 7, 8),
                (6, 7, 8, 9, 10),
                3,
            )
        if expected_municipal_data is not None:
            _verify_grouped_summary(
                municipal_summary,
                expected_municipal_data,
                (0, 1),
                (1, 2),
                (5, 6),
                (6, 7),
                4,
            )

        if monthly_sheets:
            by_month: dict[str, list[Sequence[Any]]] = defaultdict(list)
            if expected_general_data is not None:
                for row in expected_general_data:
                    by_month[str(row[0])].append(row)
            for period in selected_periods:
                base = month_label(period, include_year)
                monthly = [
                    name
                    for name in names
                    if name == base or name.startswith(base + " ")
                ]
                if not monthly:
                    raise ValueError(f"A aba mensal de {period} não foi gerada.")
                if expected_general_data is not None:
                    expected = by_month[base]
                    totals = [
                        sum_published(
                            (row[column - 1] for row in expected),
                            places=0 if column == 11 else 2,
                        )
                        for column in (7, 8, 9, 10, 11)
                    ]
                    _verify_saved_data(
                        workbook,
                        monthly,
                        RAW_HEADERS,
                        len(expected),
                        (7, 8, 9, 10, 11),
                        totals,
                        expected,
                    )
    finally:
        workbook.close()


def save_workbook_atomic(
    workbook: Any,
    output_path: Path,
    validator: Callable[[Path], None] | None = None,
) -> None:
    """Salva em arquivo temporário e substitui o destino somente após sucesso."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{output_path.stem}_",
        suffix=".xlsx",
        dir=output_path.parent,
    )
    os.close(descriptor)
    temporary_path = Path(temporary_name)
    try:
        workbook.save(temporary_path)
        if validator is not None:
            validator(temporary_path)
        os.replace(temporary_path, output_path)
    except PermissionError as exc:
        raise ValueError(
            f"Não foi possível salvar {output_path.name}. Feche o arquivo no Excel "
            "e tente novamente."
        ) from exc
    finally:
        temporary_path.unlink(missing_ok=True)


def create_report(
    output_path: Path,
    pairs: Sequence[ProductNcm],
    api_rows: Sequence[dict[str, Any]],
    start_period: str,
    end_period: str,
    monthly_sheets: bool,
    municipal_rows: Sequence[dict[str, Any]] = (),
    *,
    progress_callback: Callable[[float, str], None] | None = None,
) -> None:
    def report(percent: float, message: str) -> None:
        if progress_callback is not None:
            progress_callback(percent, message)

    report(0, "Conferindo os registros antes de montar o Excel")
    if not pairs:
        raise ValueError("A geração do relatório exige ao menos um produto com NCM.")
    # Valida o intervalo antes de percorrer as linhas. Assim, chamadas diretas
    # recebem uma mensagem determinística e nunca falham dentro da validação de
    # um registro por causa de um período de entrada malformado.
    try:
        start_period = validate_period(start_period)
        end_period = validate_period(end_period)
    except argparse.ArgumentTypeError as exc:
        raise ValueError(str(exc)) from exc
    if period_index(start_period) > period_index(end_period):
        raise ValueError("O período inicial não pode ser posterior ao período final.")
    if start_period < MIN_PERIOD:
        raise ValueError(f"Os Dados Gerais do Comex Stat começam em {MIN_PERIOD}.")
    allowed_ncms = {pair.ncm for pair in pairs}
    general_keys = set()
    for row_number, row in enumerate(api_rows, start=1):
        if not isinstance(row, dict):
            raise ValueError(
                f"A linha {row_number} dos dados não é um registro válido."
            )
        validate_api_row(
            row,
            requested_ncms=allowed_ncms,
            window_start=start_period,
            window_end=end_period,
            context=f"dados do relatório, linha {row_number}",
        )
        key = api_row_identity(row)
        if key in general_keys:
            raise ValueError(
                "Registro estatístico repetido na base geral. O relatório não foi publicado."
            )
        general_keys.add(key)

    report(8, "Organizando produtos e períodos do Excel")
    products_by_ncm: dict[str, list[str]] = defaultdict(list)
    for pair in pairs:
        if pair.product not in products_by_ncm[pair.ncm]:
            products_by_ncm[pair.ncm].append(pair.product)

    products_by_heading: dict[str, list[str]] = defaultdict(list)
    ncms_by_heading: dict[str, list[str]] = defaultdict(list)
    for pair in pairs:
        heading = heading_from_ncm(pair.ncm)
        if pair.product not in products_by_heading[heading]:
            products_by_heading[heading].append(pair.product)
        if pair.ncm not in ncms_by_heading[heading]:
            ncms_by_heading[heading].append(pair.ncm)

    selected_periods = list(iter_periods(start_period, end_period))
    include_year = start_period[:4] != end_period[:4]
    sorted_rows = sorted(
        api_rows,
        key=lambda row: (
            row_period(row),
            -validate_metric(row.get("metricCIF"), "metricCIF", "ordenação"),
            str(row.get("coNcm", "")),
            str(row.get("state", "")),
            str(row.get("urf", "")),
            str(row.get("country", "")),
        ),
    )
    excel_rows = [raw_row(row, include_year, products_by_ncm) for row in sorted_rows]
    allowed_headings = set(products_by_heading)
    municipal_keys = set()
    for row_number, row in enumerate(municipal_rows, start=1):
        if not isinstance(row, dict):
            raise ValueError(
                f"A linha {row_number} dos dados municipais não é um registro válido."
            )
        validate_municipal_api_row(
            row,
            requested_headings=allowed_headings,
            window_start=start_period,
            window_end=end_period,
            context=f"dados municipais do relatório, linha {row_number}",
        )
        key = municipal_row_identity(row)
        if key in municipal_keys:
            raise ValueError(
                "Registro estatístico repetido na base municipal. O relatório não foi publicado."
            )
        municipal_keys.add(key)
    sorted_municipal_rows = sorted(
        municipal_rows,
        key=lambda row: (
            row_period(row),
            -validate_metric(row.get("metricFOB"), "metricFOB", "ordenação municipal"),
            normalize_heading(row.get("headingCode")),
            municipal_name(row).casefold(),
            str(row.get("country", "")).casefold(),
        ),
    )
    municipal_excel_rows = [
        municipal_excel_row(
            row,
            include_year,
            products_by_heading,
            ncms_by_heading,
        )
        for row in sorted_municipal_rows
    ]

    report(18, "Preparando os resumos do Excel")
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = "LIARTE"
    workbook.properties.lastModifiedBy = "LIARTE"
    workbook.properties.title = (
        f"Importações Comex Stat — {start_period} a {end_period}"
    )
    workbook.properties.subject = (
        "Análise de importações por NCM e módulo municipal SH4"
    )
    workbook.properties.description = (
        "Relatório gerado pela LIARTE com Dados Gerais por NCM e uma base "
        "municipal SH4 separada, ambos provenientes do Comex Stat."
    )
    workbook.properties.keywords = "LIARTE, Comex Stat, importações, NCM"
    expected_totals = [
        sum_published(
            (row.get(field) for row in sorted_rows),
            places=0 if field == "metricKG" else 2,
        )
        for field in DEFAULT_METRICS
    ]
    expected_municipal_totals = [
        sum_published(
            (row.get(field) for row in sorted_municipal_rows),
            places=0 if field == "metricKG" else 2,
        )
        for field in MUNICIPAL_METRICS
    ]
    try:
        summary = workbook.create_sheet(SHEET_SUMMARY_NCM)
        summary.sheet_properties.tabColor = SUMMARY_TAB_COLOR
        build_summary(summary, sorted_rows, products_by_ncm)

        report(23, "Montando o resumo municipal")
        municipal_summary = workbook.create_sheet(SHEET_SUMMARY_MUNICIPAL)
        municipal_summary.sheet_properties.tabColor = SUMMARY_TAB_COLOR
        build_municipal_summary(municipal_summary, sorted_municipal_rows)

        report(28, f"Formatando a Base geral: {len(excel_rows):,} registros")
        create_data_sheets(workbook, SHEET_GENERAL_DATA, excel_rows)
        report(
            44, f"Formatando a Base municipal: {len(municipal_excel_rows):,} registros"
        )
        create_municipal_data_sheets(workbook, municipal_excel_rows)
        report(62, "Bases geral e municipal formatadas")

        if monthly_sheets:
            by_period: dict[str, list[list[Any]]] = defaultdict(list)
            for api_row, excel_row in zip(sorted_rows, excel_rows, strict=True):
                by_period[row_period(api_row)].append(excel_row)
            for index, period in enumerate(selected_periods):
                report(
                    62 + 18 * index / len(selected_periods),
                    f"Montando a aba mensal de {period}",
                )
                create_data_sheets(
                    workbook,
                    month_label(period, include_year),
                    by_period.get(period, []),
                    tab_color=MONTH_TAB_COLOR,
                )

        report(81, "Montando o cadastro de produtos no Excel")
        mapping = workbook.create_sheet(SHEET_PRODUCTS)
        mapping.sheet_properties.tabColor = REFERENCE_TAB_COLOR
        build_mapping(mapping, pairs)

        def verify_saved(path: Path) -> None:
            report(92, "Conferindo cada registro e os totais do Excel salvo")
            validate_generated_workbook(
                path,
                pairs=pairs,
                expected_rows=len(excel_rows),
                selected_periods=selected_periods,
                include_year=include_year,
                monthly_sheets=monthly_sheets,
                expected_totals=expected_totals,
                expected_municipal_rows=len(municipal_excel_rows),
                expected_municipal_totals=expected_municipal_totals,
                expected_general_data=excel_rows,
                expected_municipal_data=municipal_excel_rows,
            )

        report(85, "Gravando o arquivo Excel")
        save_workbook_atomic(workbook, output_path, validator=verify_saved)
        report(100, "Excel gerado e conferido")
    finally:
        workbook.close()


def default_output_path(start_period: str, end_period: str) -> Path:
    return Path(f"IMPORTACAO_{start_period}_{end_period}.xlsx")


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Consulta os Dados Gerais pelos NCMs incorporados ao aplicativo, consulta "
            "o módulo municipal por SH4 no Comex Stat e gera um Excel validado."
        )
    )
    parser.add_argument(
        "--inicio",
        type=validate_period,
        help="Mês inicial em AAAA-MM ou MM/AAAA.",
    )
    parser.add_argument(
        "--fim",
        type=validate_period,
        help="Mês final em AAAA-MM ou MM/AAAA. Padrão: igual ao início.",
    )
    parser.add_argument(
        "--ultimo-mes",
        action="store_true",
        help="Não pergunta o período e consulta o último mês publicado.",
    )
    parser.add_argument(
        "--saida",
        type=Path,
        help="Arquivo .xlsx de saída. Padrão: nome gerado a partir do período.",
    )
    parser.add_argument(
        "--lote-ncms",
        type=int,
        default=50,
        help="Quantidade de NCMs por requisição (padrão: 50).",
    )
    parser.add_argument(
        "--lote-sh4",
        type=int,
        default=25,
        help="Quantidade de posições SH4 por consulta municipal (padrão: 25).",
    )
    parser.add_argument(
        "--sem-municipios",
        action="store_true",
        help="Não consulta o módulo municipal; mantém as abas municipais vazias.",
    )
    parser.add_argument(
        "--timeout",
        type=float,
        default=120.0,
        help="Tempo limite por requisição, em segundos (padrão: 120).",
    )
    parser.add_argument(
        "--intervalo-api",
        type=float,
        default=10.5,
        help="Intervalo mínimo entre requisições, em segundos (padrão: 10.5).",
    )
    parser.add_argument(
        "--sem-abas-mensais",
        action="store_true",
        help="Gera apenas as bases consolidadas e as abas de resumo/controle.",
    )
    parser.add_argument(
        "--sem-power-bi",
        action="store_true",
        help="Não gera o pacote de projeto Power BI após criar o Excel.",
    )
    parser.add_argument(
        "--saida-power-bi",
        type=Path,
        help="Arquivo .zip do Power BI. Padrão: nome gerado ao lado do Excel.",
    )
    parser.add_argument(
        "--atualizar-cache",
        action="store_true",
        help="Consulta novamente todos os meses e substitui o cache validado.",
    )
    parser.add_argument(
        "--sem-cache",
        action="store_true",
        help="Consulta diretamente a API sem ler ou gravar cache local.",
    )
    return parser.parse_args(argv)


def run(args: argparse.Namespace) -> Path:
    if args.lote_ncms < 1 or args.lote_ncms > 200:
        raise ValueError("--lote-ncms deve estar entre 1 e 200.")
    if args.lote_sh4 < 1 or args.lote_sh4 > 200:
        raise ValueError("--lote-sh4 deve estar entre 1 e 200.")
    if not math.isfinite(args.timeout) or args.timeout <= 0:
        raise ValueError("--timeout deve ser maior que zero.")
    if not math.isfinite(args.intervalo_api) or args.intervalo_api < 0:
        raise ValueError("--intervalo-api não pode ser negativo.")
    if args.ultimo_mes and (args.inicio is not None or args.fim is not None):
        raise ValueError(
            "Use --ultimo-mes ou --inicio/--fim, não os dois modos juntos."
        )

    pairs = internal_product_ncms()
    ncms = sorted({pair.ncm for pair in pairs})
    print(
        f"Cadastro interno: {len(pairs)} produtos | {len(ncms)} NCMs únicos",
        flush=True,
    )

    client = ComexStatClient(
        timeout=args.timeout,
        min_interval=args.intervalo_api,
        use_cache=not args.sem_cache,
        refresh_cache=args.atualizar_cache,
    )
    municipal_status = "skipped" if args.sem_municipios else "complete"
    try:
        latest_period, updated = client.latest_published_period()
        if args.ultimo_mes:
            start_period = latest_period
            end_period = start_period
        elif args.inicio is None and args.fim is None and sys.stdin.isatty():
            start_period, end_period = choose_period_interactively()
        elif args.inicio is None and args.fim is None:
            # Execuções automáticas sem terminal continuam funcionando.
            start_period = latest_period
            end_period = start_period
        elif args.inicio is None:
            start_period = args.fim
            end_period = args.fim
        else:
            start_period = args.inicio
            end_period = args.fim or args.inicio

        if period_index(start_period) > period_index(end_period):
            raise ValueError(
                "O período inicial não pode ser posterior ao período final."
            )
        if end_period > latest_period:
            raise ValueError(
                f"O mês final ainda não foi publicado. Último período: {latest_period}."
            )
        print(
            f"Período selecionado: {start_period} até {end_period}",
            flush=True,
        )

        output_path = (
            args.saida or default_output_path(start_period, end_period)
        ).expanduser()
        if not output_path.is_absolute():
            output_path = (Path.cwd() / output_path).resolve()
        if output_path.suffix.lower() != ".xlsx":
            output_path = output_path.with_suffix(".xlsx")
        api_rows = client.query_imports(
            ncms=ncms,
            start_period=start_period,
            end_period=end_period,
            chunk_size=args.lote_ncms,
        )
        if args.sem_municipios:
            municipal_rows = []
        else:
            try:
                municipal_rows = client.query_municipal_imports(
                    ncms=ncms,
                    start_period=start_period,
                    end_period=end_period,
                    chunk_size=args.lote_sh4,
                )
            except Exception as exc:
                municipal_rows = []
                municipal_status = "unavailable"
                print(
                    "Aviso: os Dados Gerais serão preservados, mas o módulo "
                    "municipal ficou indisponível nesta tentativa "
                    f"({type(exc).__name__}: {exc}).",
                    file=sys.stderr,
                    flush=True,
                )
    finally:
        client.close()

    print(f"Linhas recebidas: {len(api_rows)}", flush=True)
    print(f"Linhas municipais recebidas: {len(municipal_rows)}", flush=True)
    create_report(
        output_path=output_path,
        pairs=pairs,
        api_rows=api_rows,
        start_period=start_period,
        end_period=end_period,
        monthly_sheets=not args.sem_abas_mensais,
        municipal_rows=municipal_rows,
    )
    print(f"Relatório salvo em: {output_path}", flush=True)
    powerbi_path = None
    if not args.sem_power_bi:
        try:
            from .powerbi_importacoes import (
                create_powerbi_package,
                default_powerbi_path,
            )
        except (ImportError, SystemExit) as exc:
            raise ValueError(
                "O Excel foi criado, mas o gerador de Power BI não foi encontrado. "
                "Mantenha powerbi_importacoes.py na mesma pasta ou use --sem-power-bi."
            ) from exc

        powerbi_path = args.saida_power_bi or default_powerbi_path(output_path)
        powerbi_path = powerbi_path.expanduser()
        if not powerbi_path.is_absolute():
            powerbi_path = (Path.cwd() / powerbi_path).resolve()
        if powerbi_path.suffix.casefold() != ".zip":
            powerbi_path = powerbi_path.with_suffix(".zip")
        powerbi_result = create_powerbi_package(
            excel_path=output_path,
            output_path=powerbi_path,
            start_period=start_period,
            end_period=end_period,
        )
        print(
            "Power BI salvo em: "
            f"{powerbi_result['path']} "
            f"({powerbi_result['pages']} páginas, "
            f"{powerbi_result['visuals']} visuais)",
            flush=True,
        )
    manifest_path = create_manifest(
        excel_path=output_path,
        powerbi_path=powerbi_path,
        version=GENERATOR_VERSION,
        catalog_version=INTERNAL_CATALOG_VERSION,
        pairs=pairs,
        start=start_period,
        end=end_period,
        general_rows=api_rows,
        municipal_rows=municipal_rows,
        number=to_number,
        municipal_status=municipal_status,
        publication=client.source_revision,
        cache=client.cache,
    )
    print(f"Rastreabilidade salva em: {manifest_path}", flush=True)
    return output_path


def main(argv: Sequence[str] | None = None) -> int:
    try:
        run(parse_args(argv))
    except (ComexStatError, FileNotFoundError, OSError, ValueError) as exc:
        print(f"Erro: {exc}", file=sys.stderr)
        return 1
    except KeyboardInterrupt:
        print("Operação cancelada.", file=sys.stderr)
        return 130
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
