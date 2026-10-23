"""Regressões sobre respostas oficiais capturadas; rede simulada somente nos testes."""

from __future__ import annotations
from dataclasses import replace
from decimal import Decimal
import copy
import io
import json
from pathlib import Path
from threading import Event
import tempfile
import time
import unittest
from unittest.mock import patch
from urllib.parse import urlsplit
from PySide6.QtCore import QPointF, QObject, QCoreApplication, QEvent, QMetaObject, QUrl
from PySide6.QtTest import QTest
from liarte_core.catalog import catalog
from liarte_core.comex import ComexProvider, ComexError, integral, uf
from liarte_core.domain import Period, Selection
from liarte_core.exports import ExportRequest, export_document, pdf_source
from liarte_core.query import Cancelled
from liarte_engine import comexstat_importacoes as engine
from liarte_engine.pdf_importacoes import _format_pt, consolidate_pdf_rows
from liarte_engine.powerbi_importacoes import validate_powerbi_project
from liarte_desktop.app import create_bundle
from tests.test_qt import APP, pump, settled, click as native_click


def click(window, item):
    scroll = window.findChild(QObject, "contentScroll")
    if scroll is not None:
        top = scroll.mapToScene(QPointF(0, 0)).y()
        y = item.mapToScene(QPointF(0, 0)).y()
        if y < top or y + item.height() > top + scroll.height():
            offset = scroll.property("contentY") + y - top - 60
            scroll.setProperty(
                "contentY",
                max(0, min(offset, scroll.property("contentHeight") - scroll.height())),
            )
            APP.processEvents()
    native_click(window, item)


FIXTURE = json.loads(
    (Path(__file__).parent / "fixtures/comex_2026-09.json").read_text(encoding="utf-8")
)


class Response(io.BytesIO):
    status = 200

    def __init__(self, document):
        raw = json.dumps(document).encode()
        super().__init__(raw)
        self.headers = {"Content-Length": str(len(raw))}


class Transport:
    def __init__(self):
        self.general = copy.deepcopy(FIXTURE["general"])
        self.municipal = copy.deepcopy(FIXTURE["municipal"])
        self.calls = []
        self.updated = "2026-10-06"
        self.fail_municipal = False

    def __call__(self, request, **kwargs):
        path = urlsplit(request.full_url).path
        self.calls.append((request.method, path))
        if request.method == "GET":
            return Response(
                {
                    "success": True,
                    "data": {
                        "year": "2026",
                        "monthNumber": "09",
                        "updated": self.updated,
                    },
                }
            )
        if path == "/cities" and self.fail_municipal:
            raise engine.ComexStatError("Falha municipal controlada")
        body = json.loads(request.data)
        codes = set(map(str, body["filters"][0]["values"]))
        rows = self.municipal if path == "/cities" else self.general
        field = "headingCode" if path == "/cities" else "coNcm"
        rows = [
            r
            for r in rows
            if r[field] in codes
            and body["period"]["from"] <= engine.row_period(r) <= body["period"]["to"]
        ]
        if len(body["details"]) == 1:
            grouped = {}
            metrics = (
                engine.MUNICIPAL_METRICS
                if path == "/cities"
                else engine.DEFAULT_METRICS
            )
            for row in rows:
                key = (row["year"], row["monthNumber"], row[field])
                target = grouped.setdefault(
                    key,
                    {
                        "year": key[0],
                        "monthNumber": key[1],
                        field: key[2],
                        **dict.fromkeys(metrics, 0),
                    },
                )
                for metric in metrics:
                    target[metric] += Decimal(str(row[metric]))
            rows = [
                {**row, **{metric: str(row[metric]) for metric in metrics}}
                for row in grouped.values()
            ]
        return Response({"success": True, "data": {"list": rows}})


def factory(**kwargs):
    kwargs["min_interval"] = 0
    return engine.ComexStatClient(**kwargs)


class ComexFixtureCase(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.directory = Path(self.temp.name)
        self.transport = Transport()
        self.patch = patch(
            "liarte_engine.secure_http.urlopen",
            side_effect=self.transport,
        )
        self.patch.start()
        self.provider = ComexProvider(
            self.directory,
            client_factory=factory,
            codes=["29161220"],
            verify_official_totals=getattr(self, "verify_official_totals", False),
        )
        self.selection = Selection(
            Period("2026-09", "2026-09"), sort_column=9, descending=True
        )

    def tearDown(self):
        self.patch.stop()
        self.temp.cleanup()

    def query(self, **changes):
        return self.provider.query(
            replace(self.selection, **changes), Event(), lambda *_: None
        )


class ComexTests(ComexFixtureCase):
    def test_fixture_queries_never_reach_either_platform_network_transport(self):
        with (
            patch(
                "urllib.request.urlopen",
                side_effect=AssertionError("Unexpected standard network request"),
            ),
            patch(
                "liarte_engine.secure_http._origin_opener",
                side_effect=AssertionError("Unexpected Windows network request"),
            ),
        ):
            result = self.query()
        self.assertEqual(len(result.rows), len(FIXTURE["general"]))
        self.assertTrue(self.transport.calls)

    def test_exact_general_totals_from_official_fixture(self):
        result = self.query()
        self.assertEqual(
            result.totals.value_cents,
            sum(int(r["metricCIF"]) * 100 for r in FIXTURE["general"]),
        )
        self.assertEqual(
            result.totals.kg, sum(int(r["metricKG"]) for r in FIXTURE["general"])
        )
        self.assertEqual(len(result.rows), 18)
        self.assertEqual(result.rows[0].product, "Acrilato de etila")
        self.assertIn("Ésteres", result.rows[0].official_description)

    def test_municipal_is_sh4_fob_and_never_cif(self):
        result = self.query(kind="municipal", sort_column=6)
        self.assertEqual(result.totals.basis, "FOB")
        with self.assertRaises(ValueError):
            _ = result.totals.cif_cents
        self.assertEqual(
            result.totals.value_cents,
            sum(int(r["metricFOB"]) * 100 for r in FIXTURE["municipal"]),
        )
        self.assertEqual({r.heading for r in result.rows}, {"2916"})
        self.assertFalse(hasattr(result.rows[0], "ncm"))

    def test_city_filter_matches_independent_sum(self):
        self.query()
        city = FIXTURE["municipal"][0]["noMunMinsgUf"]
        result = self.query(kind="municipal", city=city, sort_column=6)
        expected = [r for r in FIXTURE["municipal"] if r["noMunMinsgUf"] == city]
        self.assertEqual(
            result.totals.fob_cents, sum(int(r["metricFOB"]) * 100 for r in expected)
        )
        self.assertTrue(all(r.city == city for r in result.rows))

    def test_filters_pages_and_dataset_switch_do_not_call_network(self):
        self.query()
        count = len(self.transport.calls)
        self.query(country="China")
        self.query(page=2, page_size=25)
        self.query(kind="municipal", sort_column=6)
        self.assertEqual(count, len(self.transport.calls))

    def test_new_provider_refetches_api_instead_of_reusing_persistent_results(self):
        first = self.query()
        posts = sum(method == "POST" for method, _ in self.transport.calls)
        other = ComexProvider(
            self.directory, client_factory=factory, codes=["29161220"]
        )
        second = other.query(self.selection, Event(), lambda *_: None)
        self.assertEqual(first.totals, second.totals)
        self.assertEqual(
            posts + 2, sum(method == "POST" for method, _ in self.transport.calls)
        )
        self.assertEqual(other.snapshot.cache_hits, 0)

    def test_force_refresh_fetches_both_sources(self):
        self.query()
        posts = sum(m == "POST" for m, _ in self.transport.calls)
        self.provider.force_refresh()
        self.query()
        self.assertEqual(sum(m == "POST" for m, _ in self.transport.calls), posts + 2)

    def test_revision_change_invalidates_cached_rows(self):
        self.query()
        self.transport.updated = "2026-10-07"
        self.provider.availability(Event(), force=True)
        self.provider.revalidate_next()
        before = sum(m == "POST" for m, _ in self.transport.calls)
        self.query()
        self.assertEqual(sum(m == "POST" for m, _ in self.transport.calls), before + 2)

    def test_second_source_failure_preserves_complete_snapshot(self):
        self.query()
        before = self.provider.snapshot
        self.transport.fail_municipal = True
        self.provider.force_refresh()
        with self.assertRaises(ComexError):
            self.query()
        self.assertIs(self.provider.snapshot, before)

    def test_divergent_cif_does_not_replace_a_snapshot(self):
        self.query()
        before = self.provider.snapshot
        self.transport.general[0]["metricCIF"] = "1"
        self.provider.force_refresh()
        with self.assertRaises(ComexError):
            self.query()
        self.assertIs(before, self.provider.snapshot)

    def test_cancel_prevents_any_network_access(self):
        cancel = Event()
        cancel.set()
        with self.assertRaises(Cancelled):
            self.provider.query(self.selection, cancel, lambda *_: None)
        self.assertEqual(self.transport.calls, [])

    def test_backoff_is_cancellable(self):
        cancel = Event()
        cancel.set()
        client = factory(
            cancel_check=lambda: __import__(
                "liarte_core.query", fromlist=["checkpoint"]
            ).checkpoint(cancel),
            cancel_wait=cancel.wait,
        )
        start = time.monotonic()
        with self.assertRaises(Cancelled):
            client._wait(120)
        self.assertLess(time.monotonic() - start, 0.1)

    def test_future_period_is_rejected_before_dataset_post(self):
        with self.assertRaises(ComexError):
            self.query(period=Period("2026-10", "2026-10"))
        self.assertFalse(any(m == "POST" for m, _ in self.transport.calls))

    def test_offline_availability_is_explicit(self):
        self.provider.availability(Event())
        with patch(
            "liarte_engine.comexstat_importacoes.ComexStatClient.latest_published_period",
            side_effect=engine.ComexStatError("offline"),
        ):
            with self.assertRaises(ComexError):
                self.provider.availability(Event(), force=True)
            self.provider._availability_time = time.monotonic() - 301
            fallback = self.provider.availability(Event())
        self.assertTrue(fallback.cached)
        self.assertEqual(fallback.municipal_latest, "2026-09")

    def test_decimal_precision_never_rounds_imported_metrics(self):
        self.assertEqual(integral("123.45", cents=True), 12345)
        for value in ("NaN", "Infinity", "-1", "0.001"):
            with self.assertRaises(ComexError):
                integral(value, cents=True)
        with self.assertRaises(ComexError):
            integral("1.5")

    def test_unknown_uf_is_not_silently_dropped(self):
        self.assertEqual(uf("Rio de Janeiro"), "RJ")
        self.assertEqual(uf("Não Declarada"), "ND")
        with self.assertRaises(ComexError):
            uf("UF inexistente")

    def test_pdf_rounding_uses_decimal_half_up(self):
        self.assertEqual(_format_pt(Decimal("1.2345"), 3), "1,235")
        self.assertEqual(_format_pt(Decimal("12345.5"), 3), "12.345,5")

    def test_catalog_ncms_and_product_relationships_match_engine(self):
        self.assertEqual(
            {r["ncm"] for r in catalog()},
            {p.ncm for p in engine.internal_product_ncms()},
        )
        self.assertEqual(
            sum(len(r["products"]) for r in catalog()),
            len(engine.internal_product_ncms()),
        )

    def test_multiple_countries_and_states_use_union_then_intersection(self):
        result = self.query(
            country_values=("China", "Estados Unidos"), state_values=("SP", "SC")
        )
        raw = [
            r
            for r in FIXTURE["general"]
            if r["country"] in ("China", "Estados Unidos")
            and r["state"] in ("São Paulo", "Santa Catarina")
        ]
        self.assertEqual(
            result.totals.value_cents, sum(int(r["metricCIF"]) * 100 for r in raw)
        )
        self.assertEqual(result.totals.records, len(raw))

    def test_out_of_catalog_code_is_rejected(self):
        with self.assertRaises(ComexError):
            self.query(ncm="99999999")
        with self.assertRaises(ValueError):
            ComexProvider(self.directory, codes=["99999999"])

    def test_multiple_values_are_canonical_and_validated(self):
        a = replace(self.selection, country_values=("China", "Alemanha", "China"))
        b = replace(self.selection, country_values=("Alemanha", "China"))
        self.assertEqual(a, b)
        with self.assertRaises(ValueError):
            replace(self.selection, country_values=["China"])
        with self.assertRaises(ValueError):
            replace(self.selection, ncm_values=("2916",))

    def test_http_rate_limit_retries_are_bounded(self):
        from urllib.error import HTTPError

        errors = [
            HTTPError(
                "https://api-comexstat.mdic.gov.br/general",
                429,
                "limit",
                {"Retry-After": "1"},
                io.BytesIO(b"limit"),
            )
            for _ in range(3)
        ]
        client = factory(max_attempts=3)
        waits = []
        with (
            patch(
                "liarte_engine.secure_http.urlopen",
                side_effect=errors,
            ) as request,
            patch.object(client, "_wait", side_effect=waits.append),
            self.assertRaises(engine.ComexStatError),
        ):
            client.request_json("POST", "/general", json_body={})
        self.assertEqual(request.call_count, 3)
        self.assertEqual(len(waits), 2)


class ExportTests(ComexFixtureCase):
    def request(self, file_format, **changes):
        self.query()
        return ExportRequest(
            file_format,
            self.directory
            / (
                "report."
                + {"excel": "xlsx", "pdf": "pdf", "powerbi": "zip"}[file_format]
            ),
            self.provider.snapshot,
            replace(self.selection, **changes),
        )

    def export(self, request, cancel=None):
        return export_document(
            request, cancel or Event(), lambda *_: None, self.directory / "audit"
        )

    def test_excel_totals_uf_width_and_vertical_alignment(self):
        from openpyxl import load_workbook

        request = self.request("excel")
        self.export(request)
        book = load_workbook(request.destination)
        try:
            self.assertNotIn("Parâmetros", book.sheetnames)
            self.assertEqual(book["Base municipal"].column_dimensions["E"].width, 64)
            general = list(book["Base geral"].values)
            headers = list(general[0])
            self.assertEqual(
                sum(r[headers.index("CIF")] for r in general[1:]),
                sum(int(r["metricCIF"]) for r in FIXTURE["general"]),
            )
            for sheet in book:
                for row in sheet:
                    for cell in row:
                        if cell.value is not None:
                            self.assertEqual(
                                cell.alignment.vertical,
                                "center",
                                (sheet.title, cell.coordinate),
                            )
        finally:
            book.close()

    def test_excel_ignores_analytical_filters(self):
        from openpyxl import load_workbook

        request = self.request("excel", country="China", search="nao existe")
        self.export(request)
        book = load_workbook(request.destination, read_only=True)
        try:
            self.assertEqual(book["Base geral"].max_row, len(FIXTURE["general"]) + 1)
        finally:
            book.close()

    def test_pdf_uses_filters_and_search_across_all_pages(self):
        request = self.request("pdf", state="SP", search="China", page=5)
        source = pdf_source(request, Event())
        expected = [
            r
            for r in FIXTURE["general"]
            if r["state"] == "São Paulo" and r["country"] == "China"
        ]
        self.assertEqual(source, expected)

    def test_pdf_commercial_names_no_ncm_code_or_noisy_footer(self):
        from pypdf import PdfReader

        request = self.request("pdf")
        self.export(request)
        reader = PdfReader(request.destination)
        text = "".join(p.extract_text() for p in reader.pages)
        self.assertIn("Acrilato de etila", text)
        self.assertIn("Importações filtradas", text)
        self.assertNotIn("29161220", text)
        self.assertNotIn("LIARTE - Importações", text)
        self.assertNotIn("linhas -", text)
        self.assertNotIn("Ésteres de etila", text)
        for p in reader.pages:
            self.assertAlmostEqual(float(p.mediabox.width), 595.2756, places=2)
            self.assertIn("Descrição NCM", p.extract_text())

    def test_pdf_repeats_headers_when_multiple_pages(self):
        from liarte_engine.pdf_importacoes import create_import_pdf
        from pypdf import PdfReader

        raw = []
        for i in range(125):
            row = copy.deepcopy(FIXTURE["general"][0])
            row["country"] = f"País de teste {i:03d}"
            raw.append(row)
        path = self.directory / "multi.pdf"
        create_import_pdf(
            path, engine.internal_product_ncms(), raw, "2026-09", "2026-09"
        )
        reader = PdfReader(path)
        self.assertGreater(len(reader.pages), 1)
        self.assertTrue(all("Descrição NCM" in p.extract_text() for p in reader.pages))

    def test_pdf_excludes_exactly_1000kg_and_airports(self):
        raw = copy.deepcopy(FIXTURE["general"][:3])
        raw[0]["metricKG"] = "1000"
        raw[1]["metricKG"] = "1001"
        raw[1]["urf"] = "AEROPORTO INTERNACIONAL"
        raw[2]["metricKG"] = "1001"
        rows = consolidate_pdf_rows([], raw)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["kg"], 1001)

    def test_municipal_pdf_is_not_misrepresented_as_cif(self):
        request = self.request("pdf", kind="municipal", sort_column=6)
        with self.assertRaises(ValueError):
            self.export(request)
        self.assertFalse(request.destination.exists())

    def test_powerbi_package_contains_valid_pages_and_source(self):
        import zipfile

        request = self.request("powerbi")
        result = self.export(request)
        self.assertGreaterEqual(result["pages"], 6)
        with zipfile.ZipFile(request.destination) as archive:
            self.assertIsNone(archive.testzip())
            self.assertTrue(any(n.endswith(".xlsx") for n in archive.namelist()))
            self.assertTrue(any(n.endswith(".pbip") for n in archive.namelist()))
            destination = self.directory / "bi"
            archive.extractall(destination)
        project = next(destination.rglob("*.pbip")).parent
        self.assertGreater(validate_powerbi_project(project)["visuals"], 20)

    def test_cancelled_excel_preserves_existing_destination(self):
        request = self.request("excel")
        request.destination.write_bytes(b"existing file")
        cancel = Event()

        def progress(percent, *_):
            if percent >= 44:
                cancel.set()

        with self.assertRaises(Cancelled):
            export_document(request, cancel, progress, self.directory / "audit")
        self.assertEqual(request.destination.read_bytes(), b"existing file")

    def test_audit_failure_does_not_hide_a_saved_document(self):
        request = self.request("pdf")
        with patch("liarte_core.exports.atomic_json", side_effect=OSError("disk full")):
            result = self.export(request)
        self.assertTrue(request.destination.is_file())
        self.assertEqual(result["name"], request.destination.name)


class RealQtTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.directory = Path(self.temp.name)
        self.transport = Transport()
        self.patch = patch(
            "liarte_engine.secure_http.urlopen",
            side_effect=self.transport,
        )
        self.patch.start()
        self.provider = ComexProvider(
            self.directory, client_factory=factory, codes=["29161220"]
        )
        self.bundle = create_bundle(initial_module="imports", provider=self.provider)
        self.controller = self.bundle.controller
        self.window = self.bundle.engine.rootObjects()[0]
        APP.processEvents()

    def tearDown(self):
        self.controller.shutdown()
        self.window.close()
        APP.processEvents()
        self.bundle.engine.deleteLater()
        QCoreApplication.sendPostedEvents(None, QEvent.Type.DeferredDelete)
        APP.processEvents()
        self.patch.stop()
        self.temp.cleanup()

    def item(self, name):
        item = self.window.findChild(QObject, name)
        self.assertIsNotNone(item, name)
        return item

    def load(self):
        self.controller.queryPeriod("2026-09", "2026-09")
        settled(self.controller)
        self.assertTrue(self.controller.state["loaded"], self.controller.state["error"])

    def test_real_mode_loads_official_data_without_synthetic_fallback(self):
        self.assertFalse(self.controller.state["demo"])
        self.load()
        self.assertEqual(len(self.controller.table.rows), 18)
        self.assertEqual(
            self.controller.state["metrics"][0]["full"], "US$ 2.340.152,00"
        )
        self.assertEqual(self.bundle.messages, [])

    def test_municipal_tab_has_correct_headers_metrics_and_city_filter(self):
        self.load()
        click(self.window, self.item("municipalTab"))
        settled(self.controller)
        self.assertEqual(self.controller.table.columnCount(), 8)
        self.controller.table.setProfile("full")
        self.assertEqual(self.controller.table.columnCount(), 9)
        self.assertEqual(self.controller.state["metrics"][0]["label"], "Valor FOB")
        self.assertEqual(self.controller.state["metrics"][3]["label"], "SH4 ativos")
        self.assertFalse(self.item("exportPdf").property("enabled"))
        city = self.provider.options("municipal")["cities"][0]
        self.controller.setFilter("city", city)
        settled(self.controller)
        self.assertTrue(all(r.city == city for r in self.controller.table.rows))
        self.assertEqual(self.bundle.messages, [])

    def test_metadata_refresh_keeps_ui_responsive_and_updates_periods(self):
        self.controller.refreshAvailability()
        self.assertTrue(pump(lambda: not self.controller.state["metadataBusy"]))
        self.assertEqual(self.controller.periods[-1]["value"], "2026-09")
        self.assertEqual(self.controller.state["sourceUpdated"], "2026-10-06")

    def test_quick_month_uses_the_published_month(self):
        self.controller.refreshAvailability()
        pump(lambda: not self.controller.state["metadataBusy"])
        click(self.window, self.item("cancelButton"))
        settled(self.controller)
        self.assertEqual(self.controller.state["appliedStart"], "2026-09")
        self.assertEqual(self.controller.state["appliedEnd"], "2026-09")

    def test_clear_filters_resets_search_text(self):
        self.load()
        self.controller.setFilter("search", "China")
        settled(self.controller)
        self.assertEqual(self.item("searchBox").property("text"), "China")
        self.controller.clearFilters()
        settled(self.controller)
        self.assertEqual(self.item("searchBox").property("text"), "")

    def test_pending_filters_disable_export(self):
        self.load()
        self.assertTrue(self.controller.state["canExport"])
        self.controller.setFilter("country", "China")
        self.assertFalse(self.controller.state["canExport"])
        settled(self.controller)
        self.assertTrue(self.controller.state["canExport"])

    def test_native_save_dialog_opens_with_the_right_format(self):
        self.load()
        click(self.window, self.item("exportExcel"))
        self.assertTrue(pump(lambda: self.item("exportDialog").property("visible")))
        self.assertEqual(self.item("exportDialog").property("defaultSuffix"), "xlsx")
        QMetaObject.invokeMethod(self.item("exportDialog"), "close")
        pump(lambda: not self.item("exportDialog").property("visible"))
        self.assertEqual(self.bundle.messages, [])

    def test_excel_export_runs_in_a_worker_and_saves(self):
        self.load()
        path = self.directory / "generated.xlsx"
        self.controller.exportTo("excel", QUrl.fromLocalFile(str(path)))
        self.assertTrue(pump(lambda: not self.controller.state["exportBusy"], 15))
        self.assertTrue(path.is_file(), self.controller.state["exportError"])
        self.assertEqual(self.controller.state["lastExport"], path.name)

    def test_pdf_export_has_no_success_notification(self):
        self.load()
        path = self.directory / "generated.pdf"
        self.controller.exportTo("pdf", QUrl.fromLocalFile(str(path)))
        self.assertTrue(pump(lambda: not self.controller.state["exportBusy"], 15))
        self.assertTrue(path.is_file(), self.controller.state["exportError"])
        self.assertEqual(self.controller.state["exportMessage"], "")

    def test_source_failure_preserves_previous_rows_and_disables_export(self):
        self.load()
        before = self.controller.table.rows
        self.provider.force_refresh()
        self.transport.fail_municipal = True
        self.controller._queue(self.controller._selection)
        settled(self.controller)
        self.assertEqual(self.controller.table.rows, before)
        self.assertTrue(self.controller.state["error"])

    def test_headers_fit_when_switching_back_to_general(self):
        self.load()
        self.controller.setDataset("municipal")
        settled(self.controller)
        self.controller.setDataset("general")
        settled(self.controller)
        self.assertEqual(self.controller.table.columnCount(), 8)
        self.controller.table.setProfile("full")
        self.assertEqual(self.controller.table.columnCount(), 12)
        self.assertIn("CIF (US$)", self.controller.state["headers"])
        self.assertEqual(self.bundle.messages, [])

    def test_searchable_filter_accepts_accentless_typing_and_keyboard(self):
        from PySide6.QtCore import Qt

        self.load()
        control = self.item("countryFilter")
        click(self.window, control)
        search = self.item("countryFilterSearch")
        self.assertTrue(pump(lambda: search.hasActiveFocus()))
        for letter in "AFRICA":
            QTest.keyClick(self.window, Qt.Key(ord(letter)))
        APP.processEvents()
        QTest.keyClick(self.window, Qt.Key.Key_Return)
        settled(self.controller)
        self.assertEqual(self.controller.state["countryFilter"], "África do Sul")
        self.assertTrue(
            all(r.country == "África do Sul" for r in self.controller.table.rows)
        )
        self.assertEqual(self.bundle.messages, [])

    def test_period_without_a_selected_country_keeps_the_filter_visible(self):
        self.load()
        self.controller.setFilter("country", "China")
        settled(self.controller)
        self.controller.queryPeriod("2026-08", "2026-08")
        settled(self.controller)
        self.assertEqual(self.controller.state["scopeCount"], "0")
        self.assertEqual(self.controller.state["countryFilter"], "China")
        self.assertIn("China", self.item("countryFilter").property("displayText"))
        self.assertEqual(self.bundle.messages, [])

    def test_failed_refresh_disables_export_until_a_successful_query(self):
        self.load()
        self.provider.force_refresh()
        self.transport.fail_municipal = True
        self.controller._queue(self.controller._selection)
        settled(self.controller)
        self.assertFalse(self.controller.state["canExport"])

    def test_multiple_filters_can_be_cleared_without_hidden_selection(self):
        self.load()
        self.controller.setFilters("country", ["China", "Estados Unidos"])
        settled(self.controller)
        self.assertEqual(
            self.controller.state["countryValues"], ["China", "Estados Unidos"]
        )
        self.assertEqual(
            self.item("countryFilter").property("displayText"), "2 selecionados"
        )
        self.controller.setFilters("country", [])
        settled(self.controller)
        self.assertEqual(self.controller.state["countryValues"], [])
        self.assertEqual(self.controller.state["scopeCount"], "18")

    def test_cancelling_pending_multi_filters_restores_controls(self):
        self.load()
        self.controller.setFilters("country", ["China", "Estados Unidos"])
        self.controller.cancel()
        self.assertEqual(self.controller.state["countryValues"], [])
        self.assertEqual(self.item("countryFilter").property("currentIndex"), 0)
        self.assertEqual(self.bundle.messages, [])
