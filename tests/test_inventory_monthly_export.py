"""Excel mensal: histórico completo, unidades separadas e lacunas preservadas."""

from dataclasses import replace
from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory
from threading import Event
from unittest import TestCase
from unittest.mock import patch

from openpyxl import load_workbook

from liarte_core.inventory import (
    SCALE,
    Balance,
    Closing,
    InventoryData,
    Product,
    Sale,
    StockSelection,
    analyze_inventory,
)
from liarte_core.inventory_export import export_inventory
from liarte_core.inventory_monthly import analyze_monthly_inventory


def monthly_export_data():
    return InventoryData(
        date(2026, 4, 15),
        "Base sintética mensal",
        "a" * 64,
        (
            Product("A", "Produto atual", "Grupo", "PA", "KG"),
            Product("B", "=SOMA(1;2)", "Grupo", "PA", "KG"),
            Product("C", "Somente vendas", "Grupo", "PA", "L"),
            Product("D", "Produto novo no histórico", "Grupo", "MP", "UN"),
        ),
        (
            Balance("A", "01", "01", "KG", 20 * SCALE, 24000),
            Balance("A", "01", "02", "KG", 5 * SCALE, 5000),
            Balance("D", "01", "01", "UN", 7 * SCALE, 700),
        ),
        (
            Closing("A", "01", "01", "KG", "2026-01", 10 * SCALE, 10000, 2, 15),
            Closing("A", "01", "02", "KG", "2026-01", 5 * SCALE, 5000, 4, 7.5),
            Closing("B", "01", "01", "KG", "2026-01", 3 * SCALE, 6000, 1, 30),
            Closing("A", "01", "01", "KG", "2026-02", 20 * SCALE, 24000, 3, 10),
            Closing("A", "01", "02", "KG", "2026-02", 5 * SCALE, 5000, -2, -15),
            Closing("B", "01", "01", "KG", "2026-02", 2 * SCALE, 4000, 5, 6),
            Closing("D", "01", "01", "UN", "2026-02", 7 * SCALE, 700, 0, None),
        ),
        (
            Sale("A", "01", "KG", date(2026, 1, 1), SCALE, 100),
            Sale("B", "01", "KG", date(2026, 1, 10), 2 * SCALE, 200),
            Sale("C", "03", "L", date(2026, 1, 5), 4 * SCALE, 400),
            Sale("A", "01", "KG", date(2026, 2, 10), 3 * SCALE, 300),
            Sale("B", "01", "KG", date(2026, 2, 10), 8 * SCALE, 800),
            Sale("C", "03", "L", date(2026, 2, 1), 6 * SCALE, 600),
            Sale("A", "01", "KG", date(2026, 4, 1), SCALE, 100),
        ),
    )


class InventoryMonthlyExportTests(TestCase):
    def _export(self, directory, selection=None, *, monthly=None):
        data = monthly_export_data()
        selection = selection or StockSelection("2026-01", "2026-04")
        path = Path(directory) / "mensal.xlsx"
        export_inventory(
            data, analyze_inventory(data, selection), selection, path, Event(), monthly=monthly
        )
        return load_workbook(path)

    @staticmethod
    def _rows(sheet):
        return list(sheet.iter_rows(min_row=2, values_only=True))

    def test_monthly_summaries_preserve_missing_cost_zero_sales_and_separate_units(self):
        with TemporaryDirectory() as directory:
            book = self._export(directory)
            try:
                self.assertTrue(
                    {
                        "Estoque",
                        "Fechamentos",
                        "Vendas por filial",
                        "Leitura dos indicadores",
                    }.issubset(book.sheetnames)
                )
                rows = self._rows(book["Análise mensal"])
                totals = {r[0]: r for r in rows if r[1] == "Resumo mensal"}
                units = {(r[0], r[2]): r for r in rows if r[1] == "Quantidade por unidade"}
                self.assertEqual(totals["2026-01"][3], 210)
                self.assertEqual(totals["2026-02"][3], 337)
                self.assertEqual(totals["2026-02"][6:9], (3, 3, 1))
                self.assertEqual(totals["2026-02"][10:12], (120, 127))
                self.assertEqual(totals["2026-02"][14:16], (1, 0))
                self.assertEqual(totals["2026-02"][23:25], (7, 0))
                self.assertEqual(
                    totals["2026-02"][11],
                    totals["2026-02"][10] + totals["2026-02"][23] - totals["2026-02"][24],
                )
                self.assertIsNone(totals["2026-03"][3])
                self.assertEqual(totals["2026-04"][13], "Sim")
                self.assertEqual(units["2026-02", "KG"][4:6], (27, 11))
                self.assertEqual(units["2026-02", "L"][4:6], (None, 6))
                self.assertEqual(units["2026-03", "KG"][4:6], (None, 0))
                # The retained worksheet follows the current-position cohort;
                # the monthly worksheet also includes the history-only B.
                old_sales = {(r[0], r[1]): r[2] for r in self._rows(book["Vendas por filial"])}
                self.assertEqual(old_sales["2026-02", "KG"], 3)
                self.assertEqual(units["2026-02", "KG"][5], 11)
                self.assertIn(
                    "posições atuais selecionadas", book["Vendas por filial"].cell(1, 3).value
                )
                rules = dict(book["Leitura dos indicadores"].iter_rows(min_row=2, values_only=True))
                self.assertIn(
                    "Produtos sem posição atual não entram",
                    rules["Vendas por filial · posições atuais"],
                )
                self.assertIn("inclusive produtos sem posição atual", rules["Vendas mensais"])
                # R$ appears once per month; units cannot duplicate capital or
                # present a fictitious sum of liters and kilograms.
                self.assertTrue(all(r[3] is None and r[6] is None for r in units.values()))
                self.assertEqual(sum(r[3] or 0 for r in rows), 547)
            finally:
                book.close()

    def test_full_monthly_products_include_historical_and_sales_only_products_without_duplicating_sales(
        self,
    ):
        with TemporaryDirectory() as directory:
            book = self._export(directory)
            try:
                rows = self._rows(book["Produtos mensais"])
                feb = [r for r in rows if r[0] == "2026-02"]
                historical = [r for r in feb if r[1] == "Fechamento" and r[2] == "B"]
                self.assertEqual(len(historical), 1)
                self.assertEqual(historical[0][7:10], (2, None, 40))
                a_sales = [r for r in feb if r[1] == "Vendas da filial" and r[2] == "A"]
                self.assertEqual(len(a_sales), 1)
                self.assertEqual(a_sales[0][8], 3)
                c_sales = [r for r in feb if r[1] == "Vendas da filial" and r[2] == "C"]
                self.assertEqual(c_sales[0][4:10], ("03", None, "L", None, 6, None))
                jan_b = next(r for r in rows if r[:3] == ("2026-01", "Fechamento", "B"))
                self.assertEqual(jan_b[12].date(), date(2026, 1, 10))
                self.assertEqual(jan_b[13], 21)
                # Negative source indicators remain original, not sanitised.
                negative = next(
                    r for r in feb if r[1] == "Fechamento" and r[2] == "A" and r[5] == "02"
                )
                self.assertEqual(negative[10:12], (-2, -15))
                bad_name = next(
                    row[3]
                    for row in book["Produtos mensais"].iter_rows()
                    if row[3].value == "=SOMA(1;2)"
                )
                self.assertEqual(bad_name.data_type, "s")
            finally:
                book.close()

    def test_contributions_keep_paired_values_and_coverage_context(self):
        with TemporaryDirectory() as directory:
            book = self._export(directory)
            try:
                rows = self._rows(book["Variações mensais"])
                costs = {
                    r[4]: r[6]
                    for r in rows
                    if r[0] == "2026-02" and r[2] == "Custo · posições pareadas"
                }
                self.assertEqual(costs, {"A": 140, "B": -20})
                coverage = next(
                    r[7] for r in rows if r[0] == "2026-02" and r[2] == "Abrangência da comparação"
                )
                self.assertIn("1 posições novas", coverage)
                self.assertIn("não são presumidos como saldo zero", coverage)
                changes = {
                    (r[3], r[4]): r[6]
                    for r in rows
                    if r[0] == "2026-02" and r[2] == "Vendas das filiais"
                }
                self.assertEqual(changes, {("KG", "A"): 2, ("KG", "B"): 6, ("L", "C"): 2})
                self.assertFalse(
                    any(r[0] == "2026-04" and r[2] == "Vendas das filiais" for r in rows)
                )
            finally:
                book.close()

    def test_warehouse_filter_keeps_monthly_sales_blank_and_cached_report_is_reused(self):
        data = monthly_export_data()
        selection = StockSelection("2026-01", "2026-04", warehouse="01")
        monthly = analyze_monthly_inventory(data, selection)
        with (
            TemporaryDirectory() as directory,
            patch(
                "liarte_core.inventory_monthly.analyze_monthly_inventory",
                side_effect=AssertionError("Relatório mensal já calculado"),
            ),
        ):
            book = self._export(directory, selection, monthly=monthly)
            try:
                rows = self._rows(book["Análise mensal"])
                self.assertTrue(all(r[5] is None for r in rows))
                self.assertFalse(
                    any(r[1] == "Vendas da filial" for r in self._rows(book["Produtos mensais"]))
                )
                self.assertTrue(all("armazém" in r[17] for r in rows if r[1] == "Resumo mensal"))
                old_sales = {(r[0], r[1]): r[2] for r in self._rows(book["Vendas por filial"])}
                self.assertEqual(old_sales["2026-02", "KG"], 3)
                rules = dict(book["Leitura dos indicadores"].iter_rows(min_row=2, values_only=True))
                self.assertIn("filial inteira", rules["Vendas por filial · posições atuais"])
                self.assertIn("indisponibiliza as vendas nessas duas abas", rules["Vendas mensais"])
            finally:
                book.close()

    def test_export_keeps_all_products_beyond_the_visible_ranking(self):
        codes = [f"{i:04d}" for i in range(16)]
        data = replace(
            monthly_export_data(),
            products=tuple(Product(code, "Produto " + code, "Grupo", "PA", "KG") for code in codes),
            balances=(),
            closings=tuple(
                Closing(code, "01", "01", "KG", period, (i + 1) * SCALE, (i + 1) * 100, 1, 30)
                for period in ("2026-01", "2026-02")
                for i, code in enumerate(codes)
            ),
            sales=tuple(
                Sale(code, "01", "KG", date(2026, month, 1), (i + 1) * SCALE, (i + 1) * 100)
                for month in (1, 2)
                for i, code in enumerate(codes)
            ),
        )
        selection = StockSelection("2026-01", "2026-02")
        with TemporaryDirectory() as directory:
            path = Path(directory) / "todos.xlsx"
            export_inventory(data, analyze_inventory(data, selection), selection, path, Event())
            book = load_workbook(path)
            try:
                rows = self._rows(book["Produtos mensais"])
                for period in ("2026-01", "2026-02"):
                    for kind in ("Fechamento", "Vendas da filial"):
                        self.assertEqual(
                            {r[2] for r in rows if r[:2] == (period, kind)}, set(codes)
                        )
                self.assertEqual(len(rows), 64)
            finally:
                book.close()
