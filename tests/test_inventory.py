"""Contratos do estoque: cenários sintéticos, precisão e interação Qt real."""

import gzip
import hashlib
import json
import statistics
import tempfile
import unittest
import zipfile
from dataclasses import replace
from datetime import date
from pathlib import Path
from threading import Event
from unittest.mock import patch

from PySide6.QtCore import QMetaObject, QObject, Qt, QUrl
from PySide6.QtTest import QAbstractItemModelTester, QSignalSpy, QTest

from liarte_core.inventory import (
    SCALE,
    Balance,
    Closing,
    InventoryData,
    Product,
    Sale,
    StockSelection,
    analyze_inventory,
    closing_summary,
    month_range,
    scaled,
)
from liarte_core.inventory_export import export_inventory
from liarte_core.inventory_assets import verify_bundled_inventory
from liarte_core.inventory_search import ProductSearch
from liarte_core.inventory_io import read_pbix, read_snapshot, validate_data, write_snapshot
from liarte_core.query import Cancelled
from liarte_desktop.app import create_bundle
from liarte_desktop.inventory import InventoryController
from tests.test_qt import APP, click, pump


def sample():
    return InventoryData(
        date(2026, 4, 15),
        "Base sintética de teste.pbix",
        "a" * 64,
        (
            Product("0001", "Ácido acrílico", "Químicos", "MP", "KG"),
            Product("0002", "Solvente", "Solventes", "PA", "L"),
            Product("0003", "Embalagem", "Embalagens", "EM", "UN"),
        ),
        (
            Balance("0001", "01", "01", "KG", 100 * SCALE, 10000),
            Balance("0001", "01", "02", "KG", 50 * SCALE, 6000),
            Balance("0002", "01", "01", "L", 30 * SCALE, 9000),
            Balance("0003", "02", "01", "UN", -2 * SCALE, -200),
        ),
        (
            Closing("0001", "01", "01", "KG", "2026-01", 80 * SCALE, 8000, 2.0, 15.0),
            Closing("0001", "01", "01", "KG", "2026-03", 60 * SCALE, 6000, 3.0, 10.0),
            Closing("0002", "01", "01", "L", "2026-01", 30 * SCALE, 9000, None, None),
            Closing("0002", "01", "01", "L", "2026-03", 30 * SCALE, 10000, None, None),
        ),
        (
            Sale("0001", "01", "KG", date(2026, 1, 1), 10 * SCALE, 1000),
            Sale("0001", "01", "KG", date(2026, 3, 30), 30 * SCALE, 3000),
            Sale("0002", "01", "L", date(2026, 1, 1), 5 * SCALE, 100),
            Sale("0002", "01", "L", date(2026, 3, 1), 10 * SCALE, 300),
        ),
        (("01", "Filial de teste"), ("02", "Filial secundária")),
        (("01", "Principal"), ("02", "Secundário")),
    )


def selection(**kwargs):
    return StockSelection("2026-01", "2026-03", **kwargs)


class InventoryMathTests(unittest.TestCase):
    def test_completer_ranks_exact_codes_preserves_zeroes_and_ignores_accents(self):
        finder = ProductSearch(sample())
        self.assertEqual(finder.find("0001")["items"][0]["code"], "0001")
        self.assertEqual(finder.find("ACIDO acrilico")["items"][0]["name"], "Ácido acrílico")
        self.assertEqual(finder.find("quimicos 0001")["count"], 1)
        self.assertEqual(finder.find("acidosulf")["count"], 0)
        self.assertEqual(len(finder.find("", 1)["items"]), 1)
        self.assertEqual(finder.find("0001")["items"][0]["balance"], "150,000 KG")

    def test_monthly_change_never_bridges_gaps_or_uses_unavailable_october(self):
        result = analyze_inventory(sample(), selection())
        self.assertEqual(result.history_summary["period"], "Fechamento · 03/2026")
        self.assertIsNone(result.history_summary["changeValue"])
        points = [
            dict(period="2026-01", label="01/2026", raw=10000),
            dict(period="2026-02", label="02/2026", raw=12500),
            dict(period="2026-03", label="03/2026", raw=None),
        ]
        summary = closing_summary(points)
        self.assertEqual(summary["changeValue"], 25)
        self.assertIn("+25,0%", summary["change"])
        self.assertEqual(summary["period"], "Fechamento · 02/2026")
        points[0]["raw"] = 0
        self.assertIsNone(closing_summary(points)["changeValue"])

    def test_sales_months_keep_observed_zeroes_and_distinguish_missing_and_partial(self):
        data = replace(sample(), sales=sample().sales[1:])
        # Global coverage starts on January 1 because the other product has a sale.
        result = analyze_inventory(data, StockSelection("2025-12", "2026-04", product="0001"))
        months = {p["period"]: p for p in result.sales_months}
        self.assertIsNone(months["2025-12"]["quantity"])
        self.assertEqual(months["2026-02"]["quantity"], 0)
        self.assertEqual(months["2026-03"]["quantity"], 30 * SCALE)
        self.assertTrue(months["2026-04"]["partial"])
        self.assertFalse(months["2026-03"]["partial"])

    def test_future_sales_months_are_unobserved_instead_of_zero(self):
        result = analyze_inventory(sample(), StockSelection("2026-03", "2026-06", product="0001"))
        months = {p["period"]: p for p in result.sales_months}
        self.assertEqual(months["2026-03"]["quantity"], 30 * SCALE)
        self.assertEqual(months["2026-04"]["quantity"], 0)
        self.assertTrue(months["2026-04"]["partial"])
        for period in ("2026-05", "2026-06"):
            with self.subTest(period=period):
                self.assertIsNone(months[period]["quantity"])
                self.assertTrue(months[period]["partial"])

    def test_ranking_share_has_the_same_denominator_and_order_as_abc(self):
        result = analyze_inventory(sample(), selection())
        self.assertAlmostEqual(sum(p["shareValue"] for p in result.ranking), 100)
        self.assertAlmostEqual(result.ranking[-1]["cumulativeValue"], 100)
        self.assertEqual(result.ranking[0]["code"], "0001")
        self.assertEqual(result.ranking[0]["abc"], "A")

    def test_currency_rounding_and_quantities_use_decimal(self):
        self.assertEqual(scaled("1.005", 100), 101)
        self.assertEqual(scaled("-1.005", 100), -101)
        self.assertEqual(scaled("0.000001", SCALE), 1)
        for value in ("NaN", "Infinity", "-Infinity", "abc"):
            with self.subTest(value=value), self.assertRaises(ValueError):
                scaled(value, SCALE)

    def test_month_sequence_crosses_year_and_checks_interval(self):
        self.assertEqual(month_range("2025-12", "2026-02"), ("2025-12", "2026-01", "2026-02"))
        for start, end in (("2026-02", "2026-01"), ("2026-13", "2026-13"), ("2000-01", "2026-01")):
            with self.subTest(start=start), self.assertRaises(ValueError):
                month_range(start, end)

    def test_current_cost_preserves_negative_balances_and_does_not_sum_history(self):
        result = analyze_inventory(sample(), selection())
        self.assertEqual(result.total_cost, 24800)
        self.assertEqual(len(result.rows), 4)
        self.assertEqual(result.metrics[0]["value"], "R$ 248,00")
        self.assertEqual(result.trend[-1]["raw"], 16000)

    def test_mixed_units_never_become_a_fictitious_total(self):
        result = analyze_inventory(sample(), selection())
        self.assertIn("distintas", result.quantity_label)
        result = analyze_inventory(sample(), selection(unit="KG"))
        self.assertEqual(result.quantity_label, "150,000 KG")
        self.assertEqual(result.total_cost, 16000)

    def test_coverage_uses_entire_branch_without_duplicating_sales_by_warehouse(self):
        result = analyze_inventory(sample(), selection(product="0001"))
        self.assertEqual({r.sold for r in result.rows}, {40 * SCALE})
        self.assertEqual({r.coverage for r in result.rows}, {337.5})
        self.assertEqual(sum(m["quantity"] for m in result.sales_months), 40 * SCALE)

    def test_warehouse_filter_disables_unallocatable_coverage(self):
        result = analyze_inventory(sample(), selection(product="0001", warehouse="01"))
        self.assertEqual(len(result.rows), 1)
        self.assertIsNone(result.rows[0].coverage)
        self.assertEqual(result.rows[0].sold, 40 * SCALE)

    def test_months_without_sales_are_zero_and_cv_is_sample_std(self):
        result = analyze_inventory(sample(), selection(product="0001"))
        row = result.rows[0]
        self.assertEqual(row.full_months, 3)
        self.assertAlmostEqual(row.monthly_mean, 40 / 3)
        self.assertAlmostEqual(row.monthly_std, statistics.stdev([10, 0, 30]))
        self.assertAlmostEqual(row.cv, statistics.stdev([10, 0, 30]) / (40 / 3))

    def test_incomplete_month_excluded_from_cv_but_used_in_daily_coverage(self):
        selected = StockSelection("2026-02", "2026-04", product="0001")
        result = analyze_inventory(sample(), selected)
        self.assertEqual(result.rows[0].full_months, 2)
        self.assertIsNone(result.rows[0].cv)
        self.assertAlmostEqual(result.rows[0].coverage, 150 * 74 / 30)

    def test_last_positive_sale_survives_same_day_returns(self):
        data = replace(
            sample(),
            sales=sample().sales
            + (Sale("0001", "01", "KG", date(2026, 4, 1), -5 * SCALE, -500, True),),
        )
        result = analyze_inventory(data, selection(product="0001"))
        self.assertEqual(result.rows[0].last_sale, date(2026, 4, 1))
        self.assertEqual(result.rows[0].days_since_sale, 14)

    def test_source_giro_is_latest_not_sum_or_average_across_months(self):
        result = analyze_inventory(sample(), selection(product="0001"))
        self.assertEqual(result.rows[0].turnover, 3.0)
        self.assertEqual(result.rows[0].source_period, "2026-03")
        self.assertIsNone(result.rows[1].turnover)
        self.assertEqual(result.metrics[3]["value"], "3,00 x")

    def test_missing_closing_is_not_zero_or_interpolated(self):
        result = analyze_inventory(sample(), selection())
        self.assertIsNone(result.trend[1]["raw"])
        self.assertIsNone(result.trend[1]["value"])
        self.assertIsNone(result.trend[2]["movingValue"])

    def test_source_scope_is_disclosed_and_unmatched_branch_has_no_history(self):
        result = analyze_inventory(sample(), selection())
        self.assertEqual(result.historical_scope, "Filial 01 · arm. 01")
        result = analyze_inventory(sample(), selection(branch="02"))
        self.assertTrue(all(p["raw"] is None for p in result.trend))

    def test_search_without_accents_restricts_all_analysis(self):
        result = analyze_inventory(sample(), selection(search="acido 0001"))
        self.assertEqual({r.balance.product for r in result.rows}, {"0001"})
        self.assertEqual(result.total_cost, 16000)
        self.assertEqual(result.trend[-1]["raw"], 6000)

    def test_activity_and_balance_filters_preserve_positions(self):
        result = analyze_inventory(sample(), selection(activity="31–90 dias"))
        self.assertEqual({r.balance.product for r in result.rows}, {"0002"})
        result = analyze_inventory(sample(), selection(balance="negative"))
        self.assertEqual(result.total_cost, -200)
        self.assertEqual(len(result.rows), 1)
        self.assertTrue(all(p["raw"] is None for p in result.trend))

    def test_abc_uses_positive_capital_before_abc_filter(self):
        data = replace(
            sample(),
            balances=(
                Balance("0001", "01", "01", "KG", SCALE, 8000),
                Balance("0002", "01", "01", "L", SCALE, 1500),
                Balance("0003", "02", "01", "UN", SCALE, 500),
            ),
        )
        result = analyze_inventory(data, selection())
        self.assertEqual(
            {r.balance.product: r.abc for r in result.rows}, {"0001": "A", "0002": "B", "0003": "C"}
        )
        result = analyze_inventory(data, selection(abc="B"))
        self.assertEqual(result.rows[0].abc, "B")
        self.assertEqual(result.total_cost, 1500)

    def test_abc_and_product_count_use_net_product_balances(self):
        data = replace(
            sample(),
            balances=(
                Balance("0001", "01", "01", "KG", 10 * SCALE, 10000),
                Balance("0001", "01", "02", "KG", -20 * SCALE, -11000),
                Balance("0002", "01", "01", "L", SCALE, 9000),
            ),
        )
        result = analyze_inventory(data, selection())
        self.assertEqual(result.metrics[1]["value"], "1")
        self.assertEqual({r.abc for r in result.rows if r.balance.product == "0001"}, {"—"})
        self.assertEqual({r.abc for r in result.rows if r.balance.product == "0002"}, {"A"})

    def test_empty_recorte_is_valid_and_future_start_is_rejected(self):
        result = analyze_inventory(sample(), selection(search="inexistente"))
        self.assertFalse(result.rows)
        self.assertEqual(result.total_cost, 0)
        with self.assertRaises(ValueError):
            analyze_inventory(sample(), StockSelection("2026-05", "2026-06"))

    def test_nonpositive_monthly_demand_does_not_generate_cv(self):
        data = replace(
            sample(), sales=tuple(replace(s, quantity=-abs(s.quantity)) for s in sample().sales)
        )
        result = analyze_inventory(data, selection(product="0001"))
        self.assertIsNone(result.rows[0].coverage)
        self.assertIsNone(result.rows[0].cv)


class InventoryFileTests(unittest.TestCase):
    def test_embedded_manifest_rejects_corruption_and_missing_file(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "estoque.liarte"
            write_snapshot(sample(), path)
            path.with_suffix(".json").write_text(
                json.dumps({"sha256": hashlib.sha256(path.read_bytes()).hexdigest()}),
                encoding="utf-8",
            )
            verify_bundled_inventory(path)
            path.write_bytes(b"corrupto")
            with self.assertRaisesRegex(ValueError, "Reinstale"):
                verify_bundled_inventory(path)
            path.unlink()
            with self.assertRaisesRegex(ValueError, "Reinstale"):
                verify_bundled_inventory(path)

    def test_roundtrip_preserves_codes_dates_units_money_and_source(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "base.liarte"
            write_snapshot(sample(), path)
            self.assertEqual(read_snapshot(path), sample())
            self.assertEqual(read_snapshot(path).products[0].code, "0001")

    def test_duplicate_or_orphan_positions_and_invalid_values_are_rejected(self):
        for data in (
            replace(sample(), balances=sample().balances + (sample().balances[0],)),
            replace(sample(), products=(sample().products[0],)),
            replace(sample(), balances=(replace(sample().balances[0], quantity=True),)),
            replace(sample(), closings=(replace(sample().closings[0], turnover=float("nan")),)),
            replace(sample(), sales=(replace(sample().sales[0], day=date(2027, 1, 1)),)),
        ):
            with self.subTest(data=data), self.assertRaises(ValueError):
                validate_data(data)

    def test_duplicate_daily_sales_are_rejected_in_snapshots_and_preserve_previous_base(self):
        data = sample()
        duplicated = replace(data, sales=data.sales + (data.sales[0],))
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "base.liarte"
            write_snapshot(data, path)
            before = path.read_bytes()
            with self.assertRaisesRegex(ValueError, "Vendas duplicadas"):
                write_snapshot(duplicated, path)
            self.assertEqual(path.read_bytes(), before)
            self.assertEqual(list(Path(directory).glob("*.tmp")), [])

            # A portable file can have been produced outside our writer.
            payload = json.loads(gzip.decompress(before))
            payload["sales"].append(payload["sales"][0])
            path.write_bytes(gzip.compress(json.dumps(payload).encode("utf-8")))
            with self.assertRaisesRegex(ValueError, "Vendas duplicadas"):
                read_snapshot(path)

    def test_daily_sales_from_different_branches_or_units_remain_distinct(self):
        data = sample()
        first = data.sales[0]
        data = replace(
            data,
            sales=data.sales + (replace(first, branch="02"), replace(first, unit="L")),
        )
        self.assertIs(validate_data(data), data)

    def test_negative_source_ratios_are_preserved_instead_of_silently_changed(self):
        data = replace(
            sample(), closings=(replace(sample().closings[0], turnover=-2.0, days=-1.6e18),)
        )
        self.assertIs(validate_data(data), data)

    def test_cancelled_write_keeps_existing_snapshot_and_removes_temporary(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "base.liarte"
            write_snapshot(sample(), path)
            old = path.read_bytes()
            cancel = Event()
            cancel.set()
            with self.assertRaises(Cancelled):
                write_snapshot(sample(), path, cancel)
            self.assertEqual(path.read_bytes(), old)
            self.assertEqual(list(Path(directory).glob("*.tmp")), [])

    def test_atomic_replace_failure_keeps_previous_base(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "base.liarte"
            write_snapshot(sample(), path)
            old = path.read_bytes()
            with (
                patch(
                    "liarte_core.inventory_io.os.replace", side_effect=PermissionError("bloqueado")
                ),
                self.assertRaises(PermissionError),
            ):
                write_snapshot(sample(), path)
            self.assertEqual(path.read_bytes(), old)
            self.assertEqual(list(Path(directory).glob("*.tmp")), [])

    def test_corrupt_and_unknown_snapshots_rejected_without_execution(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "base.liarte"
            for body in (
                b"not gzip",
                gzip.compress(b'{"format":"other"}'),
                gzip.compress(b'{"format":'),
            ):
                path.write_bytes(body)
                with self.subTest(body=body), self.assertRaises(ValueError):
                    read_snapshot(path)

    def test_pbix_without_embedded_data_has_actionable_error(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "thin.pbix"
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr("Report/Layout", "{}")
            with self.assertRaisesRegex(ValueError, "não contém a base"):
                read_pbix(path)

    def test_pbix_projection_uses_column_names_not_schema_order_and_keeps_orphans(self):
        import pandas as pd

        tables = {
            "Dim_Atualizacao": pd.DataFrame({"Data Atualizacao": ["04/15/2026 12:00:00 PM"]}),
            "Dim_Produtos": pd.DataFrame(
                {
                    "CodProduto": ["0001"],
                    "NomeProduto": ["Ácido"],
                    "TipoProduto": ["MP"],
                    "UnidadeMedida": ["KG"],
                    "DELETED": [""],
                    "NomeGrupo": ["Químicos"],
                }
            ),
            "Dim_Filiais": pd.DataFrame(
                {"CodFilial": ["01"], "NomeFilial": ["Matriz"], "DELETED": [""]}
            ),
            "Dim_Armazem": pd.DataFrame(
                {"CodArmazem": ["01"], "Descricao": ["Principal"], "DELETED": [""]}
            ),
            "Fato_EstoqueAtual": pd.DataFrame(
                {
                    "Filial": ["01"],
                    "Produto": ["0001"],
                    "Armazem": ["01"],
                    "SaldoFechamento": [100.0],
                    "CustoTotal": [200.0],
                    "UnidadeMedida": ["KG"],
                    "DELETED": [""],
                }
            ),
            "Fact_Estoque": pd.DataFrame(
                {
                    "Filial": ["01"],
                    "Produto": ["ANTIGO"],
                    "Armazem": ["01"],
                    "SaldoFechamento": [50.0],
                    "CustoTotal": [100.0],
                    "UnidadeMedida": ["KG"],
                    "DataFechamento": [pd.Timestamp("2026-03-31")],
                    "DELETED": [""],
                    "DiasEstoqueQtd": [15.0],
                    "GiroEstoqueMensalQtd": [2.0],
                }
            ),
            "Fato_DocumentoSaida": pd.DataFrame(
                {
                    "Filial": ["01", "01"],
                    "Produto": ["0001", "0001"],
                    "UnidadeMedida": ["KG", "KG"],
                    "DataEmissao": [pd.Timestamp("2026-04-01")] * 2,
                    "Quantidade": [5.0, -10.0],
                    "CustoTotal": [10.0, -20.0],
                    "DELETED": ["", ""],
                }
            ),
        }

        class FakeModel:
            def __init__(self, path):
                self.tables = list(tables)

            def __enter__(self):
                return self

            def __exit__(self, *args):
                pass

            def get_table(self, name, columns):
                # PBIXRay mantém a ordem original do modelo na projeção.
                return tables[name].loc[:, [c for c in tables[name].columns if c in columns]]

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "test.pbix"
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr("DataModel", b"test")
            before = path.read_bytes()
            with patch("pbixray.PBIXRay", FakeModel):
                data = read_pbix(path)
            self.assertEqual(path.read_bytes(), before)
            self.assertEqual(data.balances[0].product, "0001")
            self.assertEqual(data.balances[0].cost, 20000)
            self.assertEqual(data.closings[0].turnover, 2.0)
            self.assertEqual(data.products[-1].code, "ANTIGO")
            self.assertTrue(data.notes)
            self.assertEqual(data.sales[0].quantity, -5 * SCALE)
            self.assertTrue(data.sales[0].had_positive)

    def test_excel_matches_filtered_result_units_formats_and_escapes_formula_text(self):
        from openpyxl import load_workbook

        data = replace(
            sample(), products=(replace(sample().products[0], name="=1+1"), *sample().products[1:])
        )
        selected = selection(unit="KG")
        analysis = analyze_inventory(data, selected)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "Estoque.xlsx"
            export_inventory(data, analysis, selected, path, Event())
            book = load_workbook(path)
            try:
                sheet = book["Estoque"]
                self.assertEqual(sheet.max_row, 5)
                self.assertEqual(sum(sheet.cell(r, 8).value for r in range(4, 6)), 160.0)
                self.assertEqual(sheet.cell(4, 2).data_type, "s")
                self.assertEqual(sheet.cell(4, 2).value, "=1+1")
                self.assertEqual(sheet.cell(4, 1).value, "0001")
                self.assertIn("R$", sheet.cell(4, 8).number_format)
                self.assertIsNone(book["Fechamentos"].cell(3, 2).value)
                self.assertTrue(
                    all(c.alignment.vertical == "center" for ws in book for row in ws for c in row)
                )
            finally:
                book.close()


class InventoryUiTests(unittest.TestCase):
    def test_first_use_loads_embedded_base_and_recovers_invalid_cache(self):
        self.base.with_suffix(".json").write_text(
            json.dumps({"sha256": hashlib.sha256(self.base.read_bytes()).hexdigest()}),
            encoding="utf-8",
        )
        for corrupt in (False, True):
            cache = Path(self.temp.name) / (
                "invalid/base.liarte" if corrupt else "fresh/base.liarte"
            )
            if corrupt:
                cache.parent.mkdir()
                cache.write_bytes(b"corrupto")
            controller = InventoryController(cache_path=cache, bundled_path=self.base)
            try:
                controller.ensureLoaded()
                self.assertTrue(pump(lambda: not controller.state["busy"], 15))
                self.assertTrue(controller.state["loaded"], controller.state["error"])
                self.assertEqual(read_snapshot(cache).balances, sample().balances)
                self.assertEqual(controller.state["error"], "")
                if corrupt:
                    self.assertIn("recuperada", " ".join(controller.state["notes"]))
            finally:
                controller.shutdown()

    def test_existing_user_base_takes_precedence_over_embedded_snapshot(self):
        cache = Path(self.temp.name) / "newer.liarte"
        newer = replace(
            sample(), as_of=date(2026, 5, 1), source_name="Base atualizada pelo usuário.pbix"
        )
        write_snapshot(newer, cache)
        # Integrated resource intentionally invalid: a good user cache is sufficient.
        controller = InventoryController(cache_path=cache, bundled_path=self.base)
        try:
            controller.ensureLoaded()
            self.assertTrue(pump(lambda: not controller.state["busy"], 15))
            self.assertEqual(controller._data, newer)
            self.assertEqual(read_snapshot(cache), newer)
        finally:
            controller.shutdown()

    def test_suggestions_support_arrow_selection_escape_and_clear_without_losing_text(self):
        self.load()
        QTest.keyClick(self.window, Qt.Key.Key_F, Qt.KeyboardModifier.ControlModifier)
        for letter in "ACIDO":
            QTest.keyClick(self.window, Qt.Key(ord(letter)))
        self.assertTrue(pump(lambda: self.item("stockSuggestions").property("visible")))
        QTest.keyClick(self.window, Qt.Key.Key_Down)
        QTest.keyClick(self.window, Qt.Key.Key_Return)
        self.idle()
        self.assertEqual(self.c.state["selection"]["product"], "0001")
        self.assertEqual(self.c.state["selection"]["search"], "")
        self.assertIn("Ácido acrílico", self.item("stockSearch").property("text"))
        QTest.keyClick(self.window, Qt.Key.Key_F, Qt.KeyboardModifier.ControlModifier)
        QTest.keyClick(self.window, Qt.Key.Key_S)
        self.idle()
        self.assertEqual(self.c.state["selection"]["product"], "")
        self.assertEqual(self.c.state["selection"]["search"], "s")
        QTest.keyClick(self.window, Qt.Key.Key_Escape)
        self.assertFalse(self.item("stockSuggestions").property("visible"))
        click(self.window, self.item("clearStockSearch"))
        self.idle()
        self.assertEqual(self.item("stockSearch").property("text"), "")
        self.assertEqual(self.c.state["count"], 4)
        for letter in "ZZZZZZ":
            QTest.keyClick(self.window, Qt.Key(ord(letter)))
        QTest.keyClick(self.window, Qt.Key.Key_Down)
        QTest.keyClick(self.window, Qt.Key.Key_Return)
        self.idle()
        self.assertFalse(self.item("stockSuggestions").property("visible"))
        self.assertEqual(self.c.state["count"], 0)
        self.assertEqual(self.bundle.messages, [])

    def test_chart_mode_values_average_and_nice_axis_do_not_change_the_selection(self):
        self.load()
        chart = self.item("stockTrend")
        before = self.c.state["selection"]
        chart.setProperty("mode", "column")
        chart.setProperty("showValues", True)
        chart.setProperty("movingAverage", False)
        chart.setProperty("points", [{"period": "2026-01", "label": "Jan", "value": 9}])
        APP.processEvents()
        self.assertEqual(chart.property("tickDigits"), 1)
        chart.setProperty(
            "points",
            [
                {"period": "2026-01", "label": "Jan", "value": -10},
                {"period": "2026-02", "label": "Fev", "value": None},
                {"period": "2026-03", "label": "Mar", "value": 20},
            ],
        )
        APP.processEvents()
        self.assertTrue(chart.property("hasValues"))
        self.assertEqual(self.c.state["selection"], before)
        self.assertEqual(self.bundle.messages, [])

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.base = Path(self.temp.name) / "base.liarte"
        write_snapshot(sample(), self.base)
        self.bundle = create_bundle(size=50, initial_module="inventory")
        self.c = self.bundle.inventory
        self.window = self.bundle.engine.rootObjects()[0]

    def tearDown(self):
        self.c.shutdown()
        self.bundle.controller.shutdown()
        self.window.close()
        self.bundle.engine.deleteLater()
        pump(lambda: False, 0.03)
        self.temp.cleanup()

    def item(self, name):
        result = self.window.findChild(QObject, name)
        self.assertIsNotNone(result, name)
        return result

    def idle(self):
        self.assertTrue(
            pump(lambda: not self.c.state["busy"] and not self.c._debounce.isActive(), 15),
            self.c.state,
        )

    def load(self):
        self.c.openFile(QUrl.fromLocalFile(str(self.base)))
        self.idle()
        self.assertTrue(self.c.state["loaded"], self.c.state["error"])

    def test_module_is_available_and_load_does_not_query_comex(self):
        with patch.object(self.bundle.controller._provider, "query") as query:
            self.load()
            query.assert_not_called()
        self.assertEqual(self.window.title(), "Liarte Estoque")
        self.assertEqual(self.c.state["count"], 4)
        self.assertEqual(self.bundle.messages, [])

    def test_live_filters_numeric_sort_and_reset(self):
        self.load()
        self.c.setFilter("unit", "KG")
        self.idle()
        self.assertEqual(self.c.state["count"], 2)
        self.assertEqual({r.balance.unit for r in self.c.table.rows}, {"KG"})
        self.c.sortByColumn(5)
        self.assertEqual(self.c.table.rows[0].balance.quantity, 100 * SCALE)
        self.c.sortByColumn(5)
        self.assertEqual(self.c.table.rows[0].balance.quantity, 50 * SCALE)
        self.c.resetFilters()
        self.idle()
        self.assertEqual(self.c.state["count"], 4)

    def test_latest_filter_wins_while_previous_analysis_runs(self):
        self.load()
        entered, release = Event(), Event()
        original = analyze_inventory

        def delayed(*args, **kwargs):
            entered.set()
            if not release.wait(5):
                raise RuntimeError("Teste não liberou a tarefa")
            return original(*args, **kwargs)

        try:
            with patch("liarte_desktop.inventory.analyze_inventory", side_effect=delayed):
                self.c.setFilter("unit", "KG")
                self.assertTrue(pump(entered.is_set))
                self.c.setFilter("unit", "L")
                release.set()
                self.idle()
            self.assertEqual({r.balance.unit for r in self.c.table.rows}, {"L"})
            self.assertTrue(self.c.state["current"])
        finally:
            release.set()

    def test_cancel_reverts_unapplied_filters_and_keeps_previous_rows(self):
        self.load()
        before = self.c.table.rows
        self.c.setFilter("unit", "KG")
        self.c.cancel()
        self.idle()
        self.assertEqual(self.c.table.rows, before)
        self.assertEqual(self.c.state["selection"]["unit"], "")
        self.assertTrue(self.c.state["current"])

    def test_bad_source_and_bad_interval_never_replace_valid_results(self):
        self.load()
        data = self.c._data
        before = self.c.table.rows
        self.c.openFile(QUrl.fromLocalFile(str(self.base.with_name("missing.liarte"))))
        self.idle()
        self.assertIs(self.c._data, data)
        self.assertEqual(self.c.table.rows, before)
        self.assertTrue(self.c.state["error"])
        self.c.setFilter("start", "2026-04")
        self.c.setFilter("end", "2026-01")
        self.idle()
        self.assertTrue(self.c.state["error"])
        self.assertFalse(self.c.state["current"])
        output = Path(self.temp.name) / "must-not-export.xlsx"
        self.c.exportExcel(QUrl.fromLocalFile(str(output)))
        self.assertFalse(output.exists())

    def test_module_roundtrip_preserves_stock_and_import_sessions(self):
        self.load()
        self.c.setFilter("unit", "KG")
        self.idle()
        workspace = self.item("inventoryWorkspace")
        model = self.c.table
        data = self.c._data
        reset = QSignalSpy(model.modelReset)
        self.bundle.platform.openModule("imports")
        self.bundle.controller.queryPeriod("2026-04", "2026-09")
        self.assertTrue(pump(lambda: not self.bundle.controller.state["busy"]))
        imported = self.bundle.controller._result
        self.bundle.platform.openModule("home")
        self.bundle.platform.openModule("inventory")
        APP.processEvents()
        self.assertIs(self.item("inventoryWorkspace"), workspace)
        self.assertIs(self.c._data, data)
        self.assertEqual(reset.count(), 0)
        self.assertIs(self.bundle.controller._result, imported)
        self.assertEqual(self.c.state["selection"]["unit"], "KG")
        self.assertEqual(self.bundle.messages, [])

    def test_product_detail_exposes_scope_and_demand_series(self):
        self.load()
        self.c.setFilter("product", "0001")
        self.idle()
        self.c.selectRow(0)
        detail = self.c.state["detail"]
        self.assertIn("filial inteira", detail["sold"])
        self.assertEqual(detail["code"], "0001")
        self.assertTrue(detail["salesTrend"])
        self.assertTrue(detail["giroTrend"])
        QMetaObject.invokeMethod(self.item("stockDetail"), "open")
        self.assertTrue(pump(lambda: self.item("stockDetail").property("visible")))
        self.assertEqual(self.bundle.messages, [])
        QMetaObject.invokeMethod(self.item("stockDetail"), "close")

    def test_rapid_theme_layout_and_advanced_filters_have_no_qml_warnings(self):
        self.load()
        tester = QAbstractItemModelTester(
            self.c.table, QAbstractItemModelTester.FailureReportingMode.Warning
        )
        for dark in (False, True):
            self.bundle.preferences.setPreference("dark", dark)
            for w, h in ((500, 400), (760, 600), (1380, 900)):
                self.window.resize(w, h)
                APP.processEvents()
                self.assertFalse(self.window.grabWindow().isNull())
        self.assertEqual(self.bundle.messages, [])
        self.assertIsNotNone(tester)

    def test_compact_preference_changes_actual_row_height_without_resetting_model(self):
        self.load()
        table = self.item("stockTable")

        def cell(item):
            if item.objectName() == "stockCell_0_0":
                return item
            for child in item.childItems():
                found = cell(child)
                if found is not None:
                    return found
            return None

        self.assertTrue(pump(lambda: cell(table) is not None))
        self.assertEqual(cell(table).height(), 40)
        resets = QSignalSpy(self.c.table.modelReset)
        self.bundle.preferences.setPreference("compact", True)
        self.assertTrue(pump(lambda: cell(table).height() == 32))
        self.assertEqual(resets.count(), 0)
        self.window.resize(500, 400)
        self.assertTrue(pump(lambda: table.height() <= 150))
        self.assertEqual(self.bundle.messages, [])

    def test_cached_source_restored_only_when_entering_stock(self):
        cached = InventoryController(cache_path=self.base)
        try:
            self.assertFalse(cached.state["loaded"])
            cached.ensureLoaded()
            self.assertTrue(pump(lambda: not cached.state["busy"]))
            self.assertTrue(cached.state["loaded"])
            data = cached._data
            cached.ensureLoaded()
            self.assertIs(cached._data, data)
        finally:
            cached.shutdown()

    def test_cancel_after_staging_never_replaces_previous_cached_base(self):
        cache_path = Path(self.temp.name) / "cache/base.liarte"
        write_snapshot(sample(), cache_path)
        previous = cache_path.read_bytes()
        staged, release = Event(), Event()
        original = write_snapshot

        def delayed(*args):
            original(*args)
            staged.set()
            if not release.wait(5):
                raise RuntimeError("Teste não liberou a tarefa")

        cached = InventoryController(cache_path=cache_path)
        try:
            with patch("liarte_desktop.inventory.write_snapshot", side_effect=delayed):
                cached.openFile(QUrl.fromLocalFile(str(self.base)))
                self.assertTrue(pump(staged.is_set))
                cached.cancel()
                release.set()
                self.assertTrue(pump(lambda: not cached.state["busy"]))
            self.assertFalse(cached.state["loaded"])
            self.assertEqual(cache_path.read_bytes(), previous)
            self.assertEqual(list(cache_path.parent.glob(".base-*.liarte")), [])
        finally:
            release.set()
            cached.shutdown()

    def test_keyboard_route_reaches_stock_without_loading_another_module(self):
        self.bundle.platform.openModule("home")
        QTest.keyClick(self.window, Qt.Key.Key_3, Qt.KeyboardModifier.ControlModifier)
        self.assertTrue(pump(lambda: self.bundle.platform.activeModule == "inventory"))
        self.assertFalse(self.bundle.controller.state["loaded"])
        self.assertEqual(self.bundle.messages, [])

    def test_export_preserves_selected_row_model_and_filter(self):
        self.load()
        self.c.setFilter("unit", "KG")
        self.idle()
        self.c.selectRow(0)
        before = self.c.table.rows
        spy = QSignalSpy(self.c.table.modelReset)
        path = Path(self.temp.name) / "Estoque.xlsx"
        self.c.exportExcel(QUrl.fromLocalFile(str(path)))
        self.idle()
        self.assertTrue(path.is_file())
        self.assertEqual(spy.count(), 0)
        self.assertEqual(self.c.table.rows, before)
        self.assertEqual(self.c.state["selectedRow"], 0)
        self.assertEqual(self.c.state["selection"]["unit"], "KG")

    def test_row_selection_and_sort_do_not_repaint_or_animate_the_overview_chart(self):
        self.load()
        pump(lambda: False, 0.30)
        chart = self.item("stockTrend")
        paints = chart.property("curvePaintCount")
        signals = QSignalSpy(self.c.analysisChanged)
        self.c.selectRow(0)
        self.c.sortByColumn(5)
        pump(lambda: False, 0.08)
        self.assertEqual(signals.count(), 0)
        self.assertEqual(chart.property("curvePaintCount"), paints)

    def test_search_keyboard_and_reset_keep_text_and_filters_synchronized(self):
        self.load()
        QTest.keyClick(self.window, Qt.Key.Key_F, Qt.KeyboardModifier.ControlModifier)
        for letter in "ACIDO":
            QTest.keyClick(self.window, Qt.Key(ord(letter)))
        self.idle()
        self.assertEqual({r.balance.product for r in self.c.table.rows}, {"0001"})
        self.c.resetFilters()
        self.idle()
        self.assertEqual(self.item("stockSearch").property("text"), "")
        self.assertEqual(self.bundle.messages, [])

    def test_closing_running_stock_requires_cancel_and_finishes_cleanly(self):
        self.load()
        entered, release = Event(), Event()
        original = analyze_inventory

        def delayed(*args, **kwargs):
            entered.set()
            if not release.wait(5):
                raise RuntimeError("Teste não liberou a tarefa")
            return original(*args, **kwargs)

        try:
            with patch("liarte_desktop.inventory.analyze_inventory", side_effect=delayed):
                self.c.setFilter("unit", "KG")
                self.assertTrue(pump(entered.is_set))
                self.window.close()
                self.assertTrue(pump(lambda: self.item("closeDialog").property("opened")))
                self.assertTrue(self.window.isVisible())
                release.set()
                click(self.window, self.item("confirmClose"))
                self.assertTrue(pump(lambda: not self.c.state["busy"]))
                self.assertTrue(pump(lambda: not self.window.isVisible()))
        finally:
            release.set()

    def test_open_and_save_dialogs_load_without_qml_errors(self):
        self.load()
        for name in ("stockSourceDialog", "stockSaveDialog"):
            dialog = self.item(name)
            QMetaObject.invokeMethod(dialog, "open")
            self.assertTrue(pump(lambda: dialog.property("visible")))
            QMetaObject.invokeMethod(dialog, "close")
            self.assertTrue(pump(lambda: not dialog.property("visible")))
        self.assertEqual(self.bundle.messages, [])
