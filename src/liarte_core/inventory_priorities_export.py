"""Export every classified product/unit with numerical criteria and references."""

from __future__ import annotations

from dataclasses import fields
from threading import Event

from .inventory import SCALE, StockSelection
from .query import checkpoint

_MAX_ROWS = 1_048_576
_TITLE = "Prioridades dos produtos"
_LABELS = {
    "coverage": "Cobertura < limite",
    "decline": "Variação de vendas < 0",
    "idle": "Saldo > 0 · vendas líquidas = 0",
    "capital": "Estoque a custo",
}
_HEADERS = (
    "Código", "Produto", "UM", "Critérios quantitativos",
    "Mês do fechamento", "Quantidade no fechamento",
    "Estoque a custo (R$)", "Janela anterior", "Janela recente",
    "Vendas líquidas anteriores", "Vendas líquidas recentes", "Diferença nas vendas",
    "Variação (%)", "Cobertura estimada (dias)", "Disponibilidade da cobertura",
    "Limite de cobertura (dias)", "Giro fonte · mediana (x/mês)",
    "Período do recorte", "Fonte, base e filtros", "Abrangência",
    "Observações dos dados", "Comparação disponível",
)
_WIDTHS = (
    18, 48, 12, 50, 20, 28, 27, 27, 27, 28, 28, 27, 20,
    27, 65, 24, 31, 27, 80, 70, 75, 24,
)
_DEFINITION = (
    "Uma linha por produto e unidade, mesmo quando há mais de uma categoria. "
    "As contagens das categorias não devem ser somadas. Quantidades permanecem "
    "na unidade indicada. Saldo e capital pertencem ao fechamento, não ao saldo atual. "
    "Estoque a custo não é faturamento. Vendas são líquidas: zero pode resultar de "
    "saídas compensadas por devoluções e não comprova ausência de vendas. "
    "Cobertura estimada não mede idade de lote nem confirma ruptura ou necessidade "
    "de compra. Variação de vendas não identifica sua causa. Estoque a custo é uma "
    "ordenação por custo positivo, sem limiar que defina excesso ou capital parado. "
    "Ausências permanecem vazias; zero conhecido continua numérico. O arquivo inclui "
    "todas as prioridades e unidades do recorte, sem os limites de busca, categoria "
    "ou paginação da tela."
)


def _selection_filters(report):
    if report.get("filters") is not None:
        if not isinstance(report["filters"], str):
            raise ValueError("A referência dos filtros das prioridades é inválida.")
        return report["filters"]
    key = report.get("selectionKey")
    if key is None:
        return ""
    selection_fields = fields(StockSelection)
    if not isinstance(key, (tuple, list)) or len(key) != len(selection_fields):
        raise ValueError("A referência dos filtros das prioridades é inválida.")
    return "; ".join(
        f"{field.name}={value}"
        for field, value in zip(selection_fields, key, strict=True) if value
    )


def export_product_priorities(workbook, report, *, cancel=None):
    """Append complete classified rows and return layouts for the stock exporter.

    The report is the accepted domain result, not a QML page or locally filtered
    presentation. The caller retains ownership of atomic workbook publication.
    This helper also keeps source text literal when used without that caller's
    final formatting pass.
    """
    from openpyxl.comments import Comment
    from .inventory_export import _append_values, _excel_scaled

    cancel = cancel if cancel is not None else Event()
    checkpoint(cancel)
    filters = _selection_filters(report)
    period = " a ".join(
        str(report[key]) for key in ("selectionStart", "selectionEnd") if report.get(key)
    )
    source_reference = "; ".join(value for value in (
        "Fonte: " + str(report["sourceName"]) if report.get("sourceName") else "",
        "Data da base: " + str(report["asOf"]) if report.get("asOf") else "",
        "Filtros: " + filters if filters else "",
    ) if value)
    definition = _DEFINITION
    labels = dict(_LABELS)
    for category in report.get("categories", ()):
        checkpoint(cancel)
        labels[category["key"]] = category["label"]
    context = "\n".join(str(value) for value in (
        "Fechamento: " + str(report.get("targetPeriod") or ""),
        "Janela anterior: " + str(report.get("baseLabel") or ""),
        "Janela recente: " + str(report.get("targetLabel") or ""),
        "Limite de cobertura: " + str(report.get("coverageDays") or "") + " dias",
        report.get("coveragePolicy"),
        source_reference, report.get("scope"), report.get("notice"), report.get("reason"), definition,
    ) if value)
    layouts = {}
    sheet, row_number, part = None, 0, 0

    def new_sheet():
        nonlocal sheet, row_number, part
        part += 1
        title = _TITLE + (f" {part}" if part > 1 else "")
        sheet = workbook.create_sheet(title)
        sheet.append(_HEADERS)
        sheet.cell(1, len(_HEADERS)).comment = Comment(context, "LIARTE")
        sheet.freeze_panes = "A2"
        row_number = 1
        layouts[sheet.title] = _WIDTHS

    for row in report.get("rows", ()):
        checkpoint(cancel)
        if sheet is None or row_number >= _MAX_ROWS:
            new_sheet()
        notices = " ".join(dict.fromkeys(
            str(value) for value in (
                row.get("dataNotice"), row.get("reason"), report.get("notice"),
                report.get("reason"), report.get("referenceNotice"),
            ) if value
        ))
        row_number += 1
        values = [
            row["code"], row["name"], row["unit"],
            "; ".join(labels.get(key, key) for key in row.get("categories", ())),
            report.get("targetPeriod", ""),
            _excel_scaled(row.get("stockQuantity"), SCALE),
            _excel_scaled(row.get("stockCost"), 100),
            report.get("baseLabel", ""), report.get("targetLabel", ""),
            _excel_scaled(row.get("previous"), SCALE),
            _excel_scaled(row.get("sold"), SCALE), _excel_scaled(row.get("change"), SCALE),
            row.get("percentage"), row.get("coverage"), row.get("coverageReason", ""),
            report.get("coverageDays"), row.get("turnover"), period, source_reference,
            report.get("scope", ""), notices,
            "Sim" if report.get("comparisonAvailable") is True else "Não",
        ]
        _append_values(sheet, values, row_number)
        for column in (6, 10, 11, 12):
            sheet.cell(row_number, column).number_format = "#,##0.000;[Red](#,##0.000)"
        sheet.cell(row_number, 7).number_format = '"R$" #,##0.00;[Red]("R$" #,##0.00)'
        for column in (13, 14, 17):
            sheet.cell(row_number, column).number_format = "0.00"
        for column, value in enumerate(values, 1):
            if isinstance(value, str):
                cell = sheet.cell(row_number, column)
                cell.data_type = "s"
                if cell.comment is not None and cell.comment.author == "LIARTE":
                    cell.number_format = "@"
        sheet.auto_filter.ref = f"A1:V{row_number}"
    if sheet is None:
        new_sheet()
        sheet.auto_filter.ref = "A1:V1"
    checkpoint(cancel)
    return layouts
